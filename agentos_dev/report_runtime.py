"""受限 Markdown 报表运行时；由 WorkspaceReportToolkit 在 Daytona 中执行。"""

from __future__ import annotations

import base64
import ctypes
import fcntl
import hashlib
import json
import os
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import uuid
from contextlib import contextmanager
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import unquote, urlsplit

MAX_MARKDOWN_BYTES = 1024 * 1024
MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_IMAGE_COUNT = 50
MAX_TOTAL_IMAGE_BYTES = 50 * 1024 * 1024
MAX_RESULT_BYTES = 64 * 1024
MAX_COMMAND_BYTES = 32 * 1024
MAX_ANALYSIS_OUTPUT_BYTES = 8 * 1024
MAX_DATASET_PATHS = 20
MAX_ANALYSIS_TIMEOUT = 60
MAX_PROFILE_ROWS = 100_000
MAX_PROFILE_COLUMNS = 50
MAX_PROFILE_TOTAL_COLUMNS = 100
MAX_PROFILE_TOP_VALUES = 5
MAX_PROFILE_JSON_BYTES = 25 * 1024 * 1024
IMAGE_SUFFIXES = {".gif", ".jpeg", ".jpg", ".png", ".webp"}
PROFILE_SUFFIXES = {".csv", ".json", ".jsonl", ".parquet", ".tsv", ".xls", ".xlsx"}
FAILURE_TTL_SECONDS = 24 * 60 * 60
MATPLOTLIBRC = """\
backend: Agg
font.family: sans-serif
font.sans-serif: Noto Sans CJK JP, DejaVu Sans
axes.unicode_minus: False
"""


class ReportFailure(ValueError):
    pass


def _relative_path(value: str, suffix: str | None = None) -> PurePosixPath:
    if not isinstance(value, str) or not value or "\\" in value:
        raise ReportFailure("路径必须是当前工作区的相对路径")
    path = PurePosixPath(value)
    if (
        path.is_absolute()
        or ".." in path.parts
        or any(part in ("", ".") for part in path.parts)
        or (suffix is not None and path.suffix.lower() != suffix)
    ):
        expected = f" {suffix}" if suffix is not None else ""
        raise ReportFailure(f"路径必须是当前工作区内的{expected}相对路径")
    return path


def _reject_symlinks(workspace: Path, path: Path) -> None:
    relative = path.relative_to(workspace)
    current = workspace
    for part in relative.parts:
        current /= part
        if current.is_symlink():
            raise ReportFailure("报表路径不能包含符号链接")


def _input_path(workspace: Path, value: str, suffix: str) -> Path:
    relative = _relative_path(value, suffix)
    path = workspace.joinpath(*relative.parts)
    _reject_symlinks(workspace, path)
    if not path.is_file():
        raise ReportFailure("报表源文件不存在或不是普通文件")
    return path


def _output_path(workspace: Path, value: str) -> Path:
    relative = _relative_path(value, ".pdf")
    path = workspace.joinpath(*relative.parts)
    parent = path.parent
    _reject_symlinks(workspace, parent)
    if not parent.is_dir():
        raise ReportFailure("PDF 输出目录不存在")
    if path.exists() or path.is_symlink():
        raise ReportFailure("PDF 输出文件已经存在")
    return path


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _json_scalar(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "item"):
        value = value.item()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, float) and (value != value or value in (float("inf"), float("-inf"))):
        return None
    if isinstance(value, (str, int, float, bool)):
        return value[:200] if isinstance(value, str) else value
    return str(value)[:200]


def _enable_child_subreaper() -> None:
    if sys.platform != "linux":
        return
    libc = ctypes.CDLL(None, use_errno=True)
    if libc.prctl(36, 1, 0, 0, 0) != 0:  # PR_SET_CHILD_SUBREAPER
        raise ReportFailure("无法启用分析进程隔离")


def _process_parents() -> dict[int, int]:
    result = {}
    for status in Path("/proc").glob("[0-9]*/status"):
        try:
            values = {
                key: value.strip()
                for key, value in (
                    line.split(":", 1) for line in status.read_text().splitlines() if ":" in line
                )
            }
            result[int(status.parent.name)] = int(values["PPid"])
        except (FileNotFoundError, KeyError, PermissionError, ValueError):
            continue
    return result


def _descendants(parent_pid: int, parents: dict[int, int]) -> set[int]:
    result: set[int] = set()
    changed = True
    while changed:
        changed = False
        for pid, parent in parents.items():
            if pid not in result and (parent == parent_pid or parent in result):
                result.add(pid)
                changed = True
    return result


def _terminate_descendants() -> None:
    if sys.platform != "linux":
        return
    parent_pid = os.getpid()
    for _attempt in range(8):
        parents = _process_parents()
        descendants = _descendants(parent_pid, parents)
        if not descendants:
            return
        for pid in descendants:
            try:
                os.kill(pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
        for pid, process_parent in parents.items():
            if process_parent != parent_pid or pid not in descendants:
                continue
            try:
                os.waitpid(pid, 0)
            except (ChildProcessError, ProcessLookupError):
                pass


def _execute_command(
    command: str,
    timeout: int,
    matplotlib_config_dir: str | None = None,
) -> dict[str, Any]:
    _enable_child_subreaper()
    environment = os.environ.copy()
    if matplotlib_config_dir is not None:
        environment["MPLCONFIGDIR"] = matplotlib_config_dir
        environment["MPLBACKEND"] = "Agg"
    process = subprocess.Popen(
        ["/bin/bash", "--noprofile", "--norc", "-c", command],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        start_new_session=True,
        env=environment,
    )
    try:
        output_bytes, _stderr = process.communicate(timeout=timeout)
        exit_code = process.returncode
    except subprocess.TimeoutExpired:
        os.killpg(process.pid, signal.SIGKILL)
        _terminate_descendants()
        output_bytes, _stderr = process.communicate()
        exit_code = 124
    finally:
        if process.poll() is None:
            os.killpg(process.pid, signal.SIGKILL)
            process.wait()
        _terminate_descendants()
    truncated = len(output_bytes) > MAX_ANALYSIS_OUTPUT_BYTES
    return {
        "exitCode": exit_code,
        "output": base64.b64encode(output_bytes[:MAX_ANALYSIS_OUTPUT_BYTES]).decode("ascii"),
        "truncated": truncated,
    }


def _run_supervised_command(
    command: str,
    cwd: Path,
    timeout: int,
    matplotlib_config_dir: Path,
) -> tuple[bytes, int, bool]:
    payload = json.dumps(
        {
            "command": command,
            "timeout": timeout,
            "matplotlib_config_dir": str(matplotlib_config_dir),
        },
        ensure_ascii=False,
    )
    supervisor = subprocess.Popen(
        [sys.executable, str(Path(__file__).resolve()), "_execute", payload],
        cwd=cwd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        start_new_session=True,
    )
    try:
        output_bytes, _stderr = supervisor.communicate(timeout=timeout + 5)
    except subprocess.TimeoutExpired as error:
        os.killpg(supervisor.pid, signal.SIGKILL)
        supervisor.wait()
        raise ReportFailure("分析 supervisor 未能按时退出") from error
    output = output_bytes.decode("utf-8", errors="replace")
    if supervisor.returncode != 0:
        raise ReportFailure(output or "分析 supervisor 执行失败")
    try:
        line = next(line for line in reversed(output.splitlines()) if line.strip())
        result = json.loads(line)
        command_output = base64.b64decode(result["output"], validate=True)
        exit_code = result["exitCode"]
        truncated = result["truncated"]
    except (StopIteration, KeyError, TypeError, ValueError, json.JSONDecodeError) as error:
        raise ReportFailure(output or "分析 supervisor 返回无效结果") from error
    if not isinstance(exit_code, int) or not isinstance(truncated, bool):
        raise ReportFailure("分析 supervisor 返回无效结果")
    return command_output, exit_code, truncated


def _check_image_signature(path: Path) -> None:
    header = path.read_bytes()[:12]
    suffix = path.suffix.lower()
    valid = {
        ".png": header.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": header.startswith(b"\xff\xd8\xff"),
        ".jpeg": header.startswith(b"\xff\xd8\xff"),
        ".gif": header.startswith((b"GIF87a", b"GIF89a")),
        ".webp": header.startswith(b"RIFF") and header[8:12] == b"WEBP",
    }.get(suffix, False)
    if not valid:
        raise ReportFailure("Markdown 图片格式或文件签名无效")


class ReportRuntime:
    def __init__(self, workspace: str | Path, state_root: str | Path | None = None):
        self.workspace = Path(workspace).resolve()
        self.state_root = Path(state_root or "/tmp/workspace-report").resolve()
        self.state_root.mkdir(parents=True, exist_ok=True)
        now = time.time()
        for path in self.state_root.iterdir():
            if path.is_dir() and now - path.stat().st_mtime > FAILURE_TTL_SECONDS:
                shutil.rmtree(path, ignore_errors=True)

    def _load(self, job_id: str) -> dict[str, Any]:
        state_path = self._job_path(job_id) / "state.json"
        return json.loads(state_path.read_text(encoding="utf-8"))

    def _job_path(self, job_id: str) -> Path:
        try:
            value = uuid.UUID(job_id)
        except (TypeError, ValueError) as error:
            raise ReportFailure("job_id 无效") from error
        job = self.state_root / str(value)
        if not (job / "state.json").is_file():
            raise ReportFailure("分析任务不存在")
        return job

    @contextmanager
    def _locked_state(self, job_id: str):
        job = self._job_path(job_id)
        with (job / "state.lock").open("a+b") as lock:
            fcntl.flock(lock, fcntl.LOCK_EX)
            try:
                yield self._load(job_id)
            finally:
                fcntl.flock(lock, fcntl.LOCK_UN)

    def _save(self, state: dict[str, Any]) -> None:
        job = self.state_root / state["jobId"]
        job.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(prefix="state-", suffix=".tmp", dir=job)
        temporary = Path(temporary_name)
        try:
            with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
                json.dump(state, stream, ensure_ascii=False)
            temporary.replace(job / "state.json")
        finally:
            temporary.unlink(missing_ok=True)

    def _matplotlib_config_dir(self, job_id: str) -> Path:
        try:
            config = self._job_path(job_id) / "matplotlib"
            if config.is_symlink() or (config.exists() and not config.is_dir()):
                raise ReportFailure("绘图字体配置目录无效")
            config.mkdir(mode=0o700, exist_ok=True)
            path = config / "matplotlibrc"
            if path.is_symlink() or (path.exists() and not path.is_file()):
                raise ReportFailure("绘图字体配置文件无效")
            if not path.is_file() or path.read_text(encoding="utf-8") != MATPLOTLIBRC:
                path.write_text(MATPLOTLIBRC, encoding="utf-8")
            return config
        except ReportFailure:
            raise
        except (OSError, UnicodeError) as error:
            raise ReportFailure("绘图字体配置文件无效") from error

    def _validate_datasets(self, state: dict[str, Any]) -> None:
        hashes = state.get("hashes")
        if not isinstance(hashes, dict):
            raise ReportFailure("分析任务缺少源文件校验信息")
        for value in state["paths"]:
            relative = _relative_path(value)
            path = self.workspace.joinpath(*relative.parts)
            _reject_symlinks(self.workspace, path)
            if not path.is_file() or _sha256(path) != hashes.get(value):
                raise ReportFailure("源文件发生变化，分析任务已失效")

    def _artifact(self, path: Path) -> dict[str, Any]:
        _reject_symlinks(self.workspace, path)
        if not path.is_file():
            raise ReportFailure("报表产物不存在或不是普通文件")
        return {
            "path": str(path.relative_to(self.workspace)),
            "size": path.stat().st_size,
            "sha256": _sha256(path),
        }

    @staticmethod
    def _job_status_value(state: dict[str, Any]) -> str:
        validation = state.get("validation")
        if isinstance(validation, dict):
            return "validated" if validation.get("ok") is True else "validation_failed"
        if isinstance(state.get("render"), dict):
            return "rendered"
        if state.get("successfulRoundCount", 0) > 0:
            return "analyzed"
        if state.get("roundCount", 0) > 0:
            return "analysis_failed"
        return "prepared"

    def prepare_dataset(self, paths: list[str]) -> dict[str, Any]:
        if not isinstance(paths, list) or not 1 <= len(paths) <= MAX_DATASET_PATHS:
            raise ReportFailure("paths 数量必须在 1 至 20 个之间")
        if len(set(paths)) != len(paths):
            raise ReportFailure("paths 不能重复")
        normalized = []
        hashes = {}
        for value in paths:
            relative = _relative_path(value)
            path = self.workspace.joinpath(*relative.parts)
            _reject_symlinks(self.workspace, path)
            if not path.is_file():
                raise ReportFailure("分析输入不存在或不是普通文件")
            normalized_path = str(relative)
            normalized.append(normalized_path)
            hashes[normalized_path] = _sha256(path)
        state = {
            "jobId": str(uuid.uuid4()),
            "paths": normalized,
            "hashes": hashes,
            "roundCount": 0,
            "successfulRoundCount": 0,
        }
        self._save(state)
        return {"status": "prepared", **state}

    @staticmethod
    def analysis_capabilities() -> dict[str, Any]:
        from importlib.metadata import PackageNotFoundError, version

        packages = {
            "numpy": "numpy",
            "pandas": "pandas",
            "polars": "polars",
            "pyarrow": "pyarrow",
            "scipy": "scipy",
            "statsmodels": "statsmodels",
            "scikit-learn": "scikit-learn",
            "matplotlib": "matplotlib",
            "seaborn": "seaborn",
            "plotly": "plotly",
            "openpyxl": "openpyxl",
            "xlrd": "xlrd",
            "xlsxwriter": "xlsxwriter",
            "pypdf": "pypdf",
            "pdfplumber": "pdfplumber",
            "reportlab": "reportlab",
            "weasyprint": "weasyprint",
            "markitdown": "markitdown",
            "networkx": "networkx",
            "sympy": "sympy",
        }
        available = {}
        for name, distribution in packages.items():
            try:
                available[name] = version(distribution)
            except PackageNotFoundError:
                continue
        commands = [
            name
            for name in (
                "python",
                "pandoc",
                "libreoffice",
                "pdfinfo",
                "pdftoppm",
                "pdftotext",
                "qpdf",
            )
            if shutil.which(name)
        ]
        return {
            "python": sys.version.split()[0],
            "packages": available,
            "commands": commands,
            "sql": ["sqlite3"],
        }

    @staticmethod
    def _read_profile_frame(path: Path):
        import pandas as pd

        suffix = path.suffix.lower()
        if suffix not in PROFILE_SUFFIXES:
            raise ReportFailure("确定性剖析仅支持 CSV、TSV、Excel、JSON、JSONL 和 Parquet")
        sampled = False
        if suffix in {".csv", ".tsv"}:
            chunks = pd.read_csv(
                path,
                sep="\t" if suffix == ".tsv" else ",",
                chunksize=MAX_PROFILE_ROWS,
                low_memory=False,
            )
            try:
                frame = next(chunks)
            except StopIteration:
                frame = pd.DataFrame()
            row_count = len(frame)
            for chunk in chunks:
                row_count += len(chunk)
                sampled = True
        elif suffix == ".jsonl":
            chunks = pd.read_json(path, lines=True, chunksize=MAX_PROFILE_ROWS)
            try:
                frame = next(chunks)
            except StopIteration:
                frame = pd.DataFrame()
            row_count = len(frame)
            for chunk in chunks:
                row_count += len(chunk)
                sampled = True
        elif suffix == ".parquet":
            import pyarrow.parquet as parquet

            source = parquet.ParquetFile(path)
            row_count = source.metadata.num_rows
            batch = next(source.iter_batches(batch_size=MAX_PROFILE_ROWS), None)
            frame = (
                batch.to_pandas()
                if batch is not None
                else pd.DataFrame(columns=source.schema_arrow.names)
            )
            sampled = row_count > len(frame)
        elif suffix in {".xls", ".xlsx"}:
            frame = pd.read_excel(path, nrows=MAX_PROFILE_ROWS)
            if suffix == ".xlsx":
                from openpyxl import load_workbook

                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    row_count = max(0, int(workbook.active.max_row or 0) - 1)
                finally:
                    workbook.close()
            else:
                import xlrd

                workbook = xlrd.open_workbook(path, on_demand=True)
                try:
                    row_count = max(0, int(workbook.sheet_by_index(0).nrows) - 1)
                finally:
                    workbook.release_resources()
            row_count = max(row_count, len(frame))
            sampled = row_count > len(frame)
        else:
            if path.stat().st_size > MAX_PROFILE_JSON_BYTES:
                raise ReportFailure("普通 JSON 文件超过 25 MiB，请改用 JSONL、Parquet 或分析命令")
            frame = pd.read_json(path)
            row_count = len(frame)
            if row_count > MAX_PROFILE_ROWS:
                frame = frame.head(MAX_PROFILE_ROWS)
                sampled = True
        return frame, int(row_count), sampled

    @staticmethod
    def _profile_column(series: Any, name: Any, *, sampled: bool) -> dict[str, Any]:
        import pandas as pd

        non_null = series.dropna()
        result: dict[str, Any] = {
            "name": str(name)[:200],
            "dtype": str(series.dtype),
            "nullCount": int(series.isna().sum()),
            "nullRatio": round(float(series.isna().mean()), 6) if len(series) else 0.0,
            "uniqueCount": int(non_null.nunique(dropna=True)),
            "statisticsScope": "sample" if sampled else "full",
        }
        if pd.api.types.is_numeric_dtype(series.dtype) and not pd.api.types.is_bool_dtype(
            series.dtype
        ):
            numeric = pd.to_numeric(non_null, errors="coerce").dropna()
            if len(numeric):
                quantiles = numeric.quantile([0.25, 0.5, 0.75])
                result["numeric"] = {
                    "min": _json_scalar(numeric.min()),
                    "max": _json_scalar(numeric.max()),
                    "mean": _json_scalar(round(float(numeric.mean()), 6)),
                    "median": _json_scalar(quantiles.loc[0.5]),
                    "p25": _json_scalar(quantiles.loc[0.25]),
                    "p75": _json_scalar(quantiles.loc[0.75]),
                }
        elif pd.api.types.is_datetime64_any_dtype(series.dtype):
            if len(non_null):
                result["datetime"] = {
                    "min": _json_scalar(non_null.min()),
                    "max": _json_scalar(non_null.max()),
                }
        else:
            counts = non_null.astype(str).value_counts(dropna=True).head(MAX_PROFILE_TOP_VALUES)
            result["topValues"] = [
                {"value": str(value)[:200], "count": int(count)} for value, count in counts.items()
            ]
        return result

    def profile_dataset(self, job_id: str) -> dict[str, Any]:
        with self._locked_state(job_id) as state:
            self._validate_datasets(state)
            column_limit = max(
                1,
                min(MAX_PROFILE_COLUMNS, MAX_PROFILE_TOTAL_COLUMNS // len(state["paths"])),
            )
            datasets = []
            for value in state["paths"]:
                relative = _relative_path(value)
                path = self.workspace.joinpath(*relative.parts)
                try:
                    frame, row_count, sampled = self._read_profile_frame(path)
                except ReportFailure:
                    raise
                except Exception as error:
                    raise ReportFailure(
                        f"无法剖析数据集 {relative}，请检查文件格式和沙箱分析依赖"
                    ) from error
                all_columns = list(frame.columns)
                selected_columns = all_columns[:column_limit]
                warnings = []
                if len(all_columns) > len(selected_columns):
                    warnings.append(f"列数超过返回边界，仅剖析前 {len(selected_columns)} 列")
                if sampled:
                    warnings.append(f"统计基于前 {len(frame)} 行确定性样本，rowCount 为完整行数")
                datasets.append(
                    {
                        "path": str(relative),
                        "format": path.suffix.lower().lstrip("."),
                        "size": path.stat().st_size,
                        "sha256": state["hashes"][value],
                        "rowCount": row_count,
                        "sampleRowCount": len(frame),
                        "columnCount": len(all_columns),
                        "sampled": sampled,
                        "columns": [
                            self._profile_column(frame[name], name, sampled=sampled)
                            for name in selected_columns
                        ],
                        "warnings": warnings,
                    }
                )
            result = {
                "status": "profiled",
                "jobId": state["jobId"],
                "datasetCount": len(datasets),
                "datasets": datasets,
            }
            size_warning = "剖析结果超过返回边界，已减少末尾列"
            while (
                len(json.dumps(result, ensure_ascii=False).encode("utf-8")) + 1 > MAX_RESULT_BYTES
            ):
                candidates = [dataset for dataset in datasets if dataset["columns"]]
                if not candidates:
                    raise ReportFailure("数据集元信息超过返回边界，请缩短文件路径后重试")
                dataset = max(candidates, key=lambda item: len(item["columns"]))
                dataset["columns"].pop()
                if size_warning not in dataset["warnings"]:
                    dataset["warnings"].append(size_warning)
            return result

    def job_status(self, job_id: str) -> dict[str, Any]:
        with self._locked_state(job_id) as state:
            self._validate_datasets(state)
            sources = [
                self._artifact(self.workspace.joinpath(*_relative_path(value).parts))
                for value in state["paths"]
            ]
            result: dict[str, Any] = {
                "jobId": state["jobId"],
                "status": self._job_status_value(state),
                "roundCount": state["roundCount"],
                "successfulRoundCount": state["successfulRoundCount"],
                "sources": sources,
            }
            render = state.get("render")
            if isinstance(render, dict):

                def current_artifact(recorded: dict[str, Any]) -> dict[str, Any]:
                    current = self._artifact(
                        self.workspace.joinpath(*_relative_path(recorded["path"]).parts)
                    )
                    current["changed"] = current["sha256"] != recorded["sha256"]
                    return current

                markdown_artifact = current_artifact(render["markdown"])
                pdf_artifact = current_artifact(render["pdf"])
                image_artifacts = [current_artifact(item) for item in render.get("images", [])]
                artifacts: dict[str, Any] = {
                    "markdown": markdown_artifact,
                    "pdf": pdf_artifact,
                    "images": image_artifacts,
                }
                result["artifacts"] = artifacts
                if (
                    markdown_artifact["changed"]
                    or pdf_artifact["changed"]
                    or any(item["changed"] for item in image_artifacts)
                ):
                    result["status"] = "artifact_changed"
            if isinstance(state.get("validation"), dict):
                result["validation"] = state["validation"]
            return result

    def analyze_dataset(
        self,
        job_id: str,
        command: str,
        cwd: str | None = None,
        timeout: int = 30,
    ) -> dict[str, Any]:
        if not isinstance(command, str) or not command.strip():
            raise ReportFailure("分析命令不能为空")
        if len(command.encode("utf-8")) > MAX_COMMAND_BYTES:
            raise ReportFailure("分析命令超过 32 KiB")
        if (
            isinstance(timeout, bool)
            or not isinstance(timeout, int)
            or not 1 <= timeout <= MAX_ANALYSIS_TIMEOUT
        ):
            raise ReportFailure("分析超时必须在 1 至 60 秒之间")
        working_directory = self.workspace
        if cwd is not None:
            relative = _relative_path(cwd)
            working_directory = self.workspace.joinpath(*relative.parts)
            _reject_symlinks(self.workspace, working_directory)
            if not working_directory.is_dir():
                raise ReportFailure("分析工作目录不存在")
        with self._locked_state(job_id) as state:
            self._validate_datasets(state)
            output_bytes, exit_code, truncated = _run_supervised_command(
                command,
                working_directory,
                timeout,
                self._matplotlib_config_dir(job_id),
            )
            self._validate_datasets(state)
            output = output_bytes.decode("utf-8", errors="replace")
            succeeded = exit_code == 0 and bool(output.strip())
            if exit_code == 0 and not succeeded:
                output = "分析命令退出码为 0，但没有返回分析结果；请输出本轮结果后重试。"
            state["roundCount"] += 1
            if succeeded:
                state["successfulRoundCount"] += 1
            self._save(state)
            return {
                "ok": succeeded,
                "jobId": state["jobId"],
                "status": "analyzed" if succeeded else "analysis_failed",
                "roundCount": state["roundCount"],
                "successfulRoundCount": state["successfulRoundCount"],
                "paths": state["paths"],
                "exitCode": exit_code,
                "output": output,
                "truncated": truncated,
            }

    def _images(self, markdown_path: Path, tokens: list[Any]) -> set[Path]:
        images: list[Path] = []
        for token in tokens:
            for child in token.children or []:
                if child.type != "image":
                    continue
                source = child.attrGet("src") or ""
                parsed = urlsplit(source)
                if parsed.scheme or parsed.netloc or parsed.query or parsed.fragment:
                    raise ReportFailure("Markdown 图片只能引用工作区内的相对路径")
                decoded = unquote(parsed.path)
                if not decoded or "\\" in decoded:
                    raise ReportFailure("Markdown 图片路径无效")
                relative = PurePosixPath(decoded)
                if relative.is_absolute() or ".." in relative.parts:
                    raise ReportFailure("Markdown 图片只能引用工作区内的相对路径")
                image = markdown_path.parent.joinpath(*relative.parts)
                _reject_symlinks(self.workspace, image)
                try:
                    image.relative_to(self.workspace)
                except ValueError as error:
                    raise ReportFailure("Markdown 图片路径越界") from error
                if image.suffix.lower() not in IMAGE_SUFFIXES or not image.is_file():
                    raise ReportFailure("Markdown 图片不存在或格式不受支持")
                if image.stat().st_size > MAX_IMAGE_BYTES:
                    raise ReportFailure("单张 Markdown 图片超过 10 MiB")
                _check_image_signature(image)
                images.append(image.resolve())
        unique = set(images)
        if len(images) > MAX_IMAGE_COUNT:
            raise ReportFailure("Markdown 图片数量超过 50 张")
        if sum(path.stat().st_size for path in unique) > MAX_TOTAL_IMAGE_BYTES:
            raise ReportFailure("Markdown 图片合计超过 50 MiB")
        return unique

    def render_markdown(self, job_id: str, markdown_path: str, output_path: str) -> dict[str, Any]:
        try:
            import pypdf
            from markdown_it import MarkdownIt
            from weasyprint import HTML, URLFetcher
        except ImportError as error:
            raise ReportFailure("PDF 运行时依赖不可用") from error

        with self._locked_state(job_id) as state:
            self._validate_datasets(state)
            if state["successfulRoundCount"] < 1:
                raise ReportFailure("至少成功完成一轮分析后才能渲染 PDF")
            source = _input_path(self.workspace, markdown_path, ".md")
            if source.stat().st_size > MAX_MARKDOWN_BYTES:
                raise ReportFailure("Markdown 文件超过 1 MiB")
            try:
                markdown = source.read_text(encoding="utf-8")
            except UnicodeDecodeError as error:
                raise ReportFailure("Markdown 文件必须使用 UTF-8 编码") from error

            parser = MarkdownIt("commonmark", {"html": False}).enable("table")
            tokens = parser.parse(markdown)
            allowed_images = self._images(source, tokens)
            source_artifact = self._artifact(source)
            image_artifacts = [self._artifact(path) for path in sorted(allowed_images)]
            body = parser.renderer.render(tokens, parser.options, {})
            output = _output_path(self.workspace, output_path)
            file_fetcher = URLFetcher(allowed_protocols={"file"}, fail_on_errors=True)

            def fetch_resource(url: str) -> dict[str, Any]:
                parsed = urlsplit(url)
                if parsed.scheme != "file" or parsed.netloc not in ("", "localhost"):
                    raise ReportFailure("PDF 渲染禁止访问外部资源")
                path = Path(unquote(parsed.path)).resolve()
                if path not in allowed_images:
                    raise ReportFailure("PDF 渲染引用了未校验资源")
                return file_fetcher(url)

            document = (
                "<meta charset='utf-8'>"
                "<style>"
                "@page{size:A4;margin:18mm}"
                "body{font-family:'Noto Sans CJK SC','Noto Sans CJK JP',sans-serif;"
                "font-size:10.5pt;line-height:1.65;color:#202124}"
                "h1{font-size:24pt}h2{font-size:17pt}h3{font-size:13pt}"
                "h1,h2,h3{page-break-after:avoid}"
                "table{width:100%;border-collapse:collapse;margin:10px 0}"
                "th,td{border:1px solid #c7c9cc;padding:5px 7px;text-align:left}"
                "th{background:#f1f3f4}"
                "img{display:block;max-width:100%;height:auto;margin:12px auto}"
                "pre,code{white-space:pre-wrap;overflow-wrap:anywhere}"
                "blockquote{border-left:3px solid #9aa0a6;margin-left:0;padding-left:12px}"
                "</style>"
                f"{body}"
            )
            descriptor, temporary_name = tempfile.mkstemp(
                prefix=f".{output.name}-", suffix=".tmp.pdf", dir=output.parent
            )
            os.close(descriptor)
            temporary = Path(temporary_name)
            try:
                HTML(
                    string=document,
                    base_url=str(source.parent),
                    url_fetcher=fetch_resource,
                ).write_pdf(str(temporary))
                reader = pypdf.PdfReader(str(temporary))
                if not reader.pages:
                    raise ReportFailure("PDF 校验失败")
                page_count = len(reader.pages)
                try:
                    os.link(temporary, output)
                except FileExistsError as error:
                    raise ReportFailure("PDF 输出文件已经存在") from error
            finally:
                temporary.unlink(missing_ok=True)
            if self._artifact(source)["sha256"] != source_artifact["sha256"] or any(
                self._artifact(path)["sha256"] != artifact["sha256"]
                for path, artifact in zip(sorted(allowed_images), image_artifacts, strict=True)
            ):
                output.unlink(missing_ok=True)
                raise ReportFailure("Markdown 或图片在渲染期间发生变化，请重新生成报表")
            render = {
                "markdown": source_artifact,
                "pdf": self._artifact(output),
                "images": image_artifacts,
                "pageCount": page_count,
                "imageCount": len(allowed_images),
            }
            state["render"] = render
            state.pop("validation", None)
            self._save(state)
            return {
                "status": "rendered",
                "jobId": state["jobId"],
                "roundCount": state["roundCount"],
                "successfulRoundCount": state["successfulRoundCount"],
                "markdownPath": str(source.relative_to(self.workspace)),
                "pdfPath": str(output.relative_to(self.workspace)),
                "pageCount": page_count,
                "imageCount": len(allowed_images),
                "size": output.stat().st_size,
            }

    def validate_pdf(self, job_id: str, pdf_path: str) -> dict[str, Any]:
        try:
            import pypdf
            from PIL import Image
        except ImportError as error:
            raise ReportFailure("PDF 视觉验收依赖不可用") from error
        if not shutil.which("pdftoppm"):
            raise ReportFailure("PDF 视觉验收命令不可用")

        with self._locked_state(job_id) as state:
            self._validate_datasets(state)
            render = state.get("render")
            if not isinstance(render, dict) or render.get("pdf", {}).get("path") != pdf_path:
                raise ReportFailure("PDF 未登记为当前分析任务的渲染产物")
            supporting_artifacts = [render["markdown"], *render.get("images", [])]
            for artifact in supporting_artifacts:
                supporting = self.workspace.joinpath(*_relative_path(artifact["path"]).parts)
                if self._artifact(supporting)["sha256"] != artifact["sha256"]:
                    raise ReportFailure("Markdown 或图片产物发生变化，请重新渲染后验收")
            relative = _relative_path(pdf_path, ".pdf")
            path = self.workspace.joinpath(*relative.parts)
            current = self._artifact(path)
            if current["sha256"] != render["pdf"]["sha256"]:
                raise ReportFailure("PDF 产物发生变化，请重新渲染后验收")
            pages = []
            blank_pages = []
            rendered_image_count = 0
            with tempfile.TemporaryDirectory(
                prefix="pdf-visual-", dir=self._job_path(job_id)
            ) as temp:
                prefix = Path(temp) / "page"
                try:
                    process = subprocess.run(
                        [
                            "pdftoppm",
                            "-gray",
                            "-r",
                            "72",
                            "-png",
                            str(path),
                            str(prefix),
                        ],
                        stdout=subprocess.DEVNULL,
                        stderr=subprocess.DEVNULL,
                        timeout=60,
                        check=False,
                    )
                    reader = pypdf.PdfReader(str(path))
                except (OSError, subprocess.TimeoutExpired, pypdf.errors.PdfReadError) as error:
                    raise ReportFailure("PDF 视觉验收无法打开产物") from error
                rendered_pages = sorted(
                    Path(temp).glob("page-*.png"),
                    key=lambda item: int(item.stem.rsplit("-", 1)[-1]),
                )
                if process.returncode != 0 or len(rendered_pages) != len(reader.pages):
                    raise ReportFailure("PDF 视觉验收栅格化失败")
                for index, (page, rendered_page) in enumerate(
                    zip(reader.pages, rendered_pages, strict=True), start=1
                ):
                    with Image.open(rendered_page) as image:
                        grayscale = image.convert("L")
                        samples = grayscale.tobytes()
                        width, height = grayscale.size
                    non_white = sum(value < 250 for value in samples)
                    ratio = round(non_white / len(samples), 6) if samples else 0.0
                    text_char_count = len("".join((page.extract_text() or "").split()))
                    image_count = len(page.images)
                    rendered_image_count += image_count
                    blank = ratio < 0.0005 and text_char_count == 0 and image_count == 0
                    if blank:
                        blank_pages.append(index)
                    pages.append(
                        {
                            "page": index,
                            "width": width,
                            "height": height,
                            "nonWhiteRatio": ratio,
                            "textCharCount": text_char_count,
                            "imageCount": image_count,
                            "blank": blank,
                        }
                    )
            markdown_image_count = int(render.get("imageCount") or 0)
            missing_images = max(0, markdown_image_count - rendered_image_count)
            ok = bool(pages) and not blank_pages and missing_images == 0
            validation = {
                "ok": ok,
                "status": "validated" if ok else "validation_failed",
                "pdfPath": pdf_path,
                "pdfSha256": current["sha256"],
                "pageCount": len(pages),
                "markdownImageCount": markdown_image_count,
                "renderedImageCount": rendered_image_count,
                "missingImageCount": missing_images,
                "blankPages": blank_pages,
                "pages": pages,
            }
            state["validation"] = validation
            self._save(state)
            return validation


def main(arguments: list[str] | None = None) -> int:
    values = arguments if arguments is not None else sys.argv[1:]
    try:
        if len(values) != 2:
            raise ReportFailure("报表渲染参数无效")
        action, payload_text = values
        payload = json.loads(payload_text)
        if action == "_execute":
            result = _execute_command(
                payload["command"],
                payload["timeout"],
                payload.get("matplotlib_config_dir"),
            )
        else:
            runtime = ReportRuntime(Path.cwd())
            if action == "capabilities":
                result = runtime.analysis_capabilities()
            elif action == "prepare":
                result = runtime.prepare_dataset(payload["paths"])
            elif action == "profile":
                result = runtime.profile_dataset(payload["job_id"])
            elif action == "status":
                result = runtime.job_status(payload["job_id"])
            elif action == "analyze":
                result = runtime.analyze_dataset(
                    payload["job_id"],
                    payload["command"],
                    payload.get("cwd"),
                    payload.get("timeout", 30),
                )
            elif action == "render_markdown":
                result = runtime.render_markdown(
                    payload["job_id"], payload["markdown_path"], payload["output_path"]
                )
            elif action == "validate_pdf":
                result = runtime.validate_pdf(payload["job_id"], payload["pdf_path"])
            else:
                raise ReportFailure("未知报表操作")
        encoded = json.dumps(result, ensure_ascii=False)
        if len(encoded.encode()) + 1 > MAX_RESULT_BYTES:
            raise ReportFailure("报表结果超过返回边界")
        print(encoded)
        return 0
    except (KeyError, TypeError, json.JSONDecodeError, ReportFailure) as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False))
        return 1
    except Exception:
        print(json.dumps({"error": "报表运行时执行失败"}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
