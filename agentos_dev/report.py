import hashlib
import json
import tempfile
import uuid
from contextlib import contextmanager
from io import BytesIO
from json import JSONDecodeError, JSONDecoder
from pathlib import Path, PurePosixPath
from typing import Any
from zipfile import BadZipFile, ZipFile

import pandas as pd
from agno.run import RunContext
from agno.tools.function import Function
from agno.tools.pandas import PandasTools
from openpyxl import load_workbook

from .workspace import WorkspaceError, WorkspaceService

MAX_DATASET_ROWS = 100_000
MAX_DATASET_COLUMNS = 100
MAX_MANIFEST_COLUMNS = 30
MAX_EXPANDED_BYTES = 128 * 1024 * 1024
MAX_XLSX_MEMBERS = 1000
MAX_RESULT_BYTES = 32 * 1024
MAX_SAMPLE_ROWS = 20
MAX_SAMPLE_COLUMNS = 10
MAX_MANIFEST_PARTS = 16
MAX_MANIFEST_PART_BYTES = 8 * 1024 * 1024
MAX_MANIFEST_TOTAL_BYTES = 100 * 1024 * 1024
MAX_MANIFEST_BYTES = 256 * 1024
DATASET_MANIFEST_VERSION = "agui.report.dataset.v1"
REPORT_CONFIG_VERSION = "agui.report.config.v1"
MAX_CHART_POINTS = 1000
MAX_PIE_CATEGORIES = 20
PREFLIGHT_CHUNK_ROWS = 1000
OBJECT_CELL_ESTIMATE_BYTES = 64
SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".json", ".jsonl"}
AGGREGATIONS = {"count", "sum", "avg", "min", "max"}
CHART_TYPES = {"bar", "line", "scatter", "pie", "histogram", "box"}
FINAL_REPORT_CHART_TYPES = {"bar", "line", "pie"}
FINAL_REPORT_SECTION_TYPES = {"overview", "notes", "chart", "analysis"}
NUMERIC_MANIFEST_TYPES = {"float", "integer", "monetary"}


def _thread(run_context: RunContext | None) -> str:
    if not run_context or not run_context.session_id:
        raise WorkspaceError("报表工具需要绑定对话的运行上下文")
    return run_context.session_id


def _column_names(frame: pd.DataFrame) -> list[str]:
    columns = [str(value) for value in frame.columns]
    if len(set(columns)) != len(columns):
        raise WorkspaceError("数据集字段名称转为文本后存在重复")
    return columns


def _ensure_columns(frame: pd.DataFrame) -> pd.DataFrame:
    columns = _column_names(frame)
    frame.columns = columns
    _check_shape(len(frame.index), len(frame.columns))
    expanded = int(frame.memory_usage(index=True, deep=True).sum())
    if expanded > MAX_EXPANDED_BYTES:
        raise WorkspaceError("数据集展开内存超过 128 MiB")
    return frame


def _check_shape(row_count: int, column_count: int) -> None:
    if row_count > MAX_DATASET_ROWS:
        raise WorkspaceError(f"数据集超过 {MAX_DATASET_ROWS} 行")
    if column_count > MAX_DATASET_COLUMNS:
        raise WorkspaceError(f"数据集超过 {MAX_DATASET_COLUMNS} 列")


def _storage_kind(series: pd.Series) -> str:
    if pd.api.types.is_bool_dtype(series.dtype):
        return "bool"
    if pd.api.types.is_numeric_dtype(series.dtype):
        return "numeric"
    if pd.api.types.is_datetime64_any_dtype(series.dtype):
        return "datetime"
    if pd.api.types.is_timedelta64_dtype(series.dtype):
        return "timedelta"
    return "object"


class _FramePreflight:
    def __init__(self):
        self.rows = 0
        self.column_bytes: dict[str, int] = {}
        self.column_kinds: dict[str, str] = {}

    def add(self, frame: pd.DataFrame) -> None:
        columns = _column_names(frame)
        frame.columns = columns
        row_count = len(frame.index)
        present = set(columns)
        previous_columns = set(self.column_bytes)

        for name in previous_columns - present:
            self.column_bytes[name] += row_count * 8
        for name in present - previous_columns:
            self.column_bytes[name] = self.rows * 8

        for name in columns:
            series = frame[name]
            kind = _storage_kind(series)
            previous_kind = self.column_kinds.get(name)
            size = int(series.memory_usage(index=False, deep=True))
            if previous_kind is None:
                self.column_kinds[name] = kind
            elif previous_kind == "object":
                size = max(size, row_count * OBJECT_CELL_ESTIMATE_BYTES)
            elif kind == "object" or kind != previous_kind:
                self.column_bytes[name] = max(
                    self.column_bytes[name],
                    self.rows * OBJECT_CELL_ESTIMATE_BYTES,
                )
                self.column_kinds[name] = "object"
            self.column_bytes[name] += size

        self.rows += row_count
        _check_shape(self.rows, len(self.column_bytes))
        expanded = sum(self.column_bytes.values()) + self.rows * 8
        if expanded > MAX_EXPANDED_BYTES:
            raise WorkspaceError("数据集展开内存超过 128 MiB")


def _skip_json_whitespace(value: str, index: int) -> int:
    while index < len(value) and value[index] in " \t\r\n":
        index += 1
    return index


def _preflight_json(content: bytes) -> None:
    try:
        value = content.decode("utf-8")
    except UnicodeDecodeError as error:
        raise WorkspaceError("JSON 数据集不是有效的 UTF-8 文件") from error
    index = _skip_json_whitespace(value, 0)
    if index >= len(value) or value[index] != "[":
        raise WorkspaceError("JSON 数据集仅支持顶层数组")
    index = _skip_json_whitespace(value, index + 1)
    if index < len(value) and value[index] == "]":
        return

    decoder = JSONDecoder()
    budget = _FramePreflight()
    batch: list[Any] = []
    try:
        while True:
            item, index = decoder.raw_decode(value, index)
            batch.append(item)
            if len(batch) >= PREFLIGHT_CHUNK_ROWS:
                budget.add(pd.DataFrame(batch))
                batch.clear()

            index = _skip_json_whitespace(value, index)
            if index >= len(value):
                raise JSONDecodeError("unterminated array", value, index)
            if value[index] == "]":
                index = _skip_json_whitespace(value, index + 1)
                if index != len(value):
                    raise JSONDecodeError("trailing data", value, index)
                if batch:
                    budget.add(pd.DataFrame(batch))
                return
            if value[index] != ",":
                raise JSONDecodeError("missing comma", value, index)
            index = _skip_json_whitespace(value, index + 1)
            if index >= len(value) or value[index] == "]":
                raise JSONDecodeError("missing array item", value, index)
    except JSONDecodeError as error:
        raise WorkspaceError("JSON 数据集格式无效") from error


def _preflight_delimited(path: str, suffix: str) -> None:
    budget = _FramePreflight()
    reader: Any
    if suffix == ".csv":
        reader = pd.read_csv(
            path,
            nrows=MAX_DATASET_ROWS + 1,
            chunksize=PREFLIGHT_CHUNK_ROWS,
        )
    else:
        reader = pd.read_json(
            path,
            lines=True,
            nrows=MAX_DATASET_ROWS + 1,
            chunksize=PREFLIGHT_CHUNK_ROWS,
        )
    for frame in reader:
        budget.add(frame)


def _preflight_jsonl_paths(paths: list[str]) -> None:
    budget = _FramePreflight()
    for path in paths:
        if not Path(path).stat().st_size:
            continue
        reader = pd.read_json(
            path,
            lines=True,
            nrows=MAX_DATASET_ROWS + 1,
            chunksize=PREFLIGHT_CHUNK_ROWS,
        )
        for frame in reader:
            budget.add(frame)


def _manifest_path_parts(path: str) -> tuple[str, ...] | None:
    parts = PurePosixPath(path).parts
    if (
        len(parts) == 4
        and parts[0] == "报表"
        and parts[1] == "原始数据"
        and parts[3] == "数据集.json"
    ):
        return parts
    return None


def _manifest_header(
    service: WorkspaceService,
    thread: str,
    path: str,
) -> tuple[dict[str, Any], bytes, str]:
    parts = _manifest_path_parts(path)
    if parts is None:
        raise WorkspaceError("最终报表只能使用中文原始数据目录中的数据集清单")
    dataset_id = parts[2]
    try:
        if str(uuid.UUID(dataset_id)) != dataset_id:
            raise ValueError
    except (ValueError, AttributeError):
        raise WorkspaceError("数据集清单目录不是有效的 UUID")
    content, _mime_type = service.file_bytes(thread, path)
    if len(content) > MAX_MANIFEST_BYTES:
        raise WorkspaceError("数据集清单超过 256 KiB")
    try:
        manifest = json.loads(content.decode("utf-8"))
    except (UnicodeDecodeError, JSONDecodeError) as error:
        raise WorkspaceError("数据集清单不是有效的 UTF-8 JSON") from error
    if not isinstance(manifest, dict) or manifest.get("version") != DATASET_MANIFEST_VERSION:
        raise WorkspaceError("数据集清单版本不受支持")
    if manifest.get("datasetId") != dataset_id:
        raise WorkspaceError("数据集清单与目录标识不一致")
    return manifest, content, dataset_id


def _load_dataset_manifest(
    service: WorkspaceService,
    thread: str,
    path: str,
    local_path: Path,
) -> tuple[dict[str, Any], str] | None:
    parts = _manifest_path_parts(path)
    if parts is None:
        return None
    manifest, _content, dataset_id = _manifest_header(service, thread, path)
    row_count = manifest.get("rowCount")
    total_size = manifest.get("totalSize")
    fields = manifest.get("fields")
    fragments = manifest.get("fragments")
    if (
        isinstance(row_count, bool)
        or not isinstance(row_count, int)
        or not 0 <= row_count <= MAX_DATASET_ROWS
    ):
        raise WorkspaceError("数据集清单行数无效")
    if (
        isinstance(total_size, bool)
        or not isinstance(total_size, int)
        or not 0 <= total_size <= MAX_MANIFEST_TOTAL_BYTES
    ):
        raise WorkspaceError("数据集清单总大小无效")
    if not isinstance(fields, list) or not 1 <= len(fields) <= MAX_MANIFEST_COLUMNS:
        raise WorkspaceError("数据集清单字段数必须在 1 至 30 之间")
    field_names = []
    for field in fields:
        name = field.get("name") if isinstance(field, dict) else None
        if not isinstance(name, str) or not name or len(name) > 128:
            raise WorkspaceError("数据集清单字段格式无效")
        field_names.append(name)
    if len(set(field_names)) != len(field_names):
        raise WorkspaceError("数据集清单字段不能重复")
    if not isinstance(fragments, list) or not 1 <= len(fragments) <= MAX_MANIFEST_PARTS:
        raise WorkspaceError("数据集清单分片数必须在 1 至 16 之间")

    actual_total = 0
    actual_rows = 0
    prefix = f"报表/原始数据/{dataset_id}/分片"
    with local_path.open("wb") as output:
        for index, fragment in enumerate(fragments, start=1):
            if not isinstance(fragment, dict):
                raise WorkspaceError("数据集清单分片格式无效")
            expected_path = f"{prefix}/数据-{index:04d}.jsonl"
            fragment_path = fragment.get("path")
            size = fragment.get("size")
            digest = fragment.get("sha256")
            if fragment_path != expected_path:
                raise WorkspaceError("数据集清单分片路径越界或顺序无效")
            if (
                isinstance(size, bool)
                or not isinstance(size, int)
                or not 0 <= size <= MAX_MANIFEST_PART_BYTES
            ):
                raise WorkspaceError("数据集清单分片大小无效")
            if not isinstance(digest, str) or len(digest) != 64:
                raise WorkspaceError("数据集清单分片摘要无效")
            fragment_content, _fragment_mime = service.file_bytes(thread, fragment_path)
            if len(fragment_content) != size:
                raise WorkspaceError("数据集分片大小与清单不一致")
            if hashlib.sha256(fragment_content).hexdigest() != digest:
                raise WorkspaceError("数据集分片 SHA-256 校验失败")
            for line in BytesIO(fragment_content):
                if not line.strip():
                    raise WorkspaceError("数据集分片包含空行")
                actual_rows += 1
            actual_total += len(fragment_content)
            output.write(fragment_content)
            del fragment_content
    if actual_total != total_size:
        raise WorkspaceError("数据集分片总大小与清单不一致")
    if actual_rows != row_count:
        raise WorkspaceError("数据集分片总行数与清单不一致")
    return manifest, str(local_path)


def _preflight_xlsx(path: str, sheet: str | int | None) -> None:
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise WorkspaceError(f"XLSX 压缩包成员超过 {MAX_XLSX_MEMBERS} 个")
            expanded = sum(item.file_size for item in members)
    except WorkspaceError:
        raise
    except BadZipFile as error:
        raise WorkspaceError("XLSX 数据集格式无效") from error
    if expanded > MAX_EXPANDED_BYTES:
        raise WorkspaceError(f"XLSX 展开内容超过 {MAX_EXPANDED_BYTES // (1024 * 1024)} MiB")

    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if isinstance(sheet, str):
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.worksheets[sheet or 0]
        _check_shape(max(0, worksheet.max_row - 1), worksheet.max_column)
        rows = iter(worksheet.iter_rows(values_only=True))
        next(rows, None)
        budget = _FramePreflight()
        batch = []
        for row in rows:
            batch.append(row)
            if len(batch) >= PREFLIGHT_CHUNK_ROWS:
                budget.add(pd.DataFrame(batch))
                batch.clear()
        if batch:
            budget.add(pd.DataFrame(batch))
    except WorkspaceError:
        raise
    except Exception as error:
        raise WorkspaceError("XLSX 数据集格式无效或工作表不存在") from error
    finally:
        if workbook is not None:
            workbook.close()


@contextmanager
def _load_frames(
    service: WorkspaceService,
    thread: str,
    paths: list[str],
    sheet: str | int | None = None,
):
    toolkit = PandasTools(
        enable_create_pandas_dataframe=True,
        enable_run_dataframe_operation=False,
    )
    with tempfile.TemporaryDirectory(prefix="agui-report-") as directory:
        frames: list[pd.DataFrame] = []
        try:
            for index, path in enumerate(paths):
                local_manifest_data = Path(directory) / f"dataset-{index}.jsonl"
                manifest_dataset = _load_dataset_manifest(
                    service, thread, path, local_manifest_data
                )
                if manifest_dataset is not None:
                    manifest, manifest_local_path = manifest_dataset
                    expected_columns = [item["name"] for item in manifest["fields"]]
                    _preflight_jsonl_paths([manifest_local_path])
                    if not Path(manifest_local_path).stat().st_size:
                        frame = pd.DataFrame(columns=expected_columns)
                    else:
                        name = f"dataset_{index}"
                        manifest_parameters = {
                            "path_or_buf": manifest_local_path,
                            "lines": True,
                            "nrows": MAX_DATASET_ROWS + 1,
                        }
                        toolkit.create_pandas_dataframe(name, "read_json", manifest_parameters)
                        frame = toolkit.dataframes.get(name)
                        if frame is None:
                            frame = pd.read_json(**manifest_parameters)
                            toolkit.dataframes[name] = frame
                    frame.columns = _column_names(frame)
                    if list(frame.columns) != expected_columns:
                        raise WorkspaceError("数据集分片字段与清单不一致")
                    if len(frame.index) != manifest["rowCount"]:
                        raise WorkspaceError("数据集加载行数与清单不一致")
                    frames.append(_ensure_columns(frame))
                    continue
                suffix = PurePosixPath(path).suffix.lower()
                if suffix not in SUPPORTED_SUFFIXES:
                    raise WorkspaceError("报表仅支持 CSV、XLSX、JSON 和 JSONL 文件")
                content, _mime_type = service.file_bytes(thread, path)
                local_path = PurePosixPath(directory) / (f"dataset-{index}{suffix}")
                with open(str(local_path), "wb") as output:
                    output.write(content)
                if suffix in {".csv", ".jsonl"}:
                    _preflight_delimited(str(local_path), suffix)
                elif suffix == ".json":
                    _preflight_json(content)
                elif suffix == ".xlsx":
                    _preflight_xlsx(str(local_path), sheet)
                function_name = {
                    ".csv": "read_csv",
                    ".xlsx": "read_excel",
                    ".json": "read_json",
                    ".jsonl": "read_json",
                }[suffix]
                parameters: dict[str, Any]
                if suffix == ".csv":
                    parameters = {
                        "filepath_or_buffer": str(local_path),
                        "nrows": MAX_DATASET_ROWS + 1,
                    }
                elif suffix == ".xlsx":
                    parameters = {
                        "io": str(local_path),
                        "nrows": MAX_DATASET_ROWS + 1,
                    }
                    if sheet is not None:
                        parameters["sheet_name"] = sheet
                else:
                    parameters = {"path_or_buf": str(local_path)}
                    if suffix == ".jsonl":
                        parameters["lines"] = True
                        parameters["nrows"] = MAX_DATASET_ROWS + 1
                name = f"dataset_{index}"
                toolkit.create_pandas_dataframe(name, function_name, parameters)
                frame = toolkit.dataframes.get(name)
                if frame is None:
                    # PandasTools intentionally does not retain empty frames.
                    loader = getattr(pd, function_name)
                    frame = loader(**parameters)
                    toolkit.dataframes[name] = frame
                frames.append(_ensure_columns(frame))
            yield frames
        finally:
            toolkit.dataframes.clear()
            frames.clear()


def _records(frame: pd.DataFrame) -> list[dict[str, Any]]:
    return json.loads(frame.to_json(orient="records", date_format="iso"))


def _bounded(payload: dict[str, Any]) -> dict[str, Any]:
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    if len(encoded) <= MAX_RESULT_BYTES:
        return payload
    result = dict(payload, truncated=True)
    for key in ("rows", "columns", "sample"):
        values = result.get(key)
        if not isinstance(values, list):
            continue
        values = list(values)
        result[key] = values
        while (
            values
            and len(
                json.dumps(
                    result,
                    ensure_ascii=False,
                    separators=(",", ":"),
                    default=str,
                ).encode("utf-8")
            )
            > MAX_RESULT_BYTES
        ):
            values.pop()
    if (
        len(
            json.dumps(
                result,
                ensure_ascii=False,
                separators=(",", ":"),
                default=str,
            ).encode("utf-8")
        )
        > MAX_RESULT_BYTES
    ):
        raise WorkspaceError("报表工具结果超过 32 KiB")
    return result


def _column(frame: pd.DataFrame, name: str) -> pd.Series:
    if name not in frame.columns:
        raise WorkspaceError(f"数据集不存在字段：{name}")
    return frame[name]


def _validate_dimensions(frame: pd.DataFrame, dimensions: list[str]) -> None:
    if not isinstance(dimensions, list) or len(dimensions) > 2:
        raise WorkspaceError("分组维度最多 2 个")
    if len(set(dimensions)) != len(dimensions):
        raise WorkspaceError("分组维度不能重复")
    for name in dimensions:
        _column(frame, name)


def _metric_spec(frame: pd.DataFrame, metrics: list[dict[str, str]]):
    if not isinstance(metrics, list) or not 1 <= len(metrics) <= 5:
        raise WorkspaceError("聚合指标必须为 1 至 5 个")
    result = []
    aliases = set()
    for metric in metrics:
        if not isinstance(metric, dict) or set(metric) != {"field", "aggregation"}:
            raise WorkspaceError("聚合指标格式无效")
        field = metric["field"]
        aggregation = metric["aggregation"]
        series = _column(frame, field)
        if aggregation not in AGGREGATIONS:
            raise WorkspaceError(f"不支持的聚合方式：{aggregation}")
        if aggregation in {"sum", "avg"} and not pd.api.types.is_numeric_dtype(series):
            raise WorkspaceError(f"字段 {field} 不是数值类型")
        alias = f"{field}:{aggregation}"
        if alias in aliases:
            raise WorkspaceError("聚合指标不能重复")
        aliases.add(alias)
        result.append((alias, field, "mean" if aggregation == "avg" else aggregation))
    return result


def _group_frame(
    frame: pd.DataFrame,
    dimensions: list[str],
    metrics: list[dict[str, str]],
) -> pd.DataFrame:
    _validate_dimensions(frame, dimensions)
    specs = _metric_spec(frame, metrics)
    named = {
        alias: pd.NamedAgg(column=field, aggfunc=aggregation) for alias, field, aggregation in specs
    }
    if dimensions:
        return frame.groupby(dimensions, dropna=False, observed=True).agg(**named).reset_index()
    values = {alias: getattr(frame[field], aggregation)() for alias, field, aggregation in specs}
    return pd.DataFrame([values])


def report_tools(service: WorkspaceService) -> list[Function]:
    def profile_dataset(
        path: str,
        sheet: str | int | None = None,
        run_context: RunContext | None = None,
    ):
        """概览当前对话工作区中的 CSV、XLSX、JSON 或 JSONL 数据集。"""
        with _load_frames(service, _thread(run_context), [path], sheet=sheet) as frames:
            frame = frames[0]
            columns = []
            for name in frame.columns:
                series = frame[name]
                item: dict[str, Any] = {
                    "name": name,
                    "dtype": str(series.dtype),
                    "nullCount": int(series.isna().sum()),
                    "uniqueCount": int(series.nunique(dropna=True)),
                }
                if pd.api.types.is_numeric_dtype(series):
                    item.update(
                        {
                            "min": None if series.dropna().empty else float(series.min()),
                            "max": None if series.dropna().empty else float(series.max()),
                            "mean": None if series.dropna().empty else float(series.mean()),
                        }
                    )
                columns.append(item)
            return _bounded(
                {
                    "path": path,
                    "rowCount": len(frame.index),
                    "columnCount": len(frame.columns),
                    "memoryBytes": int(frame.memory_usage(index=True, deep=True).sum()),
                    "columns": columns,
                }
            )

    def sample_dataset(
        path: str,
        columns: list[str] | None = None,
        limit: int = 10,
        run_context: RunContext | None = None,
    ):
        """返回当前对话数据集的有界原始样例，最多 20 行、10 列和 32 KiB。"""
        if (
            isinstance(limit, bool)
            or not isinstance(limit, int)
            or not 1 <= limit <= MAX_SAMPLE_ROWS
        ):
            raise WorkspaceError("样例行数必须在 1 至 20 之间")
        with _load_frames(service, _thread(run_context), [path]) as frames:
            frame = frames[0]
            selected = list(frame.columns[:MAX_SAMPLE_COLUMNS]) if columns is None else columns
            if (
                not isinstance(selected, list)
                or not 1 <= len(selected) <= MAX_SAMPLE_COLUMNS
                or not all(isinstance(name, str) for name in selected)
                or len(set(selected)) != len(selected)
            ):
                raise WorkspaceError("样例字段必须是 1 至 10 个不重复字段")
            for name in selected:
                _column(frame, name)
            return _bounded(
                {
                    "path": path,
                    "rowCount": min(limit, len(frame.index)),
                    "columnCount": len(selected),
                    "columns": selected,
                    "rows": _records(frame[selected].head(limit)),
                }
            )

    def group_dataset(
        path: str,
        dimensions: list[str],
        metrics: list[dict[str, str]],
        run_context: RunContext | None = None,
    ):
        """按最多两个维度和五个语义化指标聚合当前对话的数据集。"""
        with _load_frames(service, _thread(run_context), [path]) as frames:
            result = _group_frame(frames[0], dimensions, metrics)
            return _bounded(
                {
                    "path": path,
                    "dimensions": dimensions,
                    "metrics": metrics,
                    "rowCount": len(result.index),
                    "rows": _records(result),
                }
            )

    def pivot_dataset(
        path: str,
        rows: list[str],
        columns: list[str],
        value: str,
        aggregation: str,
        run_context: RunContext | None = None,
    ):
        """用受控聚合方式创建当前对话数据集的透视结果。"""
        with _load_frames(service, _thread(run_context), [path]) as frames:
            frame = frames[0]
            if (
                not isinstance(rows, list)
                or not isinstance(columns, list)
                or not rows
                or not columns
            ):
                raise WorkspaceError("透视表必须提供行维度和列维度")
            if len(rows) > 2 or len(columns) > 2:
                raise WorkspaceError("透视表行维度和列维度各最多 2 个")
            for name in rows + columns + [value]:
                _column(frame, name)
            if aggregation not in AGGREGATIONS:
                raise WorkspaceError("不支持的透视聚合方式")
            if aggregation in {"sum", "avg"} and not pd.api.types.is_numeric_dtype(frame[value]):
                raise WorkspaceError("透视值字段不是数值类型")
            result = pd.pivot_table(
                frame,
                index=rows,
                columns=columns,
                values=value,
                aggfunc="mean" if aggregation == "avg" else aggregation,
                observed=True,
            ).reset_index()
            result.columns = [
                " / ".join(str(part) for part in name if str(part))
                if isinstance(name, tuple)
                else str(name)
                for name in result.columns
            ]
            return _bounded(
                {
                    "path": path,
                    "rowCount": len(result.index),
                    "rows": _records(result),
                }
            )

    def concat_datasets(
        paths: list[str],
        source_labels: list[str],
        run_context: RunContext | None = None,
    ):
        """显式纵向合并结构完全一致的数据集，并标注各行来源。"""
        if not isinstance(paths, list) or not 2 <= len(paths) <= 5:
            raise WorkspaceError("纵向合并必须选择 2 至 5 个数据集")
        if not isinstance(source_labels, list) or len(source_labels) != len(paths):
            raise WorkspaceError("每个数据集都必须提供来源标签")
        thread = _thread(run_context)
        with _load_frames(service, thread, paths) as frames:
            expected = list(frames[0].columns)
            if "_source" in expected or any(
                list(frame.columns) != expected for frame in frames[1:]
            ):
                raise WorkspaceError("仅支持字段结构和顺序完全一致的数据集纵向合并")
            labeled = []
            for frame, label in zip(frames, source_labels):
                item = frame.copy()
                item.insert(0, "_source", str(label)[:120])
                labeled.append(item)
            result = _ensure_columns(pd.concat(labeled, ignore_index=True))
            path = f"报表/分析数据/{uuid.uuid4()}/合并数据.jsonl"
            content = result.to_json(orient="records", lines=True, force_ascii=False).encode(
                "utf-8"
            )
            service.upload(thread, path, content)
            return _bounded(
                {
                    "path": path,
                    "rowCount": len(result.index),
                    "columnCount": len(result.columns),
                    "sourceLabels": source_labels,
                }
            )

    def create_report_config(
        manifest_path: str,
        dataset_hash: str,
        title: str,
        dimensions: list[str],
        metrics: list[dict[str, str]],
        chart_type: str,
        chart_metric: str,
        chart_title: str | None = None,
        analysis_notes: list[str] | None = None,
        purpose: str | None = None,
        sections: list[str] | None = None,
        run_context: RunContext | None = None,
    ):
        """校验分析口径并在中文配置目录中创建最终报表配置。"""
        thread = _thread(run_context)
        manifest, manifest_content, _dataset_id = _manifest_header(service, thread, manifest_path)
        if (
            not isinstance(dataset_hash, str)
            or len(dataset_hash) != 64
            or hashlib.sha256(manifest_content).hexdigest() != dataset_hash
        ):
            raise WorkspaceError("数据集清单摘要不匹配")
        fields = manifest.get("fields")
        if not isinstance(fields, list) or not 1 <= len(fields) <= MAX_MANIFEST_COLUMNS:
            raise WorkspaceError("数据集清单字段格式无效")
        definitions = {
            item.get("name"): item
            for item in fields
            if isinstance(item, dict) and isinstance(item.get("name"), str)
        }
        if len(definitions) != len(fields):
            raise WorkspaceError("数据集清单字段格式无效")
        if not isinstance(title, str) or not title.strip() or len(title.strip()) > 160:
            raise WorkspaceError("报表标题必须是 1 至 160 个字符")
        if (
            not isinstance(dimensions, list)
            or not 1 <= len(dimensions) <= 2
            or len(set(dimensions)) != len(dimensions)
            or any(name not in definitions for name in dimensions)
        ):
            raise WorkspaceError("最终报表必须选择 1 至 2 个有效维度")
        if not isinstance(metrics, list) or not 1 <= len(metrics) <= 5:
            raise WorkspaceError("最终报表必须选择 1 至 5 个指标")
        aliases = set()
        normalized_metrics = []
        for metric in metrics:
            if not isinstance(metric, dict) or set(metric) - {"field", "aggregation", "label"}:
                raise WorkspaceError("最终报表指标格式无效")
            field = metric.get("field")
            aggregation = metric.get("aggregation")
            label = metric.get("label")
            if field not in definitions or aggregation not in AGGREGATIONS:
                raise WorkspaceError("最终报表指标字段或聚合方式无效")
            if (
                aggregation in {"sum", "avg"}
                and definitions[field].get("type") not in NUMERIC_MANIFEST_TYPES
            ):
                raise WorkspaceError(f"字段 {field} 不是可求和或平均的数值字段")
            if aggregation in {"min", "max"} and definitions[field].get(
                "type"
            ) not in NUMERIC_MANIFEST_TYPES | {"date", "datetime"}:
                raise WorkspaceError(f"字段 {field} 不支持最小值或最大值")
            alias = f"{field}:{aggregation}"
            if alias in aliases:
                raise WorkspaceError("最终报表指标不能重复")
            if label is not None and (
                not isinstance(label, str) or not label.strip() or len(label.strip()) > 80
            ):
                raise WorkspaceError("指标标签必须是 1 至 80 个字符")
            aliases.add(alias)
            normalized_metrics.append(
                {"field": field, "aggregation": aggregation, "label": label or alias}
            )
        chart_metric_definition = next(
            (
                metric
                for metric in normalized_metrics
                if f"{metric['field']}:{metric['aggregation']}" == chart_metric
            ),
            None,
        )
        if chart_type not in FINAL_REPORT_CHART_TYPES or chart_metric_definition is None:
            raise WorkspaceError("最终报表图表类型或指标无效")
        if (
            chart_metric_definition["aggregation"] != "count"
            and definitions[chart_metric_definition["field"]].get("type")
            not in NUMERIC_MANIFEST_TYPES
        ):
            raise WorkspaceError("最终报表图表指标必须是数值结果")
        if chart_title is not None and (
            not isinstance(chart_title, str) or len(chart_title.strip()) > 160
        ):
            raise WorkspaceError("图表标题不能超过 160 个字符")
        notes = analysis_notes or []
        if (
            not isinstance(notes, list)
            or len(notes) > 10
            or any(not isinstance(note, str) or len(note) > 500 for note in notes)
        ):
            raise WorkspaceError("分析说明最多 10 条且每条不超过 500 个字符")
        report_purpose = title.strip() if purpose is None else purpose
        if (
            not isinstance(report_purpose, str)
            or not report_purpose.strip()
            or len(report_purpose.strip()) > 500
        ):
            raise WorkspaceError("报表目的必须是 1 至 500 个字符")
        report_purpose = report_purpose.strip()
        selected_sections = (
            ["overview", "notes", "chart", "analysis"] if sections is None else sections
        )
        if (
            not isinstance(selected_sections, list)
            or not 1 <= len(selected_sections) <= len(FINAL_REPORT_SECTION_TYPES)
            or len(set(selected_sections)) != len(selected_sections)
            or any(value not in FINAL_REPORT_SECTION_TYPES for value in selected_sections)
        ):
            raise WorkspaceError("报表章节必须是有效且不重复的章节列表")

        report_id = str(uuid.uuid4())
        path = f"报表/配置/{report_id}/报表配置.json"
        config = {
            "version": REPORT_CONFIG_VERSION,
            "reportId": report_id,
            "manifestPath": manifest_path,
            "datasetHash": dataset_hash,
            "title": title.strip(),
            "analysis": {
                "dimensions": dimensions,
                "metrics": normalized_metrics,
                "notes": notes,
            },
            "chart": {
                "type": chart_type,
                "metric": chart_metric,
                "title": (chart_title or title).strip(),
            },
            "presentation": {
                "purpose": report_purpose,
                "sections": selected_sections,
            },
        }
        service.create_file(
            thread,
            path,
            json.dumps(config, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode(
                "utf-8"
            ),
        )
        return {"reportId": report_id, "configPath": path}

    def generate_chart(
        path: str,
        chart_type: str,
        x: str,
        y: str | None = None,
        group: str | None = None,
        aggregation: str | None = None,
        title: str | None = None,
        run_context: RunContext | None = None,
    ):
        """为当前对话数据集生成 PNG 和独立 HTML 图表产物。"""
        if chart_type not in CHART_TYPES:
            raise WorkspaceError("不支持的图表类型")
        thread = _thread(run_context)
        with _load_frames(service, thread, [path]) as frames:
            frame = frames[0]
            selected = [name for name in (x, y, group) if name]
            for name in selected:
                _column(frame, name)
            plot_frame = frame[selected].copy()
            if chart_type in {"bar", "line", "scatter"} and not y:
                raise WorkspaceError("柱状图、折线图和散点图必须提供 y 字段")
            if aggregation:
                if (
                    aggregation not in AGGREGATIONS
                    or not y
                    or chart_type in {"scatter", "histogram", "box"}
                ):
                    raise WorkspaceError("该图表不支持所选聚合方式")
                plot_frame = _group_frame(
                    plot_frame,
                    [x] + ([group] if group else []),
                    [{"field": y, "aggregation": aggregation}],
                ).rename(columns={f"{y}:{aggregation}": y})
            if chart_type in {"bar", "line"} and group and not aggregation:
                if bool(plot_frame.duplicated([x, group]).any()):
                    raise WorkspaceError("分组图存在重复坐标，请指定聚合方式")
            if chart_type == "pie":
                if y:
                    if not pd.api.types.is_numeric_dtype(plot_frame[y]):
                        raise WorkspaceError("饼图数值字段不是数值类型")
                    plot_frame = (
                        plot_frame.groupby(
                            x,
                            dropna=False,
                            observed=True,
                        )[y]
                        .sum()
                        .reset_index()
                    )
                else:
                    plot_frame = (
                        plot_frame.groupby(
                            x,
                            dropna=False,
                            observed=True,
                        )
                        .size()
                        .reset_index(name="_count")
                    )
                    y = "_count"
                if len(plot_frame.index) > MAX_PIE_CATEGORIES:
                    raise WorkspaceError("饼图分类超过 20 个")
            if len(plot_frame.index) > MAX_CHART_POINTS:
                raise WorkspaceError("图表绘制点超过 1000 个")

            import matplotlib

            matplotlib.use("Agg")
            import matplotlib.pyplot as plt
            import plotly.express as px

            if chart_type == "histogram":
                html_figure = px.histogram(plot_frame, x=y or x, title=title)
            elif chart_type == "box":
                html_figure = px.box(plot_frame, x=group, y=y or x, title=title)
            elif chart_type == "pie":
                html_figure = px.pie(plot_frame, names=x, values=y, title=title)
            else:
                html_factory = {
                    "bar": px.bar,
                    "line": px.line,
                    "scatter": px.scatter,
                }[chart_type]
                html_figure = html_factory(plot_frame, x=x, y=y, color=group, title=title)

            identifier = str(uuid.uuid4())
            png_path = f"报表/图表/{identifier}/图表.png"
            html_path = f"报表/图表/{identifier}/交互图表.html"
            figure, axis = plt.subplots(figsize=(9, 5.5))
            try:
                if chart_type in {"bar", "line"}:
                    if group:
                        table = plot_frame.pivot(index=x, columns=group, values=y)
                        table.plot(kind=chart_type, ax=axis)
                    else:
                        plot_frame.plot(x=x, y=y, kind=chart_type, ax=axis, legend=False)
                elif chart_type == "scatter":
                    if group:
                        for label, values in plot_frame.groupby(group, dropna=False):
                            axis.scatter(values[x], values[y], label=str(label))
                        axis.legend()
                    else:
                        axis.scatter(plot_frame[x], plot_frame[y])
                elif chart_type == "pie":
                    plot_frame.set_index(x)[y].plot(kind="pie", ax=axis, ylabel="")
                elif chart_type == "histogram":
                    plot_frame[y or x].plot(kind="hist", ax=axis)
                else:
                    plot_frame.boxplot(column=y or x, by=group, ax=axis)
                axis.set_title((title or "数据图表")[:160])
                figure.tight_layout()
                with tempfile.TemporaryDirectory(prefix="agui-chart-") as directory:
                    local_png = f"{directory}/chart.png"
                    figure.savefig(local_png, format="png", dpi=144)
                    with open(local_png, "rb") as source:
                        png = source.read()
            finally:
                plt.close(figure)
            html = html_figure.to_html(
                full_html=True,
                include_plotlyjs=True,
            ).encode("utf-8")
            service.upload(thread, png_path, png)
            try:
                service.upload(thread, html_path, html)
            except Exception:
                service.delete_file(thread, png_path)
                raise
            return _bounded(
                {
                    "sourcePath": path,
                    "chartType": chart_type,
                    "pointCount": len(plot_frame.index),
                    "pngPath": png_path,
                    "htmlPath": html_path,
                }
            )

    return [
        Function(name="pandas_profile_dataset", entrypoint=profile_dataset),
        Function(name="pandas_sample_dataset", entrypoint=sample_dataset),
        Function(name="pandas_group_dataset", entrypoint=group_dataset),
        Function(name="pandas_pivot_dataset", entrypoint=pivot_dataset),
        Function(name="pandas_concat_datasets", entrypoint=concat_datasets),
        Function(name="pandas_generate_chart", entrypoint=generate_chart),
        Function(name="pandas_create_report_config", entrypoint=create_report_config),
    ]
