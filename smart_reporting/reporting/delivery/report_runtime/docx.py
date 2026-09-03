"""Reporting DOCX 渲染、后处理和结构验收能力。"""

from __future__ import annotations

import html
import os
import shutil
import subprocess
import zipfile
from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path
from string import Formatter
from typing import Any
from xml.etree import ElementTree

from .markdown import _WORD_MARKERS, REPORT_VISUAL_THEME
from .pdf import MAX_PDF_PAGES, _formatted_page_text, _page_number_context
from .validation import MAX_DOCX_BYTES, ReportFailure

DOCX_RENDER_TIMEOUT_SECONDS = 540
DOCX_VALIDATION_TIMEOUT_SECONDS = 540
_DOCX_FORBIDDEN_PARTS = ("word/vbaProject.bin", "word/embeddings/", "word/activeX/")
_WORD_PAGE_FIELDS = {"page": "PAGE", "pages": "SECTIONPAGES"}


def _fit_image_dimensions(
    width: int, height: int, maximum_width: int, maximum_height: int
) -> tuple[int, int]:
    scale = min(1, maximum_width / width, maximum_height / height)
    return int(width * scale), int(height * scale)


def _render_docx(
    html_document: str,
    *,
    source_parent: Path,
    output: Path,
    context: dict[str, Any],
    layout: dict[str, str],
) -> dict[str, Any]:
    pandoc = shutil.which("pandoc")
    if pandoc is None:
        raise ReportFailure("Word 渲染命令 Pandoc 不可用")
    html_path = output.with_suffix(".html")
    html_path.write_text(html_document, encoding="utf-8")
    try:
        process = subprocess.run(
            [
                pandoc,
                "--from=html",
                "--to=docx",
                f"--resource-path={source_parent}",
                "--output",
                str(output),
                str(html_path),
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=DOCX_RENDER_TIMEOUT_SECONDS,
            check=False,
            cwd=source_parent,
        )
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ReportFailure("Word 渲染失败或超时") from error
    if process.returncode != 0 or not output.is_file():
        raise ReportFailure("Word 渲染失败")
    _postprocess_docx(output, context=context, layout=layout)
    return _validate_docx_structure(
        output,
        expected_sections=context["sections"],
        expected_headings=context["headingNumbers"],
        expected_image_count=None,
    )


def _postprocess_docx(path: Path, *, context: dict[str, Any], layout: dict[str, str]) -> None:
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.oxml import OxmlElement, parse_xml
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
    except ImportError as error:
        raise ReportFailure("Word 后处理依赖 python-docx 不可用") from error

    document = Document(str(path))
    theme_colors = {
        name: RGBColor.from_string(str(value).removeprefix("#"))
        for name, value in REPORT_VISUAL_THEME.items()
        if isinstance(value, str) and value.startswith("#")
    }
    markers: dict[str, Any] = {}
    for paragraph in document.paragraphs:
        for name, marker in _WORD_MARKERS.items():
            if paragraph.text.strip() == marker:
                if name in markers:
                    raise ReportFailure("Word 版式标记重复")
                markers[name] = paragraph
    if set(markers) != set(_WORD_MARKERS):
        raise ReportFailure("Word 版式标记缺失")

    def clear_paragraph(paragraph: Any) -> None:
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)

    def field_run(
        paragraph: Any,
        instruction: str,
        *,
        result: str = "1",
        close: bool = True,
    ) -> None:
        begin_run = OxmlElement("w:r")
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        begin.set(qn("w:dirty"), "true")
        begin_run.append(begin)
        instruction_run = OxmlElement("w:r")
        instruction_text = OxmlElement("w:instrText")
        instruction_text.set(qn("xml:space"), "preserve")
        instruction_text.text = f" {instruction} "
        instruction_run.append(instruction_text)
        separate_run = OxmlElement("w:r")
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        separate_run.append(separate)
        result_run = OxmlElement("w:r")
        result_text = OxmlElement("w:t")
        result_text.text = result
        result_run.append(result_text)
        for item in (begin_run, instruction_run, separate_run, result_run):
            paragraph._p.append(item)
        if close:
            end_run = OxmlElement("w:r")
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            end_run.append(end)
            paragraph._p.append(end_run)

    def field_end(paragraph: Any) -> None:
        end_run = OxmlElement("w:r")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        end_run.append(end)
        paragraph._p.append(end_run)

    base_sect_pr = deepcopy(document._element.body.sectPr)

    def section_properties(*, page_format: str | None) -> Any:
        value = deepcopy(base_sect_pr)
        for child in list(value):
            if child.tag in {
                qn("w:headerReference"),
                qn("w:footerReference"),
                qn("w:type"),
                qn("w:pgNumType"),
                qn("w:titlePg"),
            }:
                value.remove(child)
        section_type = OxmlElement("w:type")
        section_type.set(qn("w:val"), "nextPage")
        value.insert(0, section_type)
        if page_format is not None:
            page_number = OxmlElement("w:pgNumType")
            page_number.set(qn("w:start"), "1")
            page_number.set(qn("w:fmt"), page_format)
            value.append(page_number)
        return value

    for marker_name, page_format in (("cover_end", None), ("toc_end", "lowerRoman")):
        paragraph = markers[marker_name]
        clear_paragraph(paragraph)
        paragraph._p.get_or_add_pPr().append(section_properties(page_format=page_format))
    final_sect_pr = document._element.body.sectPr
    for child in list(final_sect_pr):
        if child.tag in {qn("w:headerReference"), qn("w:footerReference"), qn("w:pgNumType")}:
            final_sect_pr.remove(child)
    body_page_number = OxmlElement("w:pgNumType")
    body_page_number.set(qn("w:start"), "1")
    body_page_number.set(qn("w:fmt"), "decimal")
    final_sect_pr.append(body_page_number)

    sections = document.sections
    if len(sections) != 3:
        raise ReportFailure("Word 必须包含封面、目录和正文三个分节")
    for section in sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(22)
        section.left_margin = Mm(18)
        section.right_margin = Mm(18)
        section.header_distance = Mm(8)
        section.footer_distance = Mm(8)

    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = (
            document.styles[style_name]
            if style_name in document.styles
            else document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        )
        style.font.name = "Noto Sans CJK SC"
        style.font.size = Pt(10.5 if style_name == "Normal" else 14)
        style.font.color.rgb = (
            theme_colors["ink"] if style_name == "Normal" else theme_colors["primary"]
        )
        style._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Noto Sans CJK SC"
        )
    if "Report Signature" not in document.styles:
        signature_style = document.styles.add_style("Report Signature", WD_STYLE_TYPE.PARAGRAPH)
        signature_style.font.name = "Noto Sans CJK SC"
        signature_style._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Noto Sans CJK SC"
        )

    paragraphs = document.paragraphs

    def marker_index(name: str) -> int:
        return next(
            index for index, paragraph in enumerate(paragraphs) if paragraph._p is markers[name]._p
        )

    cover_end_index = marker_index("cover_end")
    toc_start_index = marker_index("toc_field_start")
    toc_end_index = marker_index("toc_field_end")
    body_start_index = marker_index("body_start")
    for paragraph in paragraphs[:cover_end_index]:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph = next(
        (item for item in paragraphs[:cover_end_index] if item.text == context["title"]), None
    )
    if title_paragraph is None:
        raise ReportFailure("Word 封面缺少报告标题")
    title_paragraph.style = document.styles["Title"]
    toc_title = next(
        (item for item in paragraphs[cover_end_index:toc_start_index] if item.text == "目录"), None
    )
    if toc_title is None:
        raise ReportFailure("Word 缺少目录标题")
    toc_title.style = document.styles["Title"]

    body_headings: list[Any] = []
    search_index = body_start_index + 1
    for heading_index, item in enumerate(context["headingNumbers"], start=1):
        expected_text = f"{item['number']} {item['title']}"
        found = next(
            (
                (index, paragraph)
                for index, paragraph in enumerate(paragraphs[search_index:], start=search_index)
                if paragraph.text.strip() == expected_text
            ),
            None,
        )
        if found is None:
            raise ReportFailure("Word 正文标题与服务端编号映射不一致")
        search_index, heading = found
        search_index += 1
        heading.style = document.styles[f"Heading {item['level'] - 1}"]
        bookmark_name = item["anchor"].replace("-", "_")
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(1000 + heading_index))
        start.set(qn("w:name"), bookmark_name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(1000 + heading_index))
        insert_at = 1 if heading._p.pPr is not None else 0
        heading._p.insert(insert_at, start)
        heading._p.append(end)
        body_headings.append(heading)

    toc_entries = [
        paragraph
        for paragraph in paragraphs[toc_start_index + 1 : toc_end_index]
        if paragraph.text.strip()
    ]
    if len(toc_entries) != len(context["headingNumbers"]):
        raise ReportFailure("Word 缓存目录与正式标题不一致")
    for paragraph, item in zip(toc_entries, context["headingNumbers"], strict=True):
        clear_paragraph(paragraph)
        toc_style = f"TOC {item['level'] - 1}"
        if toc_style in document.styles:
            paragraph.style = document.styles[toc_style]
        usable_width = sections[1].page_width - sections[1].left_margin - sections[1].right_margin
        paragraph.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
        hyperlink = OxmlElement("w:hyperlink")
        bookmark_name = item["anchor"].replace("-", "_")
        hyperlink.set(qn("w:anchor"), bookmark_name)
        hyperlink.set(qn("w:history"), "1")
        run = OxmlElement("w:r")
        run_properties = OxmlElement("w:rPr")
        run_style = OxmlElement("w:rStyle")
        run_style.set(qn("w:val"), "Hyperlink")
        run_properties.append(run_style)
        run.append(run_properties)
        text = OxmlElement("w:t")
        text.text = f"{item['number']} {item['title']}"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        paragraph.add_run("\t")
        field_run(
            paragraph,
            f"PAGEREF {bookmark_name} \\h",
        )

    clear_paragraph(markers["toc_field_start"])
    field_run(
        markers["toc_field_start"],
        'TOC \\o "1-3" \\h \\z \\u',
        result="",
        close=False,
    )
    clear_paragraph(markers["toc_field_end"])
    field_end(markers["toc_field_end"])
    clear_paragraph(markers["body_start"])

    def clear_story(story: Any) -> Any:
        for table in list(story.tables):
            story._element.remove(table._element)
        paragraph = story.paragraphs[0]
        clear_paragraph(paragraph)
        for extra in list(story.paragraphs[1:]):
            story._element.remove(extra._element)
        return paragraph

    def add_template(paragraph: Any, left: str, right: str) -> None:
        usable_width = sections[1].page_width - sections[1].left_margin - sections[1].right_margin
        paragraph.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)

        def append(value: str) -> None:
            for literal, field_name, _format_spec, _conversion in Formatter().parse(value):
                if literal:
                    paragraph.add_run(literal)
                if field_name == "title":
                    paragraph.add_run(context["title"])
                elif field_name == "organization":
                    paragraph.add_run(context["organizationName"])
                elif field_name in _WORD_PAGE_FIELDS:
                    field_run(paragraph, _WORD_PAGE_FIELDS[field_name])

        append(left)
        paragraph.add_run("\t")
        append(right)
        for run in paragraph.runs:
            run.font.name = "Noto Sans CJK SC"
            run.font.size = Pt(8)
            run.font.color.rgb = theme_colors["muted"]
            run._element.get_or_add_rPr().get_or_add_rFonts().set(
                qn("w:eastAsia"), "Noto Sans CJK SC"
            )

    def add_watermark(header: Any, text: str, shape_id: int) -> None:
        escaped = html.escape(text, quote=True)
        paragraph = header.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        watermark = parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>'
            f'<v:shape id="PowerPlusWaterMarkObject{shape_id}" '
            'type="#_x0000_t136" '
            'style="position:absolute;width:430pt;height:90pt;rotation:315;z-index:-251654144;'
            'mso-position-horizontal:center;mso-position-vertical:center" '
            f'fillcolor="{REPORT_VISUAL_THEME["muted"]}" stroked="f">'
            '<v:fill opacity="0.045"/>'
            '<v:textpath style="font-family:Noto Sans CJK SC;font-size:26pt" '
            f'string="{escaped}"/></v:shape></w:pict></w:r>'
        )
        paragraph._p.append(watermark)

    for index, section in enumerate(sections):
        if index == 0:
            continue
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        header_paragraph = clear_story(section.header)
        footer_paragraph = clear_story(section.footer)
        add_template(
            header_paragraph,
            layout["headerLeft"],
            layout["headerRight"],
        )
        add_template(
            footer_paragraph,
            layout["footerLeft"],
            layout["footerRight"],
        )
        add_watermark(section.header, context["watermarkText"], index)

    for table in document.tables:
        if "Table Grid" in document.styles:
            table.style = "Table Grid"
        table_properties = table._tbl.tblPr
        for existing in table_properties.findall(qn("w:tblBorders")):
            table_properties.remove(existing)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), str(REPORT_VISUAL_THEME["grid"]).removeprefix("#"))
            borders.append(border)
        table_properties.append(borders)
        if table.rows:
            row_properties = table.rows[0]._tr.get_or_add_trPr()
            table_header = OxmlElement("w:tblHeader")
            table_header.set(qn("w:val"), "true")
            row_properties.append(table_header)
            for cell in table.rows[0].cells:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), str(REPORT_VISUAL_THEME["surface"]).removeprefix("#"))
                cell._tc.get_or_add_tcPr().append(shading)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = theme_colors["primary"]
    available_width = sections[-1].page_width - sections[-1].left_margin - sections[-1].right_margin
    maximum_height = Mm(180)
    for shape in document.inline_shapes:
        new_width, new_height = _fit_image_dimensions(
            shape.width, shape.height, available_width, maximum_height
        )
        if (new_width, new_height) != (shape.width, shape.height):
            shape.width = new_width
            shape.height = new_height

    for paragraph in paragraphs[body_start_index + 1 :]:
        if paragraph.text.strip() in {context["organizationName"], context["generatedDate"]}:
            paragraph.style = document.styles["Report Signature"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    settings = document.settings._element
    for existing in settings.findall(qn("w:updateFields")):
        settings.remove(existing)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    document.core_properties.title = context["title"]
    document.core_properties.subject = context["periodLabel"]
    document.core_properties.author = context["organizationName"]
    document.core_properties.last_modified_by = context["generatedByLabel"]
    document.core_properties.keywords = context["generatedByLabel"]
    document.core_properties.modified = datetime.now(UTC)

    postprocessed = path.with_name("render.postprocessed.docx")
    document.save(str(postprocessed))
    if postprocessed.stat().st_size > MAX_DOCX_BYTES:
        raise ReportFailure("Word 文件不能超过 200 MiB")
    os.replace(postprocessed, path)


def _validate_docx_structure(
    path: Path,
    *,
    expected_sections: list[dict[str, str]],
    expected_headings: list[dict[str, Any]],
    expected_image_count: int | None,
) -> dict[str, Any]:
    if not path.is_file() or not 1 <= path.stat().st_size <= MAX_DOCX_BYTES:
        raise ReportFailure("Word 文件不存在或超过 200 MiB")
    try:
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            if any(
                name == forbidden or name.startswith(forbidden)
                for name in names
                for forbidden in _DOCX_FORBIDDEN_PARTS
            ):
                raise ReportFailure("Word 包含宏、OLE 或 ActiveX 内容")
            for relationship_name in (name for name in names if name.endswith(".rels")):
                relationships = ElementTree.fromstring(package.read(relationship_name))
                if any(item.attrib.get("TargetMode") == "External" for item in relationships):
                    raise ReportFailure("Word 包含外部关系")
            required = {"word/document.xml", "word/settings.xml", "[Content_Types].xml"}
            if not required.issubset(names):
                raise ReportFailure("Word OOXML 结构不完整")
            document_content = package.read("word/document.xml")
            document_root = ElementTree.fromstring(document_content)
            document_xml = document_content.decode("utf-8")
            image_names = [name for name in names if name.startswith("word/media/")]
    except (OSError, KeyError, UnicodeError, zipfile.BadZipFile, ElementTree.ParseError) as error:
        raise ReportFailure("Word OOXML 无法解析") from error
    bookmark_names = {item["anchor"].replace("-", "_") for item in expected_headings}
    native_toc_present = (
        'TOC \\o "1-3"' in document_xml
        and "PAGEREF report_" in document_xml
        and all(name in document_xml for name in bookmark_names)
    )
    toc_entry_count = sum(name in document_xml for name in bookmark_names)
    word_namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    section_properties = document_root.findall(f".//{word_namespace}sectPr")
    if expected_image_count is not None and len(image_names) < expected_image_count:
        raise ReportFailure("Word 未完整嵌入图表")
    return {
        "nativeTocPresent": native_toc_present,
        "sectionCount": len(section_properties),
        "tocEntryCount": toc_entry_count,
        "embeddedImageCount": len(image_names),
        "externalRelationshipCount": 0,
    }


def _validate_docx_rendering(
    path: Path,
    directory: Path,
    *,
    context: dict[str, Any],
    layout: dict[str, str],
) -> dict[str, Any]:
    try:
        import pypdf
        from PIL import Image
    except ImportError as error:
        raise ReportFailure("Word 可渲染性验收依赖不可用") from error
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if libreoffice is None:
        raise ReportFailure("Word 验收命令 LibreOffice 不可用")
    directory.mkdir(mode=0o700)
    output = directory / "output"
    profile = directory / "profile"
    raster = directory / "raster"
    output.mkdir()
    profile.mkdir()
    raster.mkdir()
    try:
        process = subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation={profile.resolve().as_uri()}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output),
                str(path),
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=DOCX_VALIDATION_TIMEOUT_SECONDS,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ReportFailure("Word 经 LibreOffice 转换失败或超时") from error
    converted = output / f"{path.stem}.pdf"
    if process.returncode != 0 or not converted.is_file():
        raise ReportFailure("Word 经 LibreOffice 转换失败")
    try:
        reader = pypdf.PdfReader(str(converted))
    except pypdf.errors.PdfReadError as error:
        raise ReportFailure("LibreOffice 转换结果无法打开") from error
    if not 1 <= len(reader.pages) <= MAX_PDF_PAGES:
        raise ReportFailure("Word 转换页数超出边界")
    prefix = raster / "page"
    try:
        rendered = subprocess.run(
            ["pdftoppm", "-gray", "-r", "72", "-png", str(converted), str(prefix)],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=DOCX_VALIDATION_TIMEOUT_SECONDS,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ReportFailure("Word 转换结果栅格化失败") from error
    rendered_pages = sorted(
        raster.glob("page-*.png"), key=lambda item: int(item.stem.rsplit("-", 1)[-1])
    )
    if rendered.returncode != 0 or len(rendered_pages) != len(reader.pages):
        raise ReportFailure("Word 转换结果栅格化失败")
    extracted_pages: list[str] = []
    page_image_counts: list[int] = []
    image_count = 0
    for page, rendered_page in zip(reader.pages, rendered_pages, strict=True):
        page_text = page.extract_text() or ""
        extracted_pages.append(page_text)
        current_image_count = len(page.images)
        page_image_counts.append(current_image_count)
        image_count += current_image_count
        with Image.open(rendered_page) as image:
            image.verify()
    first_section_title = "".join(context["sections"][0]["title"].split())
    report_title = "".join(context["title"].split())
    first_section_pages = [
        index
        for index, page_text in enumerate(extracted_pages, start=1)
        if index > 1
        and first_section_title in "".join(page_text.split()).replace(report_title, "", 1)
    ]
    # LibreOffice 会在转换后的 PDF 文本层为中文标题插入布局空格。该位置只用于
    # 识别页码装饰，不再把目录缓存方式和正文分页形态作为发布门禁。
    body_start_page = first_section_pages[-1] if first_section_pages else 2
    blank_pages: list[int] = []
    for index, (page_text, current_image_count) in enumerate(
        zip(extracted_pages, page_image_counts, strict=True), start=1
    ):
        substantive_text = "".join(page_text.split())
        if index > 1:
            page_value, pages_value = _page_number_context(
                index,
                body_start_page=body_start_page,
                physical_page_count=len(reader.pages),
            )
            decorations = [
                context["watermarkText"],
                *(
                    _formatted_page_text(
                        value,
                        title=context["title"],
                        organization=context["organizationName"],
                        page=page_value,
                        pages=pages_value,
                    )
                    for value in layout.values()
                    if value
                ),
            ]
            # Writer 转换后的页眉、页脚和 VML 水印也会进入 PDF 文本层；只含这些
            # 服务端装饰的页面仍是空白正文页，不能借装饰绕过发布门禁。
            for decoration in decorations:
                substantive_text = substantive_text.replace("".join(decoration.split()), "", 1)
        if not substantive_text and current_image_count == 0:
            blank_pages.append(index)
    extracted_text = "\n".join(extracted_pages)
    compact_extracted_text = "".join(extracted_text.split())
    required_text = (
        context["title"],
        context["periodLabel"],
        context["organizationName"],
        context["generatedByLabel"],
        context["generatedDate"],
        *(item["title"] for item in context["sections"]),
    )
    if blank_pages or any(
        "".join(item.split()) not in compact_extracted_text for item in required_text
    ):
        raise ReportFailure("Word 转换结果缺少正式内容或包含空白页")
    if any(marker in extracted_text for marker in _WORD_MARKERS.values()):
        raise ReportFailure("Word 显示了内部版式标记")
    return {
        "convertedPageCount": len(reader.pages),
        "blankPages": blank_pages,
        "renderedImageCount": image_count,
    }
