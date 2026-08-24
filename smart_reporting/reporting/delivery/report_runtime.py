"""受限 Markdown 报表运行时；由 Reporting Workflow 在 Daytona 中执行。"""

from __future__ import annotations

import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from collections.abc import Mapping, Sequence
from contextlib import contextmanager
from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path, PurePosixPath
from string import Formatter
from typing import Any
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree

MAX_MARKDOWN_BYTES = 1024 * 1024
MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_TOTAL_IMAGE_BYTES = 50 * 1024 * 1024
MAX_RESULT_BYTES = 64 * 1024
MAX_DATASET_PATHS = 20
MAX_PDF_BYTES = 200 * 1024 * 1024
MAX_PDF_PAGES = 200
MAX_DOCX_BYTES = 200 * 1024 * 1024
PDF_VALIDATION_TIMEOUT_SECONDS = 540
DOCX_RENDER_TIMEOUT_SECONDS = 540
DOCX_VALIDATION_TIMEOUT_SECONDS = 540
IMAGE_SUFFIXES = {".gif", ".jpeg", ".jpg", ".png", ".webp"}
DEFAULT_PAGE_LAYOUT = {
    "headerLeft": "{organization}",
    "headerRight": "{title}",
    "footerLeft": "企业智能运营报表",
    "footerRight": "第 {page} / {pages} 页",
}
# 视觉主题属于服务端渲染契约，而不是模型自由生成的正文内容。PDF、Word 与
# Coding 生成的图表共同引用这一份科技蓝颜色事实，避免封面、正文和图表各自选色；
# 图表类型、系列数量和强调对象仍由模型根据数据决定，琥珀/绿色仅用于语义强调。
REPORT_VISUAL_THEME = {
    "name": "enterprise-tech-blue",
    "primary": "#0B4F8A",
    "accent": "#007EA7",
    "highlight": "#F2B134",
    "ink": "#1B2A41",
    "muted": "#5B6B7A",
    "grid": "#C7D7E5",
    "surface": "#EDF5FC",
    "chartPalette": [
        "#0B4F8A",
        "#007EA7",
        "#2F80ED",
        "#56B4E9",
        "#F2B134",
        "#2E9F6B",
        "#7A5AF8",
        "#D66B3D",
    ],
}
PAGE_LAYOUT_FIELDS = frozenset(DEFAULT_PAGE_LAYOUT)
PAGE_LAYOUT_PLACEHOLDERS = frozenset({"title", "organization", "page", "pages"})
_DOCX_FORBIDDEN_PARTS = ("word/vbaProject.bin", "word/embeddings/", "word/activeX/")
_WORD_MARKERS = {
    "cover_end": "__REPORT_COVER_END__",
    "toc_field_start": "__REPORT_TOC_FIELD_START__",
    "toc_field_end": "__REPORT_TOC_FIELD_END__",
    "toc_end": "__REPORT_TOC_END__",
    "body_start": "__REPORT_BODY_START__",
}
_WORD_PAGE_FIELDS = {"page": "PAGE", "pages": "SECTIONPAGES"}
_CITATION_MARKER = re.compile(r"\[\[citation:([^\]\r\n]+)\]\]")
_SECTION_MARKER = re.compile(r"\[\[section:([^\]\r\n]+)\]\]")
_ANALYSIS_MARKER = re.compile(r"\[\[analysis:([^\]\r\n]+)\]\]")
_TABLE_MARKER = re.compile(r"\[\[/?table:[^\]\r\n]+\]\]")
# markdown-it 遵循 CommonMark 的 Unicode 标点边界规则。中文引号/括号紧邻
# `**` 时，模型生成的粗体标记可能不会被识别，最终会原样进入成稿；只对
# 含中文或数值的成对标记补充解析所需的边界，避免改写数学表达式或孤立星号。
_CJK_STRONG_MARKER = re.compile(
    r"(?P<left>[^\s*`])(?<!\*)(?P<open>\*\*)(?P<content>[^*\r\n`]*[\u3400-\u9fff][^*\r\n`]*?)(?P<close>\*\*)(?P<right>[^\s*`])"
)
_SPACED_CJK_STRONG_MARKER = re.compile(
    r"(?<![\*`])\*\*(?P<content>[\u3400-\u9fff“「『【（《〈〔［｛][^*\r\n`]*?)\s+\*\*(?![\*`])(?=\s|[，。；：、！？）】》〉〕］｝])"
)
_SPACED_VALUE_STRONG_MARKER = re.compile(
    r"(?<![\*`])\*\*(?P<content>[^*\r\n`]*(?:%|％|亿元|万元|元|万)[^*\r\n`]*)\*\*(?![\*`])"
)


class ReportFailure(ValueError):
    pass


def _normalize_cjk_strong_markers(markdown: str) -> str:
    """让报告中的中文/数值粗体文本进入 CommonMark 的强调解析路径。"""

    def trim_boundaries(match: re.Match[str]) -> str:
        content = match["content"]
        trimmed = content.strip()
        if not trimmed:
            return match[0]
        return f"**{trimmed}**"

    def add_boundaries(match: re.Match[str]) -> str:
        return f"{match['left']} {match['open']}{match['content']}{match['close']} {match['right']}"

    normalized = _CJK_STRONG_MARKER.sub(add_boundaries, markdown)
    normalized = _SPACED_CJK_STRONG_MARKER.sub(trim_boundaries, normalized)
    return _SPACED_VALUE_STRONG_MARKER.sub(trim_boundaries, normalized)


def _pdf_markdown(
    markdown: str,
    presentations: Any,
) -> tuple[str, list[dict[str, Any]]]:
    marker_ids = tuple(dict.fromkeys(_CITATION_MARKER.findall(markdown)))
    if not marker_ids:
        visible_markdown = _ANALYSIS_MARKER.sub("", _SECTION_MARKER.sub("", markdown))
        return _TABLE_MARKER.sub("", visible_markdown), []
    if not isinstance(presentations, list):
        raise ReportFailure("PDF 缺少服务端实际引用展示信息")
    normalized: list[dict[str, Any]] = []
    for index, item in enumerate(presentations, start=1):
        if not isinstance(item, dict) or set(item) != {
            "citationId",
            "label",
            "coverageItems",
        }:
            raise ReportFailure("PDF 实际引用展示信息无效")
        citation_id = item.get("citationId")
        label = item.get("label")
        coverage_items = item.get("coverageItems")
        if (
            not isinstance(citation_id, str)
            or not citation_id
            or not isinstance(label, str)
            or not 1 <= len(label) <= 200
            or not isinstance(coverage_items, list)
            or len(coverage_items) > 100
        ):
            raise ReportFailure("PDF 实际引用展示信息与 Markdown 顺序不一致")
        normalized_coverage: list[dict[str, Any]] = []
        for coverage_index, coverage in enumerate(coverage_items, start=1):
            if not isinstance(coverage, dict) or set(coverage) != {"label", "periods"}:
                raise ReportFailure("PDF 实际引用数据覆盖信息无效")
            coverage_label = coverage.get("label")
            periods = coverage.get("periods")
            if (
                not isinstance(coverage_label, str)
                or not 1 <= len(coverage_label) <= 200
                or not isinstance(periods, list)
                or len(periods) > 1200
                or any(not isinstance(period, str) or len(period) > 32 for period in periods)
            ):
                raise ReportFailure("PDF 实际引用数据覆盖信息无效")
            normalized_coverage.append(
                {
                    "label": coverage_label or f"来源项 {coverage_index}",
                    "periods": periods,
                }
            )
        normalized.append(
            {
                "citationId": citation_id,
                "alias": f"[引用 {index:03d}]",
                "label": label,
                "coverageItems": normalized_coverage,
            }
        )
    presentation_ids = [item["citationId"] for item in normalized]
    # Manifest 绑定全部授权 DatasetLineage，而专题正文可以只直接使用其中一部分。
    # 展示 ledger 因而允许是 Markdown marker 的超集，但 marker 仍必须逐个来自服务端
    # ledger，且 ledger 自身不能重复，防止模型注入未知 citation 或伪造展示标签。
    if len(presentation_ids) != len(set(presentation_ids)) or not set(marker_ids).issubset(
        presentation_ids
    ):
        raise ReportFailure("PDF 实际引用展示信息与 Markdown 引用不一致")
    # Citation 绑定仍由服务端校验并写入渲染回执，但 PDF 展示层不泄露机器 marker、
    # 可读别名或来源附录；权威 Markdown 保持原样，供产物协议和血缘验收使用。
    visible_markdown = _ANALYSIS_MARKER.sub(
        "", _SECTION_MARKER.sub("", _CITATION_MARKER.sub("", markdown))
    )
    return _TABLE_MARKER.sub("", visible_markdown), normalized


def _relative_path(value: str, suffix: str | None = None) -> PurePosixPath:
    if not isinstance(value, str) or not value or "\\" in value:
        raise ReportFailure("路径必须是当前工作区的相对路径")
    path = PurePosixPath(value)
    if (
        path.is_absolute()
        or ".." in path.parts
        or any(part in ("", ".") for part in path.parts)
        or (suffix is not None and path.suffix.lower() != suffix)
    ):
        expected = f" {suffix}" if suffix is not None else ""
        raise ReportFailure(f"路径必须是当前工作区内的{expected}相对路径")
    return path


def _reject_symlinks(workspace: Path, path: Path) -> None:
    relative = path.relative_to(workspace)
    current = workspace
    for part in relative.parts:
        current /= part
        if current.is_symlink():
            raise ReportFailure("报表路径不能包含符号链接")


def _input_path(workspace: Path, value: str, suffix: str) -> Path:
    relative = _relative_path(value, suffix)
    path = workspace.joinpath(*relative.parts)
    _reject_symlinks(workspace, path)
    if not path.is_file():
        raise ReportFailure("报表源文件不存在或不是普通文件")
    return path


def _output_path(workspace: Path, value: str) -> Path:
    relative = _relative_path(value, ".pdf")
    path = workspace.joinpath(*relative.parts)
    parent = path.parent
    _reject_symlinks(workspace, parent)
    if not parent.is_dir() and not parent.parent.is_dir():
        raise ReportFailure("PDF revision 父目录不存在")
    if path.exists() or path.is_symlink():
        raise ReportFailure("PDF 输出文件已经存在")
    return path


def _word_output_path(workspace: Path, value: str) -> Path:
    relative = _relative_path(value, ".docx")
    path = workspace.joinpath(*relative.parts)
    _reject_symlinks(workspace, path.parent)
    if not path.parent.is_dir() and not path.parent.parent.is_dir():
        raise ReportFailure("Word revision 父目录不存在")
    if path.exists() or path.is_symlink():
        raise ReportFailure("Word 输出文件已经存在")
    return path


def _temporary_pdf_path(value: str) -> Path:
    if not isinstance(value, str) or not value.startswith("/tmp/workspace-report-"):
        raise ReportFailure("PDF 临时路径无效")
    path = Path(value)
    if path.name != "render.pdf" or path.parent.parent != Path("/tmp"):
        raise ReportFailure("PDF 临时路径无效")
    if path.parent.is_symlink() or path.exists() or path.is_symlink():
        raise ReportFailure("PDF 临时路径已经存在")
    try:
        path.parent.mkdir(mode=0o700)
    except FileExistsError as error:
        raise ReportFailure("PDF 临时路径已经存在") from error
    return path


def _temporary_docx_path(pdf_path: Path) -> Path:
    path = pdf_path.with_name("render.docx")
    if path.exists() or path.is_symlink():
        raise ReportFailure("Word 临时路径已经存在")
    return path


def _validation_directory(value: str) -> Path:
    if not isinstance(value, str) or not value.startswith("/tmp/workspace-report-"):
        raise ReportFailure("PDF 验收临时路径无效")
    path = Path(value)
    if path.parent != Path("/tmp") or not path.name.endswith("-validate"):
        raise ReportFailure("PDF 验收临时路径无效")
    if path.exists() or path.is_symlink():
        raise ReportFailure("PDF 验收临时路径已经存在")
    try:
        path.mkdir(mode=0o700)
    except FileExistsError as error:
        raise ReportFailure("PDF 验收临时路径已经存在") from error
    return path


@contextmanager
def _cleanup_directory(path: Path):
    try:
        yield
    finally:
        shutil.rmtree(path, ignore_errors=True)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _check_image_signature(path: Path) -> None:
    header = path.read_bytes()[:12]
    suffix = path.suffix.lower()
    valid = {
        ".png": header.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": header.startswith(b"\xff\xd8\xff"),
        ".jpeg": header.startswith(b"\xff\xd8\xff"),
        ".gif": header.startswith((b"GIF87a", b"GIF89a")),
        ".webp": header.startswith(b"RIFF") and header[8:12] == b"WEBP",
    }.get(suffix, False)
    if not valid:
        raise ReportFailure("Markdown 图片格式或文件签名无效")


def _page_layout(value: Any) -> dict[str, str]:
    if value is None:
        return dict(DEFAULT_PAGE_LAYOUT)
    if not isinstance(value, dict) or set(value) - PAGE_LAYOUT_FIELDS:
        raise ReportFailure("PDF 页面格式无效")
    layout = dict(DEFAULT_PAGE_LAYOUT)
    for key, item in value.items():
        if not isinstance(item, str) or len(item) > 200:
            raise ReportFailure("PDF 页面格式无效")
        if any(ord(character) < 32 and character != "\t" for character in item):
            raise ReportFailure("PDF 页面格式包含控制字符")
        try:
            parsed = tuple(Formatter().parse(item))
        except ValueError as error:
            raise ReportFailure("PDF 页面格式无效") from error
        if any(
            field_name not in PAGE_LAYOUT_PLACEHOLDERS or format_spec or conversion
            for _literal, field_name, format_spec, conversion in parsed
            if field_name is not None
        ):
            raise ReportFailure("PDF 页面格式包含不受支持的占位符")
        layout[key] = item
    footer = f"{layout['footerLeft']}\n{layout['footerRight']}"
    if "{page}" not in footer or "{pages}" not in footer:
        raise ReportFailure("PDF 页脚必须包含当前页和总页数")
    return layout


def _document_context(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict) or set(value) != {
        "title",
        "periodLabel",
        "organizationName",
        "generatedByLabel",
        "watermarkText",
        "generatedDate",
        "sectionNumbers",
        "sections",
        "headingNumbers",
    }:
        raise ReportFailure("报告缺少服务端文档展示契约")
    normalized: dict[str, Any] = {}
    for key, maximum in (
        ("title", 300),
        ("periodLabel", 100),
        ("organizationName", 200),
        ("generatedByLabel", 200),
        ("watermarkText", 200),
        ("generatedDate", 32),
    ):
        item = value.get(key)
        if (
            not isinstance(item, str)
            or not item.strip()
            or len(item) > maximum
            or any(ord(character) < 32 for character in item)
        ):
            raise ReportFailure("服务端文档展示契约无效")
        normalized[key] = item.strip()
    try:
        datetime.strptime(normalized["generatedDate"], "%Y-%m-%d")
    except ValueError as error:
        raise ReportFailure("服务端生成日期无效") from error
    sections = value.get("sections")
    if not isinstance(sections, list) or not 1 <= len(sections) <= 100:
        raise ReportFailure("服务端正式章节契约无效")
    normalized_sections: list[dict[str, str]] = []
    for item in sections:
        if not isinstance(item, dict) or set(item) != {"code", "sectionNumber", "title"}:
            raise ReportFailure("服务端正式章节契约无效")
        code = item.get("code")
        title = item.get("title")
        section_number = item.get("sectionNumber")
        if (
            not isinstance(code, str)
            or re.fullmatch(r"[a-z][a-z0-9_]{0,127}", code) is None
            or not isinstance(title, str)
            or not title.strip()
            or len(title) > 300
            or any(ord(character) < 32 for character in title)
            or not isinstance(section_number, str)
            or re.fullmatch(r"[1-9][0-9]*", section_number) is None
        ):
            raise ReportFailure("服务端正式章节契约无效")
        normalized_sections.append(
            {"code": code, "sectionNumber": section_number, "title": title.strip()}
        )
    codes = [item["code"] for item in normalized_sections]
    titles = [item["title"] for item in normalized_sections]
    if len(codes) != len(set(codes)) or len(titles) != len(set(titles)):
        raise ReportFailure("服务端正式章节契约包含重复项")
    section_numbers = value.get("sectionNumbers")
    expected_section_numbers = [str(index) for index in range(1, len(sections) + 1)]
    if (
        section_numbers != expected_section_numbers
        or [item["sectionNumber"] for item in normalized_sections] != expected_section_numbers
    ):
        raise ReportFailure("服务端正式章节编号不连续")
    headings = value.get("headingNumbers")
    if not isinstance(headings, list) or not headings:
        raise ReportFailure("报告缺少服务端标题编号映射")
    normalized_headings: list[dict[str, Any]] = []
    for item in headings:
        if not isinstance(item, dict) or set(item) != {
            "level",
            "number",
            "title",
            "sectionCode",
            "anchor",
        }:
            raise ReportFailure("服务端标题编号映射无效")
        if (
            item["level"] not in {2, 3, 4}
            or not isinstance(item["number"], str)
            or re.fullmatch(r"[1-9][0-9]*(?:\.[1-9][0-9]*){0,2}", item["number"]) is None
            or not isinstance(item["title"], str)
            or not item["title"].strip()
            or item["sectionCode"] not in codes
            or not isinstance(item["anchor"], str)
            or re.fullmatch(r"report-(?:section|heading)-[a-z0-9_-]+", item["anchor"]) is None
        ):
            raise ReportFailure("服务端标题编号映射无效")
        normalized_headings.append({**item, "title": item["title"].strip()})
    primary = [item for item in normalized_headings if item["level"] == 2]
    if (
        [item["sectionCode"] for item in primary] != codes
        or [item["number"] for item in primary] != expected_section_numbers
        or len({item["anchor"] for item in normalized_headings}) != len(normalized_headings)
    ):
        raise ReportFailure("服务端标题编号映射与正式章节不一致")
    normalized["sectionNumbers"] = expected_section_numbers
    normalized["sections"] = normalized_sections
    normalized["headingNumbers"] = normalized_headings
    return normalized


def _markdown_title(tokens: list[Any]) -> str:
    for index, token in enumerate(tokens[:-1]):
        if token.type == "heading_open" and token.tag == "h1":
            title = str(getattr(tokens[index + 1], "content", "") or "").strip()
            if title:
                return title[:300]
    return "智能运营报表"


def _bind_heading_anchors(tokens: list[Any], headings_contract: list[dict[str, Any]]) -> None:
    headings: list[tuple[Any, int, str]] = []
    for index, token in enumerate(tokens[:-1]):
        if token.type == "heading_open" and token.tag in {"h2", "h3", "h4"}:
            headings.append(
                (
                    token,
                    int(token.tag[1:]),
                    "".join(
                        str(item.content or "")
                        for item in getattr(tokens[index + 1], "children", ()) or ()
                        if item.type in {"text", "code_inline", "image"}
                    ).strip(),
                )
            )
    expected = [(item["level"], f"{item['number']} {item['title']}") for item in headings_contract]
    if [(level, title) for _token, level, title in headings] != expected:
        raise ReportFailure("Markdown 标题顺序与服务端编号映射不一致")
    for (token, _level, _title), item in zip(headings, headings_contract, strict=True):
        token.attrSet("id", item["anchor"])


def _body_tokens(tokens: list[Any]) -> list[Any]:
    for index, token in enumerate(tokens[:-2]):
        if token.type == "heading_open" and token.tag == "h1":
            return [*tokens[:index], *tokens[index + 3 :]]
    return tokens


def _formatted_page_text(
    template: str,
    *,
    title: str,
    organization: str,
    page: str | int,
    pages: str | int,
) -> str:
    return template.format(title=title, organization=organization, page=page, pages=pages)


def _has_page_layout(
    text: str,
    layout: dict[str, str],
    *,
    title: str,
    organization: str,
    page: str | int,
    pages: str | int,
) -> bool:
    compact = "".join(text.split())
    expected = (
        _formatted_page_text(
            value,
            title=title,
            organization=organization,
            page=page,
            pages=pages,
        )
        for value in layout.values()
        if value
    )
    return all("".join(value.split()) in compact for value in expected)


def _semantic_documents(
    body: str,
    *,
    context: dict[str, Any],
    layout: dict[str, str],
    toc_page_numbers: Mapping[str, int] | None = None,
) -> tuple[str, str]:
    theme = REPORT_VISUAL_THEME
    title = html.escape(context["title"])
    period = html.escape(context["periodLabel"])
    organization = html.escape(context["organizationName"])
    generated_label = html.escape(context["generatedByLabel"])
    generated_date = html.escape(context["generatedDate"])
    toc = "".join(
        f'<p class="toc-entry toc-level-{item["level"]}"><a href="#{item["anchor"]}">'
        f'<span class="toc-title">{html.escape(item["number"] + " " + item["title"])}</span>'
        '<span class="toc-leader"></span>'
        f'<span class="toc-page">{toc_page_numbers.get(item["anchor"], "") if toc_page_numbers is not None else ""}</span>'
        "</a></p>"
        for item in context["headingNumbers"]
    )
    shared = (
        f'<section class="report-cover"><h1>{title}</h1>'
        f'<p class="report-period">分析期间：{period}</p>'
        f'<p class="report-organization">{organization}</p>'
        f'<p class="report-generated">{generated_label}</p></section>'
        f'<section class="report-toc"><h1>目录</h1>{toc}</section>'
        f'<main class="report-body">{body}'
        f'<footer class="report-signature"><p>{organization}</p><p>{generated_date}</p>'
        "</footer></main>"
    )
    pdf_css = (
        "@page cover{size:A4;margin:20mm 18mm 22mm;}"
        "@page toc{size:A4;margin:20mm 18mm 22mm;}"
        "@page body{size:A4;margin:20mm 18mm 22mm;}"
        "body{font-family:'Noto Sans CJK SC','Noto Sans CJK JP',sans-serif;"
        "font-size:10.5pt;line-height:1.65;color:"
        f"{theme['ink']}"
        ";margin:0}"
        ".report-cover{page:cover;break-after:page;min-height:250mm;position:relative;"
        "display:flex;flex-direction:column;align-items:center;text-align:center}"
        ".report-cover h1{color:"
        f"{theme['primary']}"
        ";margin-top:58mm;font-size:28pt;border:0;"
        "padding:0;max-width:150mm}"
        ".report-period{margin-top:18mm;font-size:13pt;color:"
        f"{theme['accent']}"
        "}"
        ".report-organization{margin-top:24mm;font-size:12pt;color:"
        f"{theme['ink']}"
        "}"
        ".report-generated{position:absolute;bottom:8mm;font-size:9pt;color:"
        f"{theme['muted']}"
        "}"
        ".report-toc{page:toc;break-after:page;min-height:240mm}"
        ".report-toc h1{color:"
        f"{theme['primary']}"
        ";font-size:22pt;border-bottom:1.5pt solid "
        f"{theme['accent']}"
        ";"
        "padding-bottom:5mm;margin-bottom:8mm}"
        ".toc-entry{margin:0 0 3mm}.toc-entry a{color:"
        f"{theme['ink']}"
        ";text-decoration:none;display:flex;align-items:baseline;gap:2mm}"
        ".toc-level-3{padding-left:6mm}.toc-level-4{padding-left:12mm}"
        ".toc-title{min-width:0}.toc-leader{flex:1;border-bottom:0.5pt dotted "
        f"{theme['grid']}"
        ";transform:translateY(-1.5mm)}.toc-page{min-width:3ch;text-align:right}"
        ".report-body{page:body}"
        "h2{color:"
        f"{theme['primary']}"
        ";border-left:3pt solid "
        f"{theme['accent']}"
        ";font-size:16pt;padding-left:3mm}"
        "h3{font-size:12.5pt;color:"
        f"{theme['ink']}"
        "}h2,h3{page-break-after:avoid;break-after:avoid}"
        "p,li{orphans:3;widows:3}table{width:100%;border-collapse:collapse;margin:10px 0}"
        "thead{display:table-header-group}tr{break-inside:avoid}"
        "th,td{border:0.6pt solid "
        f"{theme['grid']}"
        ";padding:5px 7px;text-align:left}"
        "th{background:"
        f"{theme['surface']}"
        ";color:"
        f"{theme['primary']}"
        "}tbody tr:nth-child(even){background:#F9FAFB}"
        "img{display:block;max-width:100%;height:auto;margin:12px auto;break-inside:avoid}"
        "p:has(>img){break-after:avoid;margin-bottom:1mm}"
        "p:has(>img)+p{break-before:avoid;margin-top:0;text-align:center;color:"
        f"{theme['muted']}"
        "}"
        "pre,code{white-space:pre-wrap;overflow-wrap:anywhere}"
        "blockquote{border-left:3px solid "
        f"{theme['accent']}"
        ";background:"
        f"{theme['surface']}"
        ";margin-left:0;padding:3mm 4mm;color:"
        f"{theme['ink']}"
        "}"
        ".report-signature{margin-top:18mm;text-align:right;break-inside:avoid}"
        ".report-signature p{margin:0 0 2mm}"
    )
    pdf_document = (
        f"<html lang='zh-CN'><head><meta charset='utf-8'><title>{title}</title>"
        f"<style>{pdf_css}</style></head>"
        f"<body>{shared}</body></html>"
    )
    word_document = (
        "<meta charset='utf-8'><body>"
        f"<h1>{title}</h1><p>分析期间：{period}</p><p>{organization}</p>"
        f"<p>{generated_label}</p><p>{_WORD_MARKERS['cover_end']}</p>"
        f"<h1>目录</h1><p>{_WORD_MARKERS['toc_field_start']}</p>{toc}"
        f"<p>{_WORD_MARKERS['toc_field_end']}</p><p>{_WORD_MARKERS['toc_end']}</p>"
        f"<p>{_WORD_MARKERS['body_start']}</p>{body}"
        f"<p>{organization}</p><p>{generated_date}</p></body>"
    )
    return pdf_document, word_document


def _toc_page_numbers(
    pages: Sequence[Any], headings: Sequence[Mapping[str, Any]]
) -> dict[str, int]:
    """从 WeasyPrint 页面锚点计算正文从 1 开始的稳定目录页码。"""

    anchor_pages: dict[str, int] = {}
    expected_anchors = {str(item["anchor"]): str(item["anchor"]) for item in headings}
    for physical_page, page in enumerate(pages, start=1):
        anchors = getattr(page, "anchors", None)
        if not isinstance(anchors, Mapping):
            continue
        for anchor in anchors:
            heading_anchor = expected_anchors.get(anchor)
            if heading_anchor is None:
                continue
            if heading_anchor in anchor_pages:
                raise ReportFailure("PDF 正文标题锚点重复")
            anchor_pages[heading_anchor] = physical_page
    if set(anchor_pages) != set(expected_anchors):
        raise ReportFailure("PDF 正文标题锚点不完整")
    body_start_page = anchor_pages[str(headings[0]["anchor"])]
    page_numbers = {anchor: page - body_start_page + 1 for anchor, page in anchor_pages.items()}
    if any(number < 1 for number in page_numbers.values()):
        raise ReportFailure("PDF 正文章节页码顺序无效")
    return page_numbers


def _apply_pdf_page_decorations(
    path: Path,
    *,
    context: dict[str, Any],
    layout: dict[str, str],
) -> None:
    try:
        import pypdf
        from weasyprint import HTML
    except ImportError as error:
        raise ReportFailure("PDF 页面装饰运行时依赖不可用") from error

    overlay = path.with_name("render.decorations.pdf")
    decorated = path.with_name("render.decorated.pdf")
    try:
        reader = pypdf.PdfReader(str(path))
        if not reader.pages:
            raise ReportFailure("PDF 页面装饰处理缺少页面")
        section_pages = _pdf_section_pages(reader, context["sections"])
        body_start_page = min(section_pages.values())
        page_count = len(reader.pages)
        decoration_pages: list[str] = []
        for page_number in range(2, page_count + 1):
            page_value, pages_value = _page_number_context(
                page_number,
                body_start_page=body_start_page,
                physical_page_count=page_count,
            )
            values = {
                key: html.escape(
                    _formatted_page_text(
                        template,
                        title=context["title"],
                        organization=context["organizationName"],
                        page=page_value,
                        pages=pages_value,
                    )
                )
                for key, template in layout.items()
            }
            decoration_pages.append(
                '<section class="decoration-page">'
                '<header><span class="left">'
                f'{values["headerLeft"]}</span><span class="right">'
                f"{values['headerRight']}</span></header>"
                f'<div class="watermark">{html.escape(context["watermarkText"])}</div>'
                '<footer><span class="left">'
                f'{values["footerLeft"]}</span><span class="right">'
                f"{values['footerRight']}</span></footer></section>"
            )
        document = (
            "<meta charset='utf-8'><style>"
            "@page{size:A4;margin:0}html,body{margin:0}body{"
            "font-family:'Noto Sans CJK SC','Noto Sans CJK JP',sans-serif;color:"
            f"{REPORT_VISUAL_THEME['muted']}"
            "}"
            ".decoration-page{position:relative;box-sizing:border-box;width:210mm;height:297mm;"
            "break-after:page;overflow:hidden}.decoration-page:last-child{break-after:auto}"
            "header,footer{position:absolute;left:18mm;right:18mm;display:grid;"
            "grid-template-columns:minmax(0,1fr) minmax(0,1fr);column-gap:8mm;"
            "font-size:8pt;line-height:1.2}header{top:7mm}footer{bottom:8mm;"
            "border-top:.5pt solid "
            f"{REPORT_VISUAL_THEME['grid']}"
            ";padding-top:2mm}"
            ".left{text-align:left;white-space:pre-wrap;overflow-wrap:anywhere}"
            ".right{text-align:right;white-space:pre-wrap;overflow-wrap:anywhere}"
            ".watermark{position:absolute;left:25mm;right:25mm;top:122mm;text-align:center;"
            "transform:rotate(-32deg);font-size:26pt;line-height:1.15;font-weight:600;"
            "color:rgba(71,84,103,.045);overflow-wrap:anywhere}"
            f"</style><body>{''.join(decoration_pages)}</body>"
        )
        HTML(string=document).write_pdf(str(overlay), pdf_variant="pdf/ua-1")
        overlay_pages = pypdf.PdfReader(str(overlay)).pages
        if len(overlay_pages) != page_count - 1:
            raise ReportFailure("PDF 页面装饰页数不一致")
        writer = pypdf.PdfWriter(clone_from=str(path))
        # pypdf 默认将合并产物写成 1.3，即使源文档是 PDF/UA-1（1.7）；保留
        # PDF/UA 所需的版本声明，避免页面装饰步骤让结构化标签变成不一致的协议。
        writer.pdf_header = "%PDF-1.7"
        # WeasyPrint 的命名页沿用绝对页计数。服务端按章节锚点一次生成全部装饰页，
        # 目录使用罗马数字，正文从 1 重启；封面明确不合并任何页面元素。
        for page, overlay_page in zip(writer.pages[1:], overlay_pages, strict=True):
            page.merge_page(overlay_page, over=False)
        with decorated.open("wb") as stream:
            writer.write(stream)
        os.replace(decorated, path)
    except ReportFailure:
        raise
    except Exception as error:
        raise ReportFailure("PDF 页面装饰处理失败") from error
    finally:
        overlay.unlink(missing_ok=True)
        decorated.unlink(missing_ok=True)


def _render_docx(
    html_document: str,
    *,
    source_parent: Path,
    output: Path,
    context: dict[str, Any],
    layout: dict[str, str],
) -> dict[str, Any]:
    pandoc = shutil.which("pandoc")
    if pandoc is None:
        raise ReportFailure("Word 渲染命令 Pandoc 不可用")
    html_path = output.with_suffix(".html")
    html_path.write_text(html_document, encoding="utf-8")
    try:
        process = subprocess.run(
            [
                pandoc,
                "--from=html",
                "--to=docx",
                f"--resource-path={source_parent}",
                "--output",
                str(output),
                str(html_path),
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=DOCX_RENDER_TIMEOUT_SECONDS,
            check=False,
            cwd=source_parent,
        )
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ReportFailure("Word 渲染失败或超时") from error
    if process.returncode != 0 or not output.is_file():
        raise ReportFailure("Word 渲染失败")
    _postprocess_docx(output, context=context, layout=layout)
    return _validate_docx_structure(
        output,
        expected_sections=context["sections"],
        expected_headings=context["headingNumbers"],
        expected_image_count=None,
        watermark_text=context["watermarkText"],
    )


def _postprocess_docx(path: Path, *, context: dict[str, Any], layout: dict[str, str]) -> None:
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.oxml import OxmlElement, parse_xml
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
    except ImportError as error:
        raise ReportFailure("Word 后处理依赖 python-docx 不可用") from error

    document = Document(str(path))
    theme_colors = {
        name: RGBColor.from_string(str(value).removeprefix("#"))
        for name, value in REPORT_VISUAL_THEME.items()
        if isinstance(value, str) and value.startswith("#")
    }
    markers: dict[str, Any] = {}
    for paragraph in document.paragraphs:
        for name, marker in _WORD_MARKERS.items():
            if paragraph.text.strip() == marker:
                if name in markers:
                    raise ReportFailure("Word 版式标记重复")
                markers[name] = paragraph
    if set(markers) != set(_WORD_MARKERS):
        raise ReportFailure("Word 版式标记缺失")

    def clear_paragraph(paragraph: Any) -> None:
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)

    def field_run(
        paragraph: Any,
        instruction: str,
        *,
        result: str = "1",
        close: bool = True,
    ) -> None:
        begin_run = OxmlElement("w:r")
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        begin.set(qn("w:dirty"), "true")
        begin_run.append(begin)
        instruction_run = OxmlElement("w:r")
        instruction_text = OxmlElement("w:instrText")
        instruction_text.set(qn("xml:space"), "preserve")
        instruction_text.text = f" {instruction} "
        instruction_run.append(instruction_text)
        separate_run = OxmlElement("w:r")
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        separate_run.append(separate)
        result_run = OxmlElement("w:r")
        result_text = OxmlElement("w:t")
        result_text.text = result
        result_run.append(result_text)
        for item in (begin_run, instruction_run, separate_run, result_run):
            paragraph._p.append(item)
        if close:
            end_run = OxmlElement("w:r")
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            end_run.append(end)
            paragraph._p.append(end_run)

    def field_end(paragraph: Any) -> None:
        end_run = OxmlElement("w:r")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        end_run.append(end)
        paragraph._p.append(end_run)

    base_sect_pr = deepcopy(document._element.body.sectPr)

    def section_properties(*, page_format: str | None) -> Any:
        value = deepcopy(base_sect_pr)
        for child in list(value):
            if child.tag in {
                qn("w:headerReference"),
                qn("w:footerReference"),
                qn("w:type"),
                qn("w:pgNumType"),
                qn("w:titlePg"),
            }:
                value.remove(child)
        section_type = OxmlElement("w:type")
        section_type.set(qn("w:val"), "nextPage")
        value.insert(0, section_type)
        if page_format is not None:
            page_number = OxmlElement("w:pgNumType")
            page_number.set(qn("w:start"), "1")
            page_number.set(qn("w:fmt"), page_format)
            value.append(page_number)
        return value

    for marker_name, page_format in (("cover_end", None), ("toc_end", "lowerRoman")):
        paragraph = markers[marker_name]
        clear_paragraph(paragraph)
        paragraph._p.get_or_add_pPr().append(section_properties(page_format=page_format))
    final_sect_pr = document._element.body.sectPr
    for child in list(final_sect_pr):
        if child.tag in {qn("w:headerReference"), qn("w:footerReference"), qn("w:pgNumType")}:
            final_sect_pr.remove(child)
    body_page_number = OxmlElement("w:pgNumType")
    body_page_number.set(qn("w:start"), "1")
    body_page_number.set(qn("w:fmt"), "decimal")
    final_sect_pr.append(body_page_number)

    sections = document.sections
    if len(sections) != 3:
        raise ReportFailure("Word 必须包含封面、目录和正文三个分节")
    for section in sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(22)
        section.left_margin = Mm(18)
        section.right_margin = Mm(18)
        section.header_distance = Mm(8)
        section.footer_distance = Mm(8)

    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = (
            document.styles[style_name]
            if style_name in document.styles
            else document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        )
        style.font.name = "Noto Sans CJK SC"
        style.font.size = Pt(10.5 if style_name == "Normal" else 14)
        style.font.color.rgb = (
            theme_colors["ink"] if style_name == "Normal" else theme_colors["primary"]
        )
        style._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Noto Sans CJK SC"
        )
    if "Report Signature" not in document.styles:
        signature_style = document.styles.add_style("Report Signature", WD_STYLE_TYPE.PARAGRAPH)
        signature_style.font.name = "Noto Sans CJK SC"
        signature_style._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Noto Sans CJK SC"
        )

    paragraphs = document.paragraphs

    def marker_index(name: str) -> int:
        return next(
            index for index, paragraph in enumerate(paragraphs) if paragraph._p is markers[name]._p
        )

    cover_end_index = marker_index("cover_end")
    toc_start_index = marker_index("toc_field_start")
    toc_end_index = marker_index("toc_field_end")
    body_start_index = marker_index("body_start")
    for paragraph in paragraphs[:cover_end_index]:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph = next(
        (item for item in paragraphs[:cover_end_index] if item.text == context["title"]), None
    )
    if title_paragraph is None:
        raise ReportFailure("Word 封面缺少报告标题")
    title_paragraph.style = document.styles["Title"]
    toc_title = next(
        (item for item in paragraphs[cover_end_index:toc_start_index] if item.text == "目录"), None
    )
    if toc_title is None:
        raise ReportFailure("Word 缺少目录标题")
    toc_title.style = document.styles["Title"]

    body_headings: list[Any] = []
    search_index = body_start_index + 1
    for heading_index, item in enumerate(context["headingNumbers"], start=1):
        expected_text = f"{item['number']} {item['title']}"
        found = next(
            (
                (index, paragraph)
                for index, paragraph in enumerate(paragraphs[search_index:], start=search_index)
                if paragraph.text.strip() == expected_text
            ),
            None,
        )
        if found is None:
            raise ReportFailure("Word 正文标题与服务端编号映射不一致")
        search_index, heading = found
        search_index += 1
        heading.style = document.styles[f"Heading {item['level'] - 1}"]
        bookmark_name = item["anchor"].replace("-", "_")
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(1000 + heading_index))
        start.set(qn("w:name"), bookmark_name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(1000 + heading_index))
        insert_at = 1 if heading._p.pPr is not None else 0
        heading._p.insert(insert_at, start)
        heading._p.append(end)
        body_headings.append(heading)

    toc_entries = [
        paragraph
        for paragraph in paragraphs[toc_start_index + 1 : toc_end_index]
        if paragraph.text.strip()
    ]
    if len(toc_entries) != len(context["headingNumbers"]):
        raise ReportFailure("Word 缓存目录与正式标题不一致")
    for paragraph, item in zip(toc_entries, context["headingNumbers"], strict=True):
        clear_paragraph(paragraph)
        toc_style = f"TOC {item['level'] - 1}"
        if toc_style in document.styles:
            paragraph.style = document.styles[toc_style]
        usable_width = sections[1].page_width - sections[1].left_margin - sections[1].right_margin
        paragraph.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
        hyperlink = OxmlElement("w:hyperlink")
        bookmark_name = item["anchor"].replace("-", "_")
        hyperlink.set(qn("w:anchor"), bookmark_name)
        hyperlink.set(qn("w:history"), "1")
        run = OxmlElement("w:r")
        run_properties = OxmlElement("w:rPr")
        run_style = OxmlElement("w:rStyle")
        run_style.set(qn("w:val"), "Hyperlink")
        run_properties.append(run_style)
        run.append(run_properties)
        text = OxmlElement("w:t")
        text.text = f"{item['number']} {item['title']}"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        paragraph.add_run("\t")
        field_run(
            paragraph,
            f"PAGEREF {bookmark_name} \\h",
        )

    clear_paragraph(markers["toc_field_start"])
    field_run(
        markers["toc_field_start"],
        'TOC \\o "1-3" \\h \\z \\u',
        result="",
        close=False,
    )
    clear_paragraph(markers["toc_field_end"])
    field_end(markers["toc_field_end"])
    clear_paragraph(markers["body_start"])

    def clear_story(story: Any) -> Any:
        for table in list(story.tables):
            story._element.remove(table._element)
        paragraph = story.paragraphs[0]
        clear_paragraph(paragraph)
        for extra in list(story.paragraphs[1:]):
            story._element.remove(extra._element)
        return paragraph

    def add_template(paragraph: Any, left: str, right: str) -> None:
        usable_width = sections[1].page_width - sections[1].left_margin - sections[1].right_margin
        paragraph.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)

        def append(value: str) -> None:
            for literal, field_name, _format_spec, _conversion in Formatter().parse(value):
                if literal:
                    paragraph.add_run(literal)
                if field_name == "title":
                    paragraph.add_run(context["title"])
                elif field_name == "organization":
                    paragraph.add_run(context["organizationName"])
                elif field_name in _WORD_PAGE_FIELDS:
                    field_run(paragraph, _WORD_PAGE_FIELDS[field_name])

        append(left)
        paragraph.add_run("\t")
        append(right)
        for run in paragraph.runs:
            run.font.name = "Noto Sans CJK SC"
            run.font.size = Pt(8)
            run.font.color.rgb = theme_colors["muted"]
            run._element.get_or_add_rPr().get_or_add_rFonts().set(
                qn("w:eastAsia"), "Noto Sans CJK SC"
            )

    def add_watermark(header: Any, text: str, shape_id: int) -> None:
        escaped = html.escape(text, quote=True)
        paragraph = header.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        watermark = parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>'
            f'<v:shape id="PowerPlusWaterMarkObject{shape_id}" '
            'type="#_x0000_t136" '
            'style="position:absolute;width:430pt;height:90pt;rotation:315;z-index:-251654144;'
            'mso-position-horizontal:center;mso-position-vertical:center" '
            f'fillcolor="{REPORT_VISUAL_THEME["muted"]}" stroked="f">'
            '<v:fill opacity="0.045"/>'
            '<v:textpath style="font-family:Noto Sans CJK SC;font-size:26pt" '
            f'string="{escaped}"/></v:shape></w:pict></w:r>'
        )
        paragraph._p.append(watermark)

    for index, section in enumerate(sections):
        if index == 0:
            continue
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        header_paragraph = clear_story(section.header)
        footer_paragraph = clear_story(section.footer)
        add_template(
            header_paragraph,
            layout["headerLeft"],
            layout["headerRight"],
        )
        add_template(
            footer_paragraph,
            layout["footerLeft"],
            layout["footerRight"],
        )
        add_watermark(section.header, context["watermarkText"], index)

    for table in document.tables:
        if "Table Grid" in document.styles:
            table.style = "Table Grid"
        table_properties = table._tbl.tblPr
        for existing in table_properties.findall(qn("w:tblBorders")):
            table_properties.remove(existing)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), str(REPORT_VISUAL_THEME["grid"]).removeprefix("#"))
            borders.append(border)
        table_properties.append(borders)
        if table.rows:
            row_properties = table.rows[0]._tr.get_or_add_trPr()
            table_header = OxmlElement("w:tblHeader")
            table_header.set(qn("w:val"), "true")
            row_properties.append(table_header)
            for cell in table.rows[0].cells:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), str(REPORT_VISUAL_THEME["surface"]).removeprefix("#"))
                cell._tc.get_or_add_tcPr().append(shading)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = theme_colors["primary"]
    available_width = sections[-1].page_width - sections[-1].left_margin - sections[-1].right_margin
    for shape in document.inline_shapes:
        if shape.width > available_width:
            ratio = available_width / shape.width
            shape.width = int(shape.width * ratio)
            shape.height = int(shape.height * ratio)

    for paragraph in paragraphs[body_start_index + 1 :]:
        if paragraph.text.strip() in {context["organizationName"], context["generatedDate"]}:
            paragraph.style = document.styles["Report Signature"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    settings = document.settings._element
    for existing in settings.findall(qn("w:updateFields")):
        settings.remove(existing)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    document.core_properties.title = context["title"]
    document.core_properties.subject = context["periodLabel"]
    document.core_properties.author = context["organizationName"]
    document.core_properties.last_modified_by = context["generatedByLabel"]
    document.core_properties.keywords = context["generatedByLabel"]
    document.core_properties.modified = datetime.now(UTC)

    postprocessed = path.with_name("render.postprocessed.docx")
    document.save(str(postprocessed))
    if postprocessed.stat().st_size > MAX_DOCX_BYTES:
        raise ReportFailure("Word 文件不能超过 200 MiB")
    os.replace(postprocessed, path)


def _validate_docx_structure(
    path: Path,
    *,
    expected_sections: list[dict[str, str]],
    expected_headings: list[dict[str, Any]],
    expected_image_count: int | None,
    watermark_text: str,
) -> dict[str, Any]:
    if not path.is_file() or not 1 <= path.stat().st_size <= MAX_DOCX_BYTES:
        raise ReportFailure("Word 文件不存在或超过 200 MiB")
    try:
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            if any(
                name == forbidden or name.startswith(forbidden)
                for name in names
                for forbidden in _DOCX_FORBIDDEN_PARTS
            ):
                raise ReportFailure("Word 包含宏、OLE 或 ActiveX 内容")
            for relationship_name in (name for name in names if name.endswith(".rels")):
                relationships = ElementTree.fromstring(package.read(relationship_name))
                if any(item.attrib.get("TargetMode") == "External" for item in relationships):
                    raise ReportFailure("Word 包含外部关系")
            required = {"word/document.xml", "word/settings.xml", "[Content_Types].xml"}
            if not required.issubset(names):
                raise ReportFailure("Word OOXML 结构不完整")
            document_content = package.read("word/document.xml")
            document_root = ElementTree.fromstring(document_content)
            document_xml = document_content.decode("utf-8")
            image_names = [name for name in names if name.startswith("word/media/")]
    except (OSError, KeyError, UnicodeError, zipfile.BadZipFile, ElementTree.ParseError) as error:
        raise ReportFailure("Word OOXML 无法解析") from error
    bookmark_names = {item["anchor"].replace("-", "_") for item in expected_headings}
    native_toc_present = (
        'TOC \\o "1-3"' in document_xml
        and "PAGEREF report_" in document_xml
        and all(name in document_xml for name in bookmark_names)
    )
    toc_entry_count = sum(name in document_xml for name in bookmark_names)
    word_namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    section_properties = document_root.findall(f".//{word_namespace}sectPr")
    if expected_image_count is not None and len(image_names) < expected_image_count:
        raise ReportFailure("Word 未完整嵌入图表")
    return {
        "nativeTocPresent": native_toc_present,
        "sectionCount": len(section_properties),
        "tocEntryCount": toc_entry_count,
        "embeddedImageCount": len(image_names),
        "externalRelationshipCount": 0,
    }


def _roman(value: int) -> str:
    if not 1 <= value <= 3_999:
        raise ReportFailure("罗马页码超出支持范围")
    result: list[str] = []
    remaining = value
    for number, numeral in (
        (1000, "m"),
        (900, "cm"),
        (500, "d"),
        (400, "cd"),
        (100, "c"),
        (90, "xc"),
        (50, "l"),
        (40, "xl"),
        (10, "x"),
        (9, "ix"),
        (5, "v"),
        (4, "iv"),
        (1, "i"),
    ):
        count, remaining = divmod(remaining, number)
        result.extend(numeral for _index in range(count))
    return "".join(result)


def _page_number_context(
    physical_page: int,
    *,
    body_start_page: int,
    physical_page_count: int,
) -> tuple[str | int, str | int]:
    """返回当前分节内的页码和分节总页数，封面不进入编号体系。"""

    if (
        not 2 <= physical_page <= physical_page_count
        or not 2 <= body_start_page <= physical_page_count
    ):
        raise ReportFailure("报表分节页码边界无效")
    if physical_page < body_start_page:
        return _roman(physical_page - 1), _roman(body_start_page - 2)
    return physical_page - body_start_page + 1, physical_page_count - body_start_page + 1


def _pdf_section_pages(reader: Any, sections: list[dict[str, str]]) -> dict[str, int]:
    destinations = getattr(reader, "named_destinations", {})
    pages: dict[str, int] = {}
    for section in sections:
        name = f"report-section-{section['code']}"
        destination = destinations.get(name)
        if destination is None:
            raise ReportFailure("PDF 缺少稳定章节锚点")
        try:
            pages[section["code"]] = int(reader.get_destination_page_number(destination)) + 1
        except Exception as error:
            raise ReportFailure("PDF 章节锚点无法解析") from error
    if list(pages) != [item["code"] for item in sections] or list(pages.values()) != sorted(
        pages.values()
    ):
        raise ReportFailure("PDF 章节锚点顺序与已批准提纲不一致")
    return pages


def _pdf_heading_pages(reader: Any, headings: list[dict[str, Any]]) -> dict[str, int]:
    destinations = getattr(reader, "named_destinations", {})
    pages: dict[str, int] = {}
    for item in headings:
        destination = destinations.get(item["anchor"])
        if destination is None:
            raise ReportFailure("PDF 缺少稳定标题锚点")
        try:
            pages[item["anchor"]] = int(reader.get_destination_page_number(destination)) + 1
        except Exception as error:
            raise ReportFailure("PDF 标题锚点无法解析") from error
    if list(pages.values()) != sorted(pages.values()):
        raise ReportFailure("PDF 标题锚点顺序与编号映射不一致")
    return pages


def _pdf_link_count(reader: Any, *, start_page: int, end_page: int) -> int:
    count = 0
    for page in reader.pages[start_page - 1 : end_page]:
        annotations = page.get("/Annots") or ()
        for reference in annotations:
            try:
                annotation = reference.get_object()
            except Exception:
                continue
            if annotation.get("/Subtype") == "/Link":
                count += 1
    return count


def _validate_docx_rendering(
    path: Path,
    directory: Path,
    *,
    context: dict[str, Any],
    layout: dict[str, str],
) -> dict[str, Any]:
    try:
        import pypdf
        from PIL import Image
    except ImportError as error:
        raise ReportFailure("Word 可渲染性验收依赖不可用") from error
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if libreoffice is None:
        raise ReportFailure("Word 验收命令 LibreOffice 不可用")
    directory.mkdir(mode=0o700)
    output = directory / "output"
    profile = directory / "profile"
    raster = directory / "raster"
    output.mkdir()
    profile.mkdir()
    raster.mkdir()
    try:
        process = subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation={profile.resolve().as_uri()}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output),
                str(path),
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=DOCX_VALIDATION_TIMEOUT_SECONDS,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ReportFailure("Word 经 LibreOffice 转换失败或超时") from error
    converted = output / f"{path.stem}.pdf"
    if process.returncode != 0 or not converted.is_file():
        raise ReportFailure("Word 经 LibreOffice 转换失败")
    try:
        reader = pypdf.PdfReader(str(converted))
    except pypdf.errors.PdfReadError as error:
        raise ReportFailure("LibreOffice 转换结果无法打开") from error
    if not 1 <= len(reader.pages) <= MAX_PDF_PAGES:
        raise ReportFailure("Word 转换页数超出边界")
    prefix = raster / "page"
    try:
        rendered = subprocess.run(
            ["pdftoppm", "-gray", "-r", "72", "-png", str(converted), str(prefix)],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=DOCX_VALIDATION_TIMEOUT_SECONDS,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ReportFailure("Word 转换结果栅格化失败") from error
    rendered_pages = sorted(
        raster.glob("page-*.png"), key=lambda item: int(item.stem.rsplit("-", 1)[-1])
    )
    if rendered.returncode != 0 or len(rendered_pages) != len(reader.pages):
        raise ReportFailure("Word 转换结果栅格化失败")
    extracted_pages: list[str] = []
    page_image_counts: list[int] = []
    image_count = 0
    for page, rendered_page in zip(reader.pages, rendered_pages, strict=True):
        page_text = page.extract_text() or ""
        extracted_pages.append(page_text)
        current_image_count = len(page.images)
        page_image_counts.append(current_image_count)
        image_count += current_image_count
        with Image.open(rendered_page) as image:
            image.verify()
    first_section_title = "".join(context["sections"][0]["title"].split())
    report_title = "".join(context["title"].split())
    first_section_pages = [
        index
        for index, page_text in enumerate(extracted_pages, start=1)
        if index > 1
        and first_section_title in "".join(page_text.split()).replace(report_title, "", 1)
    ]
    # LibreOffice 会在转换后的 PDF 文本层为中文标题插入布局空格。该位置只用于
    # 识别页码装饰，不再把目录缓存方式和正文分页形态作为发布门禁。
    body_start_page = first_section_pages[-1] if first_section_pages else 2
    blank_pages: list[int] = []
    for index, (page_text, current_image_count) in enumerate(
        zip(extracted_pages, page_image_counts, strict=True), start=1
    ):
        substantive_text = "".join(page_text.split())
        if index > 1:
            page_value, pages_value = _page_number_context(
                index,
                body_start_page=body_start_page,
                physical_page_count=len(reader.pages),
            )
            decorations = [
                context["watermarkText"],
                *(
                    _formatted_page_text(
                        value,
                        title=context["title"],
                        organization=context["organizationName"],
                        page=page_value,
                        pages=pages_value,
                    )
                    for value in layout.values()
                    if value
                ),
            ]
            # Writer 转换后的页眉、页脚和 VML 水印也会进入 PDF 文本层；只含这些
            # 服务端装饰的页面仍是空白正文页，不能借装饰绕过发布门禁。
            for decoration in decorations:
                substantive_text = substantive_text.replace("".join(decoration.split()), "", 1)
        if not substantive_text and current_image_count == 0:
            blank_pages.append(index)
    extracted_text = "\n".join(extracted_pages)
    compact_extracted_text = "".join(extracted_text.split())
    required_text = (
        context["title"],
        context["periodLabel"],
        context["organizationName"],
        context["generatedByLabel"],
        context["generatedDate"],
        *(item["title"] for item in context["sections"]),
    )
    if blank_pages or any(
        "".join(item.split()) not in compact_extracted_text for item in required_text
    ):
        raise ReportFailure("Word 转换结果缺少正式内容或包含空白页")
    if any(marker in extracted_text for marker in _WORD_MARKERS.values()):
        raise ReportFailure("Word 显示了内部版式标记")
    return {
        "convertedPageCount": len(reader.pages),
        "blankPages": blank_pages,
        "renderedImageCount": image_count,
    }


class ReportRuntime:
    def __init__(self, workspace: str | Path):
        self.workspace = Path(workspace).resolve()

    def _validate_datasets(self, state: dict[str, Any]) -> None:
        sources = state.get("sources")
        if not isinstance(sources, list) or not 1 <= len(sources) <= MAX_DATASET_PATHS:
            raise ReportFailure("分析任务缺少源文件校验信息")
        for source in sources:
            if not isinstance(source, dict):
                raise ReportFailure("分析任务缺少源文件校验信息")
            value = source.get("path")
            if not isinstance(value, str):
                raise ReportFailure("分析任务缺少源文件校验信息")
            relative = _relative_path(value)
            path = self.workspace.joinpath(*relative.parts)
            _reject_symlinks(self.workspace, path)
            if (
                not path.is_file()
                or path.stat().st_size != source.get("size")
                or _sha256(path) != source.get("sha256")
            ):
                raise ReportFailure("源文件发生变化，分析任务已失效")

    def _artifact(self, path: Path) -> dict[str, Any]:
        _reject_symlinks(self.workspace, path)
        if not path.is_file():
            raise ReportFailure("报表产物不存在或不是普通文件")
        return {
            "path": str(path.relative_to(self.workspace)),
            "size": path.stat().st_size,
            "sha256": _sha256(path),
        }

    def _images(self, markdown_path: Path, tokens: list[Any]) -> set[Path]:
        images: list[Path] = []
        for token in tokens:
            for child in token.children or []:
                if child.type != "image":
                    continue
                source = child.attrGet("src") or ""
                parsed = urlsplit(source)
                if parsed.scheme or parsed.netloc or parsed.query or parsed.fragment:
                    raise ReportFailure("Markdown 图片只能引用工作区内的相对路径")
                decoded = unquote(parsed.path)
                if not decoded or "\\" in decoded:
                    raise ReportFailure("Markdown 图片路径无效")
                relative = PurePosixPath(decoded)
                if relative.is_absolute() or ".." in relative.parts:
                    raise ReportFailure("Markdown 图片只能引用工作区内的相对路径")
                image = markdown_path.parent.joinpath(*relative.parts)
                _reject_symlinks(self.workspace, image)
                try:
                    image.relative_to(self.workspace)
                except ValueError as error:
                    raise ReportFailure("Markdown 图片路径越界") from error
                if image.suffix.lower() not in IMAGE_SUFFIXES or not image.is_file():
                    raise ReportFailure("Markdown 图片不存在或格式不受支持")
                if image.stat().st_size > MAX_IMAGE_BYTES:
                    raise ReportFailure("单张 Markdown 图片超过 10 MiB")
                _check_image_signature(image)
                images.append(image.resolve())
        unique = set(images)
        if sum(path.stat().st_size for path in unique) > MAX_TOTAL_IMAGE_BYTES:
            raise ReportFailure("Markdown 图片合计超过 50 MiB")
        return unique

    def render_markdown(
        self,
        state: dict[str, Any],
        markdown_path: str,
        output_path: str,
        temporary_path: str,
        page_layout: dict[str, str] | None = None,
        word_output_path: str | None = None,
    ) -> dict[str, Any]:
        try:
            import pypdf
            from markdown_it import MarkdownIt
            from weasyprint import HTML, URLFetcher
        except ImportError as error:
            raise ReportFailure("PDF 运行时依赖不可用") from error

        self._validate_datasets(state)
        temporary: Path | None = None
        succeeded = False
        try:
            source = _input_path(self.workspace, markdown_path, ".md")
            if source.stat().st_size > MAX_MARKDOWN_BYTES:
                raise ReportFailure("Markdown 文件超过 1 MiB")
            try:
                markdown = source.read_text(encoding="utf-8")
            except UnicodeDecodeError as error:
                raise ReportFailure("Markdown 文件必须使用 UTF-8 编码") from error

            pdf_markdown, citation_presentations = _pdf_markdown(
                markdown, state.get("_citationPresentations")
            )
            pdf_markdown = _normalize_cjk_strong_markers(pdf_markdown)
            parser = MarkdownIt("commonmark", {"html": False}).enable("table")
            tokens = parser.parse(pdf_markdown)
            layout = _page_layout(page_layout)
            context = _document_context(state.get("_documentContext"))
            title = _markdown_title(tokens)
            if title != context["title"]:
                raise ReportFailure("Markdown 标题与服务端文档展示契约不一致")
            marker_sections = _SECTION_MARKER.findall(markdown)
            expected_section_codes = [item["code"] for item in context["sections"]]
            if marker_sections != expected_section_codes:
                raise ReportFailure("Markdown 正式章节标识与已批准提纲不一致")
            _bind_heading_anchors(tokens, context["headingNumbers"])
            allowed_images = self._images(source, tokens)
            source_artifact = self._artifact(source)
            image_artifacts = [self._artifact(path) for path in sorted(allowed_images)]
            body = parser.renderer.render(_body_tokens(tokens), parser.options, {})
            output = _output_path(self.workspace, output_path)
            word_output = _word_output_path(
                self.workspace,
                word_output_path or str(PurePosixPath(output_path).with_suffix(".docx")),
            )
            if output.parent != word_output.parent:
                raise ReportFailure("PDF 和 Word 必须发布到同一 revision 目录")
            temporary = _temporary_pdf_path(temporary_path)
            temporary_docx = _temporary_docx_path(temporary)
            file_fetcher = URLFetcher(allowed_protocols={"file"}, fail_on_errors=True)

            def fetch_resource(url: str) -> dict[str, Any]:
                parsed = urlsplit(url)
                if parsed.scheme != "file" or parsed.netloc not in ("", "localhost"):
                    raise ReportFailure("PDF 渲染禁止访问外部资源")
                path = Path(unquote(parsed.path)).resolve()
                if path not in allowed_images:
                    raise ReportFailure("PDF 渲染引用了未校验资源")
                return file_fetcher(url)

            pdf_document, word_document = _semantic_documents(
                body,
                context=context,
                layout=layout,
            )
            preflight_document = HTML(
                string=pdf_document,
                base_url=str(source.parent),
                url_fetcher=fetch_resource,
            ).render()
            toc_page_numbers = _toc_page_numbers(
                preflight_document.pages, context["headingNumbers"]
            )
            pdf_document, _ = _semantic_documents(
                body,
                context=context,
                layout=layout,
                toc_page_numbers=toc_page_numbers,
            )
            final_document = HTML(
                string=pdf_document,
                base_url=str(source.parent),
                url_fetcher=fetch_resource,
            ).render()
            if (
                _toc_page_numbers(final_document.pages, context["headingNumbers"])
                != toc_page_numbers
            ):
                # 目录页码使用固定宽度，正常不会改变分页；若字体或渲染器升级导致
                # 锚点漂移，则不能发布目录与正文不一致的产物。
                raise ReportFailure("PDF 目录页码在最终渲染时发生漂移")
            final_document.write_pdf(str(temporary), pdf_variant="pdf/ua-1")
            if temporary.stat().st_size > MAX_PDF_BYTES:
                raise ReportFailure("PDF 文件不能超过 200 MiB")
            base_page_count = len(final_document.pages)
            if not base_page_count:
                raise ReportFailure("PDF 校验失败")
            if base_page_count > MAX_PDF_PAGES:
                raise ReportFailure("PDF 页数不能超过 200 页")
            _apply_pdf_page_decorations(
                temporary,
                context=context,
                layout=layout,
            )
            word_structure = _render_docx(
                word_document,
                source_parent=source.parent,
                output=temporary_docx,
                context=context,
                layout=layout,
            )
            pdf_size = temporary.stat().st_size
            if pdf_size > MAX_PDF_BYTES:
                raise ReportFailure("PDF 文件不能超过 200 MiB")
            reader = pypdf.PdfReader(str(temporary))
            page_count = len(reader.pages)
            if not page_count:
                raise ReportFailure("PDF 校验失败")
            if page_count > MAX_PDF_PAGES:
                raise ReportFailure("PDF 页数不能超过 200 页")
            docx_size = temporary_docx.stat().st_size
            if docx_size > MAX_DOCX_BYTES:
                raise ReportFailure("Word 文件不能超过 200 MiB")
            word_structure = _validate_docx_structure(
                temporary_docx,
                expected_sections=context["sections"],
                expected_headings=context["headingNumbers"],
                expected_image_count=len(allowed_images),
                watermark_text=context["watermarkText"],
            )
            if self._artifact(source)["sha256"] != source_artifact["sha256"] or any(
                self._artifact(path)["sha256"] != artifact["sha256"]
                for path, artifact in zip(sorted(allowed_images), image_artifacts, strict=True)
            ):
                raise ReportFailure("Markdown 或图片在渲染期间发生变化，请重新生成报表")
            pdf_artifact = {
                "path": str(output.relative_to(self.workspace)),
                "size": pdf_size,
                "sha256": _sha256(temporary),
            }
            word_artifact = {
                "path": str(word_output.relative_to(self.workspace)),
                "size": docx_size,
                "sha256": _sha256(temporary_docx),
            }
            render = {
                "markdown": source_artifact,
                "pdf": pdf_artifact,
                "word": word_artifact,
                "images": image_artifacts,
                "pageCount": page_count,
                "imageCount": len(allowed_images),
                "pageLayout": layout,
                "reportTitle": title,
                "documentContext": context,
                "visualTheme": deepcopy(REPORT_VISUAL_THEME),
                "wordStructure": word_structure,
                "citationPresentations": citation_presentations,
                "citationAppendixPresent": False,
            }
            result = {
                "status": "rendered",
                "jobId": state["jobId"],
                "markdownPath": str(source.relative_to(self.workspace)),
                "pdfPath": str(output.relative_to(self.workspace)),
                "wordPath": str(word_output.relative_to(self.workspace)),
                "pageCount": page_count,
                "imageCount": len(allowed_images),
                "size": pdf_size,
                "wordSize": docx_size,
                "render": render,
            }
            succeeded = True
            return result
        finally:
            if temporary is not None and not succeeded:
                shutil.rmtree(temporary.parent, ignore_errors=True)

    def validate_pdf(
        self,
        state: dict[str, Any],
        pdf_path: str,
        temporary_directory: str,
        artifact_manifest: dict[str, Any] | None = None,
        word_path: str | None = None,
    ) -> dict[str, Any]:
        try:
            import pypdf
            from PIL import Image
        except ImportError as error:
            raise ReportFailure("PDF 视觉验收依赖不可用") from error
        if not shutil.which("pdftoppm"):
            raise ReportFailure("PDF 视觉验收命令不可用")

        self._validate_datasets(state)
        temp_path: Path | None = None
        try:
            render = state.get("render")
            if not isinstance(render, dict) or render.get("pdf", {}).get("path") != pdf_path:
                raise ReportFailure("PDF 未登记为当前分析任务的渲染产物")
            registered_word = render.get("word")
            current_word_path = word_path or (
                registered_word.get("path") if isinstance(registered_word, dict) else None
            )
            if (
                not isinstance(registered_word, dict)
                or not isinstance(current_word_path, str)
                or registered_word.get("path") != current_word_path
            ):
                raise ReportFailure("Word 未登记为当前分析任务的渲染产物")
            supporting_artifacts = [render["markdown"], *render.get("images", [])]
            for artifact in supporting_artifacts:
                supporting = self.workspace.joinpath(*_relative_path(artifact["path"]).parts)
                if self._artifact(supporting)["sha256"] != artifact["sha256"]:
                    raise ReportFailure("Markdown 或图片产物发生变化，请重新渲染后验收")
            relative = _relative_path(pdf_path, ".pdf")
            path = self.workspace.joinpath(*relative.parts)
            self._check_pdf_bounds(path)
            current = self._artifact(path)
            if current["sha256"] != render["pdf"]["sha256"]:
                raise ReportFailure("PDF 产物发生变化，请重新渲染后验收")
            word_relative = _relative_path(current_word_path, ".docx")
            word = self.workspace.joinpath(*word_relative.parts)
            word_current = self._artifact(word)
            if word_current["sha256"] != registered_word.get("sha256"):
                raise ReportFailure("Word 产物发生变化，请重新渲染后验收")
            pages: list[dict[str, Any]] = []
            blank_pages: list[int] = []
            missing_page_layout: list[int] = []
            rendered_image_count = 0
            temp_path = _validation_directory(temporary_directory)
            with _cleanup_directory(temp_path):
                prefix = temp_path / "page"
                try:
                    process = subprocess.run(
                        [
                            "pdftoppm",
                            "-gray",
                            "-r",
                            "72",
                            "-png",
                            str(path),
                            str(prefix),
                        ],
                        stdout=subprocess.DEVNULL,
                        stderr=subprocess.DEVNULL,
                        timeout=PDF_VALIDATION_TIMEOUT_SECONDS,
                        check=False,
                    )
                    reader = pypdf.PdfReader(str(path))
                except (OSError, subprocess.TimeoutExpired, pypdf.errors.PdfReadError) as error:
                    raise ReportFailure("PDF 视觉验收无法打开产物") from error
                rendered_pages = sorted(
                    temp_path.glob("page-*.png"),
                    key=lambda item: int(item.stem.rsplit("-", 1)[-1]),
                )
                if process.returncode != 0 or len(rendered_pages) != len(reader.pages):
                    raise ReportFailure("PDF 视觉验收栅格化失败")
                extracted_pages: list[str] = []
                layout = _page_layout(render.get("pageLayout"))
                title = str(render.get("reportTitle") or "智能运营报表")
                context = _document_context(render.get("documentContext"))
                section_pages = _pdf_section_pages(reader, context["sections"])
                _pdf_heading_pages(reader, context["headingNumbers"])
                body_start_page = min(section_pages.values())
                toc_link_count = (
                    _pdf_link_count(
                        reader,
                        start_page=2,
                        end_page=body_start_page - 1,
                    )
                    if body_start_page > 2
                    else 0
                )
                for index, (page, rendered_page) in enumerate(
                    zip(reader.pages, rendered_pages, strict=True), start=1
                ):
                    page_text = page.extract_text() or ""
                    extracted_pages.append(page_text)
                    with Image.open(rendered_page) as image:
                        grayscale = image.convert("L")
                        samples = grayscale.tobytes()
                        width, height = grayscale.size
                    non_white = sum(value < 250 for value in samples)
                    ratio = round(non_white / len(samples), 6) if samples else 0.0
                    image_count = len(page.images)
                    rendered_image_count += image_count
                    if index == 1:
                        role = "cover"
                        page_value: str | int = 1
                        pages_value: str | int = 1
                        layout_present = False
                    else:
                        role = "toc" if index < body_start_page else "body"
                        page_value, pages_value = _page_number_context(
                            index,
                            body_start_page=body_start_page,
                            physical_page_count=len(reader.pages),
                        )
                        layout_present = _has_page_layout(
                            page_text,
                            layout,
                            title=title,
                            organization=context["organizationName"],
                            page=page_value,
                            pages=pages_value,
                        )
                    watermark_present = index > 1 and context["watermarkText"] in page_text
                    substantive_text = "".join(page_text.split())
                    if index > 1:
                        decorations = [
                            context["watermarkText"],
                            *(
                                _formatted_page_text(
                                    value,
                                    title=title,
                                    organization=context["organizationName"],
                                    page=page_value,
                                    pages=pages_value,
                                )
                                for value in layout.values()
                                if value
                            ),
                        ]
                        # 页眉、页脚和水印即使存在，也不能把无正文页面伪装成非空白页。
                        for decoration in decorations:
                            substantive_text = substantive_text.replace(
                                "".join(decoration.split()), "", 1
                            )
                    text_char_count = len(substantive_text)
                    if index > 1 and not layout_present:
                        missing_page_layout.append(index)
                    blank = (
                        text_char_count == 0 and image_count == 0 and (index > 1 or ratio < 0.0005)
                    )
                    if blank:
                        blank_pages.append(index)
                    pages.append(
                        {
                            "page": index,
                            "width": width,
                            "height": height,
                            "nonWhiteRatio": ratio,
                            "textCharCount": text_char_count,
                            "imageCount": image_count,
                            "pageLayoutPresent": layout_present,
                            "pageLayoutExpected": index > 1,
                            "watermarkPresent": watermark_present,
                            "role": role,
                            "blank": blank,
                        }
                    )
                extracted_text = "\n".join(extracted_pages)
                cover_text = extracted_pages[0]
                toc_text = "\n".join(extracted_pages[1 : body_start_page - 1])
                final_text = extracted_pages[-1]
                cover_values = (
                    context["title"],
                    context["periodLabel"],
                    context["organizationName"],
                    context["generatedByLabel"],
                )
                cover_compact = "".join(cover_text.split())
                watermark_compact = "".join(context["watermarkText"].split())
                expected_cover_occurrences = sum(
                    "".join(value.split()).count(watermark_compact) for value in cover_values
                )
                # generatedByLabel 与默认水印文字相同，不能用简单的“不包含水印文字”判断封面。
                # 这里只允许封面固定事实本身贡献的出现次数，额外重复即视为水印泄漏。
                cover_ok = (
                    all("".join(value.split()) in cover_compact for value in cover_values)
                    and cover_compact.count(watermark_compact) == expected_cover_occurrences
                )
                toc_ok = "目录" in toc_text and all(
                    f"{item['number']} {item['title']}" in toc_text
                    for item in context["headingNumbers"]
                )
                signature_ok = (
                    context["organizationName"] in final_text
                    and context["generatedDate"] in final_text
                )
                watermark_pages = [
                    item["page"] for item in pages if item["page"] > 1 and item["watermarkPresent"]
                ]
                word_structure = _validate_docx_structure(
                    word,
                    expected_sections=context["sections"],
                    expected_headings=context["headingNumbers"],
                    expected_image_count=int(render.get("imageCount") or 0),
                    watermark_text=context["watermarkText"],
                )
                word_rendering = _validate_docx_rendering(
                    word,
                    temp_path / "word-validation",
                    context=context,
                    layout=layout,
                )
            markdown_image_count = int(render.get("imageCount") or 0)
            missing_images = max(0, markdown_image_count - rendered_image_count)
            chart_ids, citation_ids, section_ids = self._validate_manifest_markers(
                artifact_manifest,
                render=render,
                extracted_text=extracted_text,
            )
            ok = (
                bool(pages)
                and not blank_pages
                and missing_images == 0
                and not word_rendering["blankPages"]
                and word_structure["embeddedImageCount"] >= markdown_image_count
                and cover_ok
                and toc_ok
                and toc_link_count >= len(context["headingNumbers"])
                and signature_ok
                and word_structure["nativeTocPresent"]
                and word_structure["tocEntryCount"] == len(context["headingNumbers"])
            )
            validation = {
                "ok": ok,
                "status": "validated" if ok else "validation_failed",
                "pdfPath": pdf_path,
                "pdfSha256": current["sha256"],
                "wordPath": current_word_path,
                "wordSha256": word_current["sha256"],
                "pageCount": len(pages),
                "markdownImageCount": markdown_image_count,
                "renderedImageCount": rendered_image_count,
                "missingImageCount": missing_images,
                "chartIds": chart_ids,
                "citationIds": citation_ids,
                "sectionIds": section_ids,
                "blankPages": blank_pages,
                "missingPageLayoutPages": missing_page_layout,
                "coverPresent": cover_ok,
                "tocPresent": toc_ok,
                "tocLinkCount": toc_link_count,
                "bodyStartPage": body_start_page,
                "watermarkPages": watermark_pages,
                "signaturePresent": signature_ok,
                "generatedByLabelPresent": context["generatedByLabel"] in cover_text,
                "word": {**word_structure, **word_rendering},
                "pages": pages,
            }
            if _sha256(path) != current["sha256"] or path.stat().st_size != current["size"]:
                raise ReportFailure("PDF 产物在验收期间发生变化，请重新验收")
            if (
                _sha256(word) != word_current["sha256"]
                or word.stat().st_size != word_current["size"]
            ):
                raise ReportFailure("Word 产物在验收期间发生变化，请重新验收")
            return validation
        finally:
            if temp_path is not None:
                shutil.rmtree(temp_path, ignore_errors=True)

    def _validate_manifest_markers(
        self,
        manifest: dict[str, Any] | None,
        *,
        render: dict[str, Any],
        extracted_text: str,
    ) -> tuple[list[str], list[str], list[str]]:
        if manifest is None:
            return [], [], []
        charts = manifest.get("charts")
        citations = manifest.get("citations")
        sections = manifest.get("sections")
        section_numbers = manifest.get("sectionNumbers")
        heading_numbers = manifest.get("headingNumbers")
        context = render.get("documentContext")
        if (
            not isinstance(charts, list)
            or not isinstance(citations, list)
            or not isinstance(sections, list)
            or not isinstance(context, dict)
            or section_numbers != context.get("sectionNumbers")
            or heading_numbers != context.get("headingNumbers")
        ):
            raise ReportFailure("报告产物清单无效")
        chart_paths = {
            item.get("path")
            for item in charts
            if isinstance(item, dict) and isinstance(item.get("path"), str)
        }
        rendered_paths = {
            item.get("path")
            for item in render.get("images", [])
            if isinstance(item, dict) and isinstance(item.get("path"), str)
        }
        if len(chart_paths) != len(charts) or chart_paths != rendered_paths:
            raise ReportFailure("PDF 图表与产物清单不一致")
        chart_ids: list[str] = []
        for item in charts:
            value = item.get("chartId") if isinstance(item, dict) else None
            if isinstance(value, str):
                chart_ids.append(value)
        citation_ids: list[str] = []
        for item in citations:
            value = item.get("citationId") if isinstance(item, dict) else None
            if isinstance(value, str):
                citation_ids.append(value)
        presentations = render.get("citationPresentations")
        if (
            len(citation_ids) != len(citations)
            or not isinstance(presentations, list)
            or len(presentations) != len(citation_ids)
            or any(
                not isinstance(item, dict)
                or item.get("citationId") != citation_id
                or item.get("alias") != f"[引用 {index:03d}]"
                or not isinstance(item.get("label"), str)
                for index, (citation_id, item) in enumerate(
                    zip(citation_ids, presentations, strict=True), start=1
                )
            )
            or any(item["alias"] in extracted_text for item in presentations)
            or "实际引用附录" in extracted_text
            or "[[citation:" in extracted_text
        ):
            raise ReportFailure("PDF 不应显示引用标识或实际引用附录")
        section_ids = [item for item in sections if isinstance(item, str)]
        if len(section_ids) != len(sections) or "[[section:" in extracted_text:
            raise ReportFailure("PDF 不应显示关键章节标识")
        return chart_ids, citation_ids, section_ids

    @staticmethod
    def _check_pdf_bounds(path: Path) -> None:
        try:
            import pypdf
        except ImportError as error:
            raise ReportFailure("PDF 视觉验收依赖不可用") from error
        if not path.is_file() or path.stat().st_size > MAX_PDF_BYTES:
            raise ReportFailure("PDF 文件不能超过 200 MiB")
        try:
            page_count = len(pypdf.PdfReader(str(path)).pages)
        except pypdf.errors.PdfReadError as error:
            raise ReportFailure("PDF 视觉验收无法打开产物") from error
        if not 1 <= page_count <= MAX_PDF_PAGES:
            raise ReportFailure("PDF 页数必须在 1 至 200 页之间")


def main(arguments: list[str] | None = None) -> int:
    values = arguments if arguments is not None else sys.argv[1:]
    try:
        if len(values) != 2:
            raise ReportFailure("报表渲染参数无效")
        action, payload_text = values
        payload = json.loads(payload_text)
        runtime = ReportRuntime(Path.cwd())
        if action == "render_markdown":
            result = runtime.render_markdown(
                payload["job"],
                payload["markdown_path"],
                payload["output_path"],
                payload["temporary_path"],
                payload.get("page_layout"),
                payload.get("word_output_path"),
            )
        elif action == "validate_pdf":
            result = runtime.validate_pdf(
                payload["job"],
                payload["pdf_path"],
                payload["temporary_directory"],
                payload.get("artifact_manifest"),
                payload.get("word_path"),
            )
        else:
            raise ReportFailure("未知报表操作")
        encoded = json.dumps(result, ensure_ascii=False)
        if len(encoded.encode()) + 1 > MAX_RESULT_BYTES:
            raise ReportFailure("报表结果超过返回边界")
        print(encoded)
        return 0
    except (KeyError, TypeError, json.JSONDecodeError, ReportFailure) as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False))
        return 1
    except Exception:
        print(json.dumps({"error": "报表运行时执行失败"}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
