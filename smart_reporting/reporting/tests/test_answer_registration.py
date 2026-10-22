"""Normal publication registers only application-approved, durable mappings."""

import hashlib
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from agno.agent import Agent
from agno.run import RunContext

from smart_reporting.reporting.agent import ReportingPhaseOpenAIChat
from smart_reporting.reporting.data_sources import DatasetHandle
from smart_reporting.reporting.delivery.artifacts_v1 import ArtifactFile, DatasetLineage
from smart_reporting.reporting.trace.contracts_v1 import RevisionTraceIndexV1
from smart_reporting.reporting.trace.quality_bundle import QualityBundle
from smart_reporting.reporting.workflow.runtime import ReportWorkflowRuntime
from smart_reporting.reporting.workflow.state import ReportingRunState, apply

from .test_quality_answer_review import case


async def publication_case(tmp_path, *, provider=None, fault=None):
    editor, context, index, markdown, _, questions, plan = await case(tmp_path)
    durable = ReportingRunState.initial(report_run_id=index.report_id, external_run_id=context.scope["externalRunId"],
        thread_id=context.scope["callerThreadId"], owner_user_id=context.scope["userId"])
    durable = apply(durable, {"name": "set_analysis_plan", "commandId": "plan", "payload": {
        "analysisIds": list(questions.analysis_to_question),
        "analysisPlans": {aid: {"questionId": qid} for aid, qid in questions.analysis_to_question.items()},
        "questionSpecs": questions.model_dump(mode="json", by_alias=True),
    }}, durable.state_version).state
    commands = []

    class Repository:
        async def get(self, run_id):
            assert run_id == index.report_id
            return durable

        async def get_by_external_run_id(self, external_id):
            assert external_id == context.scope["externalRunId"]
            return durable

        async def apply(self, run_id, command, *, expected_version):
            nonlocal durable
            commands.append(command.name)
            if command.name == "freeze_answer_review_plan" and fault == "freeze":
                raise OSError("injected durable freeze failure")
            result = apply(durable, command, expected_version)
            durable = result.state
            if command.name == "freeze_answer_review_plan" and fault == "receipt":
                return SimpleNamespace(state=durable.model_copy(update={"report_run_id": "foreign"}))
            return result

    runtime = ReportWorkflowRuntime(db=SimpleNamespace(),
        reporting_agent_template=Agent(model=ReportingPhaseOpenAIChat(id="deepseek-v4-flash-0731", api_key="test")),
        task_runner=SimpleNamespace(), workspace_service=editor.workspace, registry=SimpleNamespace(),
        profiles=SimpleNamespace(), planner_enable_thinking=False, state_repository=Repository(),
        answer_review_plan_provider=provider)
    runtime._scope = lambda _: context.scope
    runtime.workspace_service = editor.workspace
    runtime._profile = lambda _: SimpleNamespace(effective_profile_hash="0" * 64)
    runtime._build_drilldown_metrics = AsyncMock(return_value=())
    run_context = RunContext(run_id=index.report_id, session_id=context.scope["threadId"], session_state={})
    files = {file.resource_id: file for file in index.files}
    dataset = index.datasets[0]
    csv = files[dataset.file_resource_id]
    handle = DatasetHandle(dataset_id=dataset.dataset_id, source_id="fixture-url", source_type=dataset.source_type,
        path=csv.path, size=csv.size, sha256=csv.sha256, row_count=dataset.row_count,
        filename=dataset.filename, requirement_id=dataset.requirement_id, sql_hash=dataset.sql_hash or "a" * 64)
    lineage = DatasetLineage(datasetId=handle.dataset_id, sourceId=handle.source_id, sourceType=handle.source_type,
        size=handle.size, sha256=handle.sha256, rowCount=handle.row_count,
        requirementId=handle.requirement_id, sqlHash=handle.sql_hash)
    args = {"handles": (handle,), "lineage": (lineage,), "revision": 1, "run_context": run_context,
        "markdown_artifact": ArtifactFile.model_validate(files[index.markdown_file_resource_id].model_dump(exclude={"resource_id"})),
        "markdown": markdown, "subject_bindings": index.subject_bindings,
        "fact_files": {entry.analysis_id: ArtifactFile(path=files[entry.file_resource_id].path,
            mediaType=files[entry.file_resource_id].media_type, size=files[entry.file_resource_id].size,
            sha256=files[entry.file_resource_id].sha256) for entry in index.fact_files}}
    return runtime, run_context, context, questions, plan, args, commands, editor


async def run_publication(runtime, context, args):
    identity = await runtime._write_trace_index(context.artifact_manifest.path, **args)
    raw = await runtime.workspace_service.read_limited_regular_file(context.scope["threadId"], identity.path, max_bytes=16 * 1024 * 1024)
    index = RevisionTraceIndexV1.model_validate_json(raw)
    quality_file = next(file for file in index.files if file.path.endswith("/quality-bundle-v1.json"))
    raw = await runtime.workspace_service.read_limited_regular_file(context.scope["threadId"], quality_file.path, max_bytes=16 * 1024 * 1024)
    assert hashlib.sha256(raw).hexdigest() == quality_file.sha256
    return QualityBundle.model_validate_json(raw)


@pytest.mark.anyio
@pytest.mark.parametrize("anyio_backend", ["asyncio"])
async def test_configured_hook_respects_disabled_quality_policy(tmp_path):
    async def provider(_):
        pytest.fail("Quality-off must not call an approval integration")

    runtime, run_context, context, _, _, args, commands, _ = await publication_case(tmp_path, provider=provider)
    assert runtime.answer_review_plan_provider is provider
    run_context.session_state["report_request_context"] = {"assertionPolicy": {"mode": "off"}}
    identity = await runtime._write_trace_index(context.artifact_manifest.path, **args)
    index = RevisionTraceIndexV1.model_validate_json(await runtime.workspace_service.read_limited_regular_file(
        context.scope["threadId"], identity.path, max_bytes=16 * 1024 * 1024))
    assert not any(file.path.endswith("quality-bundle-v1.json") for file in index.files)
    assert not commands


@pytest.mark.anyio
@pytest.mark.parametrize("anyio_backend", ["asyncio"])
async def test_application_mapping_freezes_before_current_answer_and_is_reused_on_resume(tmp_path):
    runtime, run_context, context, questions, plan, args, commands, _ = await publication_case(tmp_path)
    seen = []

    async def provider(registration):
        seen.append(registration)
        assert registration.report_id == "report-1" and registration.questions == questions
        assert registration.candidates[0].statement_verified
        assert registration.candidates[0].assertion.request == plan.criteria[0].request
        return plan

    runtime.answer_review_plan_provider = provider
    quality = await run_publication(runtime, context, args)
    assert commands == ["freeze_answer_review_plan", "record_answer_review"]
    assert len(seen) == 1
    assert quality.answer_review_plan == plan
    assert quality.question_coverage["answered"] == 1
    assert quality.management_overview.rows[0].status == "met"
    assert quality.management_overview.rows[1].status == "unknown"
    resumed = await run_publication(runtime, context, args)
    assert len(seen) == 1 and resumed.input_sha256 == quality.input_sha256
    assert commands.count("freeze_answer_review_plan") == 1
    assert (await runtime.state_repository.get("report-1")).payload["answerReviewPlan"] == plan.model_dump(mode="json", by_alias=True)


@pytest.mark.anyio
@pytest.mark.parametrize("anyio_backend", ["asyncio"])
async def test_approval_of_mapping_cannot_upgrade_an_unknown_fact(tmp_path):
    import json

    runtime, _, context, _, plan, args, commands, _ = await publication_case(tmp_path)
    analysis_id, file = next(iter(args["fact_files"].items()))
    raw = await runtime.workspace_service.read_limited_regular_file(context.scope["threadId"], file.path, max_bytes=16 * 1024 * 1024)
    document = json.loads(raw)
    document["metrics"][0]["scope"] = {}
    raw = json.dumps(document, ensure_ascii=False).encode()
    await runtime.workspace_service.awrite_bytes(context.scope["threadId"], file.path, raw, overwrite=True)
    args["fact_files"][analysis_id] = file.model_copy(update={"size": len(raw), "sha256": hashlib.sha256(raw).hexdigest()})

    async def provider(registration):
        assert registration.candidates[0].assertion.status == "unknown"
        assert not registration.candidates[0].statement_verified
        return plan

    runtime.answer_review_plan_provider = provider
    quality = await run_publication(runtime, context, args)
    assert "freeze_answer_review_plan" in commands and quality.answer_review_plan == plan
    assert quality.question_coverage["answered"] == 0
    assert quality.management_overview.rows[0].status == "unknown"
    assert not quality.management_overview.rows[0].evidence


@pytest.mark.anyio
@pytest.mark.parametrize("anyio_backend", ["asyncio"])
async def test_registration_context_does_not_hide_an_unverifiable_required_subject(tmp_path):
    runtime, _, context, _, plan, args, commands, _ = await publication_case(tmp_path)
    changed = args["markdown"].replace("2,400", "9,999")
    assert changed != args["markdown"]
    args.update(markdown=changed, markdown_artifact=args["markdown_artifact"].model_copy(update={
        "size": len(changed.encode()), "sha256": hashlib.sha256(changed.encode()).hexdigest()}))
    await runtime.workspace_service.awrite_text(context.scope["threadId"], context.markdown_path, changed, overwrite=True)
    observed = []

    async def provider(registration):
        observed.extend(registration.candidates)
        assert len(registration.candidates) == 1
        assert registration.candidates[0].assertion is None
        assert registration.candidates[0].number_reference.token == plan.criteria[0].request.operands[0]
        assert not registration.candidates[0].statement_verified
        return plan

    runtime.answer_review_plan_provider = provider
    quality = await run_publication(runtime, context, args)
    assert observed and "freeze_answer_review_plan" not in commands
    assert quality.question_coverage["answered"] == 0
    assert quality.answer_review_plan is None
    assert any(issue.code == "answer_review_registration_unavailable" for issue in quality.issues)


@pytest.mark.anyio
@pytest.mark.parametrize("anyio_backend", ["asyncio"])
@pytest.mark.parametrize("fault", ["disabled", "decline", "provider", "freeze", "receipt", "subject", "selector", "question", "mutated_context", "timeout", "budget", "context_budget", "invalid_model", "invalid_frozen"])
async def test_registration_cannot_manufacture_approval_or_erase_independent_facts(tmp_path, monkeypatch, fault):
    from smart_reporting.reporting.workflow import answer_registration

    runtime, run_context, context, questions, plan, args, commands, _ = await publication_case(tmp_path, fault=fault)
    if fault == "timeout":
        monkeypatch.setattr(answer_registration, "PROVIDER_TIMEOUT_SECONDS", 0.001)
    elif fault == "budget":
        monkeypatch.setattr(answer_registration, "MAX_REGISTRATION_BYTES", 1)
    elif fault == "context_budget":
        monkeypatch.setattr(answer_registration, "MAX_CONTEXT_BYTES", 1)
    elif fault == "invalid_frozen":
        (await runtime.state_repository.get("report-1")).payload["answerReviewPlan"] = {"version": "bad"}

    async def provider(registration):
        if fault == "invalid_frozen":
            pytest.fail("A damaged frozen mapping must not be silently replaced")
        if fault == "timeout":
            import anyio

            await anyio.sleep(0.1)
        if fault == "provider":
            raise OSError("injected approval service failure")
        if fault == "decline":
            return None
        candidate = plan.criteria[0]
        if fault == "subject":
            candidate = candidate.model_copy(update={"subject_ids": ("sub-" + "f" * 16,)})
        elif fault == "selector":
            candidate = candidate.model_copy(update={"request": candidate.request.model_copy(update={
                "operands": ("{{value:fact-aaaaaaaaaaaaaaaa:periodValues.0:万元}}",)})})
        elif fault == "question":
            return plan.model_copy(update={"question_set_sha256": "f" * 64})
        elif fault == "invalid_model":
            candidate = candidate.model_copy(update={"subject_ids": ()})
        elif fault == "mutated_context":
            registration.candidates[0].assertion.request.__dict__["operands"] = ("{{value:fact-aaaaaaaaaaaaaaaa:periodValues.0:万元}}",)
            candidate = candidate.model_copy(update={"request": registration.candidates[0].assertion.request})
        return plan.model_copy(update={"criteria": (candidate,)})

    # Model/session content cannot install the application hook or a passing plan.
    run_context.session_state.update({"answerReviewPlan": plan.model_dump(mode="json", by_alias=True),
        "answer_review_plan_provider": provider})
    if fault != "disabled":
        runtime.answer_review_plan_provider = provider
    quality = await run_publication(runtime, context, args)
    assert quality.assertions[0].status == "supported"
    assert quality.answer_review_plan is None
    assert quality.question_coverage["answered"] == 0
    assert not quality.management_overview.rows[0].evidence
    assert ("freeze_answer_review_plan" in commands) == (fault in {"freeze", "receipt"})
    assert any(issue.code == "answer_review_registration_unavailable" for issue in quality.issues) == (fault not in {"disabled", "decline"})


@pytest.mark.anyio
@pytest.mark.parametrize("anyio_backend", ["asyncio"])
async def test_registered_mapping_reaches_real_exports_and_withdraws_after_scope_edit(tmp_path):
    import pypdf
    from docx import Document

    from smart_reporting.reporting.delivery.report_runtime.runtime import ReportRuntime
    from smart_reporting.reporting.trace.contracts_v1 import canonical_sha256
    from smart_reporting.reporting.workflow.management_overview import (
        ManagementOverview,
        management_overview_projection,
    )
    from smart_reporting.reporting.workflow.overview_sources import registered_management_overview

    from .lineage_fixtures.manifest import register_trace_manifest
    from .test_report_runtime_full_render import _DEPENDENCIES_OK, _SKIP_REASON, _build_workspace

    if not _DEPENDENCIES_OK:
        pytest.skip(_SKIP_REASON)
    runtime, _, context, _, plan, args, _, editor = await publication_case(tmp_path / "publication")

    async def provider(_registration):
        return plan  # Trusted synthetic domain approval, not model-authored criteria.

    runtime.answer_review_plan_provider = provider
    markdown = "# 智能运营全链路渲染验证报告\n\n[[section:section_002]]\n\n## 1. Report\n\n" + args["markdown"]
    args.update(markdown=markdown, markdown_artifact=args["markdown_artifact"].model_copy(update={
        "size": len(markdown.encode()), "sha256": hashlib.sha256(markdown.encode()).hexdigest()}))
    await runtime.workspace_service.awrite_text(context.scope["threadId"], context.markdown_path, markdown, overwrite=True)
    quality = await run_publication(runtime, context, args)
    index_path = context.markdown_path.replace("report.md", "trace-index-v1.json")
    raw = await runtime.workspace_service.read_limited_regular_file(context.scope["threadId"], index_path, max_bytes=16 * 1024 * 1024)
    index = RevisionTraceIndexV1.model_validate_json(raw)
    manifest_identity = await register_trace_manifest(runtime.workspace_service, context.scope["threadId"], index, overwrite=True)
    context = context.model_copy(update={"artifact_manifest": manifest_identity})
    manifest = await editor.trace.load_manifest(context)
    # Register the same profile as the real runtime/index, not the fixture default.
    manifest = manifest.model_copy(update={"effective_profile_hash": "0" * 64})
    raw = manifest.model_dump_json(by_alias=True).encode()
    await runtime.workspace_service.awrite_bytes(context.scope["threadId"], manifest_identity.path, raw, overwrite=True)
    context = context.model_copy(update={"artifact_manifest": manifest_identity.model_copy(update={
        "size": len(raw), "sha256": hashlib.sha256(raw).hexdigest()})})
    approved = await registered_management_overview(runtime.workspace_service, context.scope["threadId"], manifest,
        canonical_sha256({key: context.scope[key] for key in ("userId", "threadId")}))
    assert approved.rows[0].status == "met" and quality.answer_review_plan == plan
    workspace = tmp_path / "render"
    state = _build_workspace(workspace)
    state["_documentContext"].update({
        "periodLabel": "2025-10",
        "sections": [{"code": "section_002", "sectionNumber": "1", "title": "Report"}],
        "headingNumbers": [{"level": 2, "number": "1", "title": "Report", "sectionCode": "section_002",
                            "anchor": "report-section-section_002"}],
    })
    for stage, text in enumerate((markdown, markdown.replace("北院", "南院"))):
        if stage:
            checked = await editor.trace.validate(context, text, hashlib.sha256(text.encode()).hexdigest())
            approved = ManagementOverview.model_validate(checked["quality"]["managementOverview"])
            assert approved.rows[0].status == "unknown"
        (workspace / "report.md").write_text(text)
        state["_managementOverviewProjection"] = management_overview_projection(approved, hashlib.sha256(text.encode()).hexdigest())
        result = ReportRuntime(workspace).render_markdown(state, "report.md", f"revision-{stage}/report.pdf",
            f".reporting-tmp/workspace-report-registration-{stage}-render/render.pdf")
        assert result["status"] == "rendered"
        pdf = pypdf.PdfReader(workspace / result["pdfPath"])
        word = Document(workspace / result["wordPath"])
        pdf_text = "".join((page.extract_text() or "") for page in pdf.pages)
        table_text = "".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
        assert ("2400.0万元" in table_text) == (stage == 0)
        assert ("该项已核验" in pdf_text) == (stage == 0)
        assert "该项待核实" in table_text
        assert "智能运营验证机构" in pdf_text
        assert any("智能运营验证机构" in paragraph.text for paragraph in word.paragraphs)
        state["render"] = result["render"]
        validation = ReportRuntime(workspace).validate_pdf(state, result["pdfPath"],
            f".reporting-tmp/workspace-report-registration-{stage}-validate", word_path=result["wordPath"])
        assert validation["ok"] is True  # Structure/identity, not business acceptance.
        assert validation["projectionCheck"]["status"] == "verified"
