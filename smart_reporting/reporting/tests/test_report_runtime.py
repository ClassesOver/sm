from __future__ import annotations

import json
import shutil
import uuid
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from agno.run import RunContext

from smart_reporting.reporting.delivery.report_runtime import cli as runtime_cli
from smart_reporting.reporting.delivery.report_runtime import docx as docx_module
from smart_reporting.reporting.delivery.report_runtime import runtime as runtime_module
from smart_reporting.reporting.delivery.report_runtime.docx import (
    _WORD_PAGE_FIELDS,
    _fit_image_dimensions,
    _postprocess_docx,
)
from smart_reporting.reporting.delivery.report_runtime.markdown import (
    _bind_heading_anchors,
    _document_context,
    _normalize_cjk_strong_markers,
    _strip_strong_boundaries,
    normalize_report_markdown_strong_spacing,
)
from smart_reporting.reporting.delivery.report_runtime.pdf import (
    DEFAULT_PAGE_LAYOUT,
    _apply_pdf_page_decorations,
    _page_number_context,
)
from smart_reporting.reporting.delivery.report_runtime.validation import (
    ReportFailure,
    _temporary_pdf_path,
    _validation_directory,
)
from smart_reporting.reporting.workspace import WorkspaceReportService, _report_runtime_digest
from smart_reporting.workspace import WorkspaceError


@pytest.mark.parametrize("duplicate_code", [False, True])
def test_document_context_allows_repeated_titles_with_distinct_section_identity(duplicate_code):
    codes = ["appendix_a", "appendix_a" if duplicate_code else "appendix_b"]
    context = {
        "title": "运营报告", "periodLabel": "2025年", "organizationName": "医院",
        "generatedByLabel": "智能报告", "watermarkText": "内部资料", "generatedDate": "2026-10-15",
        "sections": [{"code": code, "sectionNumber": str(i), "title": "补充管理发现"}
                     for i, code in enumerate(codes, 1)],
        "sectionNumbers": ["1", "2"],
        "headingNumbers": [{"level": 2, "number": str(i), "title": "补充管理发现",
                            "sectionCode": code, "anchor": f"report-section-{i}"}
                           for i, code in enumerate(codes, 1)],
    }
    if duplicate_code:
        with pytest.raises(ReportFailure, match="重复项"):
            _document_context(context)
    else:
        assert _document_context(context)["sections"] == context["sections"]


def test_render_docx_uses_libreoffice_when_pandoc_is_unavailable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    output = tmp_path / "render.docx"
    calls: list[list[str]] = []

    monkeypatch.setattr(
        docx_module.shutil,
        "which",
        lambda name: "/usr/bin/libreoffice" if name == "libreoffice" else None,
    )

    def fake_run(command, **_kwargs):
        calls.append(command)
        Path(command[command.index("--outdir") + 1], output.name).write_bytes(
            b"PK\x03\x04fake-docx"
        )
        return type("Completed", (), {"returncode": 0})()

    monkeypatch.setattr(docx_module.subprocess, "run", fake_run)
    monkeypatch.setattr(docx_module, "_postprocess_docx", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(docx_module, "_validate_docx_structure", lambda *_args, **_kwargs: {})

    result = docx_module._render_docx(
        "<html><body>报告</body></html>",
        source_parent=tmp_path,
        output=output,
        context={"sections": [], "headingNumbers": []},
        layout={},
    )

    assert result == {}
    assert calls and calls[0][0] == "/usr/bin/libreoffice"
    assert calls[0][1].startswith("-env:UserInstallation=file://")
    assert calls[0][2:6] == ["--headless", "--convert-to", "docx:MS Word 2007 XML", "--outdir"]
    assert "--outdir" in calls[0]
    assert output.is_file()


def test_render_docx_removes_temp_dir_when_libreoffice_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    output = tmp_path / "render.docx"
    calls: list[list[str]] = []
    monkeypatch.setattr(
        docx_module.shutil,
        "which",
        lambda name: "/usr/bin/libreoffice" if name == "libreoffice" else None,
    )

    def fake_run(command, **_kwargs):
        calls.append(command)
        return type("Completed", (), {"returncode": 1, "stderr": b"fail", "stdout": b""})()

    monkeypatch.setattr(docx_module.subprocess, "run", fake_run)

    with pytest.raises(docx_module.ReportFailure, match="Word 渲染失败"):
        docx_module._render_docx(
            "<html><body>报告</body></html>",
            source_parent=tmp_path,
            output=output,
            context={"sections": [], "headingNumbers": []},
            layout={},
        )

    outdir = Path(calls[0][calls[0].index("--outdir") + 1])
    assert not outdir.exists()
    assert not output.is_file()


@pytest.mark.parametrize("kind", ["render", "validate"])
def test_report_temporary_paths_are_private_workspace_directories(tmp_path: Path, kind: str):
    relative = f".reporting-tmp/workspace-report-test-{kind}"
    create = _temporary_pdf_path if kind == "render" else _validation_directory
    value = relative + "/render.pdf" if kind == "render" else relative
    path = create(tmp_path, value)
    directory = path.parent if kind == "render" else path
    assert directory == tmp_path / relative
    assert directory.is_dir()
    assert directory.stat().st_mode & 0o777 == 0o700
    with pytest.raises(ReportFailure):
        create(tmp_path, value)


@pytest.mark.parametrize("kind", ["render", "validate"])
@pytest.mark.parametrize("prefix", ["/tmp", "../outside", "reports", ".reporting-tmp/nested"])
def test_report_temporary_paths_reject_unscoped_directories(tmp_path: Path, kind: str, prefix: str):
    create = _temporary_pdf_path if kind == "render" else _validation_directory
    value = f"{prefix}/workspace-report-test-{kind}"
    if kind == "render":
        value += "/render.pdf"
    with pytest.raises(ReportFailure):
        create(tmp_path, value)


@pytest.mark.parametrize("kind", ["render", "validate"])
def test_report_temporary_paths_reject_symlink_root(tmp_path: Path, kind: str):
    outside = tmp_path / "outside"
    outside.mkdir()
    (tmp_path / ".reporting-tmp").symlink_to(outside, target_is_directory=True)
    create = _temporary_pdf_path if kind == "render" else _validation_directory
    value = f".reporting-tmp/workspace-report-test-{kind}"
    if kind == "render":
        value += "/render.pdf"
    with pytest.raises(ReportFailure):
        create(tmp_path, value)
    assert list(outside.iterdir()) == []


@pytest.mark.parametrize(
    ("width", "height", "maximum_width", "maximum_height", "expected"),
    [
        (2_000, 1_000, 1_000, 1_500, (1_000, 500)),
        (1_000, 2_000, 1_500, 1_000, (500, 1_000)),
        (1_000, 500, 1_500, 1_000, (1_000, 500)),
    ],
)
def test_fit_image_dimensions_preserves_aspect_ratio_within_bounds(
    width: int,
    height: int,
    maximum_width: int,
    maximum_height: int,
    expected: tuple[int, int],
) -> None:
    assert _fit_image_dimensions(width, height, maximum_width, maximum_height) == expected


def test_pdf_page_decorations_are_merged_above_report_content(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import pypdf

    path = tmp_path / "report.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=595.276, height=841.89)
    writer.add_blank_page(width=595.276, height=841.89)
    with path.open("wb") as stream:
        writer.write(stream)

    monkeypatch.setattr(
        "smart_reporting.reporting.delivery.report_runtime.pdf._pdf_section_pages",
        lambda _reader, _sections: {"section_001": 2},
    )
    merge_layers: list[bool] = []
    original_merge_page = pypdf.PageObject.merge_page

    def track_merge_page(self, page2, expand=False, over=True):
        merge_layers.append(over)
        return original_merge_page(self, page2, expand=expand, over=over)

    monkeypatch.setattr(pypdf.PageObject, "merge_page", track_merge_page)

    _apply_pdf_page_decorations(
        path,
        context={
            "title": "测试报告",
            "organizationName": "测试机构",
            "watermarkText": "内部资料",
            "sections": [{"code": "section_001"}],
        },
        layout=DEFAULT_PAGE_LAYOUT,
    )

    assert merge_layers == [True]


def test_cli_forwards_only_pdf_and_word_output_paths(
    monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    captured: dict[str, object] = {}

    class FakeRuntime:
        def __init__(self, _workspace: Path) -> None:
            pass

        def render_markdown(self, *args: object) -> dict[str, object]:
            captured["args"] = args
            return {"status": "rendered", "wordPath": "reports/report.docx"}

    monkeypatch.setattr(runtime_cli, "ReportRuntime", FakeRuntime)
    assert (
        runtime_cli.main(
            [
                "render_markdown",
                '{"job": {}, "markdown_path": "report.md", "output_path": "report.pdf", '
                '"temporary_path": ".reporting-tmp/workspace-report-test-render/render.pdf", '
                '"word_output_path": "report.docx"}',
            ]
        )
        == 0
    )

    assert captured["args"][-1] == "report.docx"
    assert len(captured["args"]) == 6
    assert '"wordPath": "reports/report.docx"' in capsys.readouterr().out


@pytest.mark.parametrize("failed_format", [None, "pdf", "word", "both", "pdf_copy", "word_copy"])
def test_render_markdown_returns_only_pdf_and_word_artifact_identity(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, failed_format: str | None,
) -> None:
    source = tmp_path / "report.md"
    source.write_text(
        "# 测试报告\n\n## 1. 经营概览\n\n[[section:overview]]\n正文。\n",
        encoding="utf-8",
    )
    output = tmp_path / "revision" / "report.pdf"
    output.parent.mkdir()
    state = {
        "jobId": "job-1",
        "sources": [],
        "_documentContext": {
            "title": "测试报告",
            "periodLabel": "2026 年",
            "organizationName": "测试机构",
            "generatedByLabel": "Reporting Agent",
            "watermarkText": "内部资料",
            "generatedDate": "2026-08-17",
            "sections": [{"code": "overview", "sectionNumber": "1", "title": "经营概览"}],
            "sectionNumbers": ["1"],
            "headingNumbers": [
                {
                    "level": 2,
                    "number": "1",
                    "title": "经营概览",
                    "sectionCode": "overview",
                    "anchor": "report-heading-overview",
                }
            ],
        },
    }
    monkeypatch.setattr(runtime_module.ReportRuntime, "_validate_datasets", lambda *_: None)
    monkeypatch.setattr(runtime_module.ReportRuntime, "_images", lambda *_: set())
    monkeypatch.setattr(
        runtime_module, "_apply_pdf_page_decorations", lambda *_args, **_kwargs: None
    )
    monkeypatch.setattr(
        runtime_module,
        "_render_docx",
        lambda *_args, **kwargs: kwargs["output"].write_bytes(b"docx") or {},
    )
    monkeypatch.setattr(runtime_module, "_validate_docx_structure", lambda *_args, **_kwargs: {})
    monkeypatch.setattr(
        Path,
        "replace",
        lambda *_args, **_kwargs: pytest.fail("私有 /tmp 与 workspace 之间不得使用 rename"),
    )

    if failed_format in {"pdf_copy", "word_copy"}:
        original_copy = shutil.copyfile
        def fail_copy(source, destination, *args, **kwargs):
            if str(destination).endswith(".pdf" if failed_format == "pdf_copy" else ".docx"):
                raise OSError("injected format copy failure")
            return original_copy(source, destination, *args, **kwargs)
        monkeypatch.setattr(shutil, "copyfile", fail_copy)
    if failed_format in {"pdf", "both"}:
        def fail_pdf(*_args, **_kwargs):
            raise runtime_module.ReportFailure("injected PDF failure")
        monkeypatch.setattr(runtime_module, "_apply_pdf_page_decorations", fail_pdf)
    if failed_format in {"word", "both"}:
        def fail_word(*_args, **_kwargs):
            raise runtime_module.ReportFailure("injected Word failure")
        monkeypatch.setattr(runtime_module, "_render_docx", fail_word)
    if failed_format == "both":
        with pytest.raises(runtime_module.ReportFailure, match="均未成功"):
            runtime_module.ReportRuntime(tmp_path).render_markdown(state, "report.md", "revision/report.pdf", ".reporting-tmp/workspace-report-failure-render/render.pdf")
        assert not output.exists() and not output.with_suffix(".docx").exists()
        return
    temporary_root = tmp_path / ".reporting-tmp" / f"workspace-report-{uuid.uuid4().hex}-render"
    try:
        result = runtime_module.ReportRuntime(tmp_path).render_markdown(
            state,
            "report.md",
            "revision/report.pdf",
            str((temporary_root / "render.pdf").relative_to(tmp_path)),
        )
    finally:
        shutil.rmtree(temporary_root, ignore_errors=True)

    assert "html" not in result["render"]
    assert "htmlPath" not in result
    failed_label = failed_format.removesuffix("_copy") if failed_format else None
    assert output.is_file() is (failed_label != "pdf")
    assert output.with_suffix(".docx").is_file() is (failed_label != "word")
    assert ("pdf" in result["render"]) is (failed_label != "pdf")
    assert ("word" in result["render"]) is (failed_label != "word")
    assert result["status"] == ("rendered" if failed_format is None else "partial_rendered")
    assert bool(result["renderIssues"]) is (failed_format is not None)
    assert not output.with_suffix(".html").exists()


@pytest.mark.anyio
async def test_report_runtime_uploads_verified_package_for_fixed_python_entrypoint(
    tmp_path: Path,
) -> None:
    del tmp_path

    class Workspace:
        def __init__(self) -> None:
            self.writes: list[tuple[str, str, bytes]] = []
            self.commands: list[tuple[str, list[str], int, int]] = []
            self.deletes: list[tuple[str, str]] = []

        async def awrite_bytes(self, thread: str, path: str, content: bytes) -> None:
            self.writes.append((thread, path, content))

        async def arun_command(
            self, thread: str, args: list[str], *, timeout: int, tail: int
        ) -> str:
            self.commands.append((thread, args, timeout, tail))
            return '{"status":"ok"}\n{"__reportExitCode":0}'

        async def adelete_file(self, thread: str, path: str) -> None:
            self.deletes.append((thread, path))

    workspace = Workspace()
    report_workspace = WorkspaceReportService(workspace)  # type: ignore[arg-type]

    result = await report_workspace._run_report_runtime(
        "validate_pdf",
        {"job": {}},
        RunContext(run_id="report-runtime-run", session_id="report-runtime-package"),
    )

    assert len(workspace.commands) == 1
    thread, args, timeout, tail = workspace.commands[0]
    assert thread == "report-runtime-package"
    assert len(workspace.writes) == 2
    (_thread, payload_path, payload), (_thread, archive_path, archive) = workspace.writes
    assert payload_path.endswith("-payload.json")
    assert payload == b'{"job":{}}'
    assert archive_path.startswith(".workspace-report-runtime-")
    assert archive_path.endswith(f"-{_report_runtime_digest()}.zip")
    assert workspace.deletes == [(thread, archive_path), (thread, payload_path)]
    with zipfile.ZipFile(BytesIO(archive)) as package:
        assert set(package.namelist()) == {
            "report_runtime/__init__.py",
            "report_runtime/cli.py",
            "report_runtime/docx.py",
            "report_runtime/markdown.py",
            "report_runtime/overview.py",
            "report_runtime/pdf.py",
            "report_runtime/projection.py",
            "report_runtime/runtime.py",
            "report_runtime/theme.py",
            "report_runtime/validation.py",
        }
    script = args[2]
    assert args[1] == "-c"
    assert "from report_runtime.cli import main" in script
    assert "sys.path.insert(0" in script
    assert archive_path in script
    assert _report_runtime_digest() in script
    assert "报表运行时版本不匹配" in script
    assert "validate_pdf" in script
    assert payload_path in script
    assert timeout == 600
    assert tail == 200
    assert result == {"status": "ok"}


@pytest.mark.anyio
async def test_report_runtime_passes_large_payload_without_argv_limit(tmp_path: Path) -> None:
    import subprocess

    class Workspace:
        async def awrite_bytes(self, _thread: str, path: str, content: bytes) -> None:
            (tmp_path / path).write_bytes(content)

        async def arun_command(
            self, _thread: str, args: list[str], *, timeout: int, tail: int
        ) -> str:
            del tail
            process = subprocess.run(
                args, cwd=tmp_path, capture_output=True, text=True, timeout=timeout, check=False
            )
            return process.stdout

        async def adelete_file(self, _thread: str, path: str) -> None:
            (tmp_path / path).unlink(missing_ok=True)

    # 真实渲染回执与产物清单可超过 Linux 单个命令行参数的 128 KiB 上限。
    payload = {"job": {"notes": ["引用覆盖说明" * 20] * 800}}
    assert len(json.dumps(payload, ensure_ascii=False).encode()) > 128 * 1024

    with pytest.raises(WorkspaceError) as caught:
        await WorkspaceReportService(Workspace())._run_report_runtime(  # type: ignore[arg-type]
            "validate_pdf",
            payload,
            RunContext(run_id="report-runtime-run", session_id="report-runtime-package"),
        )

    # 进程成功启动并由 runtime 返回结构化业务错误，而不是 Argument list too long。
    assert "Argument list too long" not in str(caught.value)
    assert "报表运行时返回无效结果" not in str(caught.value)
    assert list(tmp_path.iterdir()) == []


@pytest.mark.anyio
async def test_report_runtime_preserves_structured_error_before_stderr_warning(
    tmp_path: Path,
) -> None:
    del tmp_path

    class Workspace:
        async def awrite_bytes(self, *_args: object) -> None:
            return None

        async def arun_command(self, *_args: object, **_kwargs: object) -> str:
            return '{"error":"Markdown 正式章节标识与已批准提纲不一致"}\n{"__reportExitCode":1}'

        async def adelete_file(self, *_args: object) -> None:
            return None

    report_workspace = WorkspaceReportService(Workspace())  # type: ignore[arg-type]

    with pytest.raises(
        WorkspaceError,
        match="Markdown 正式章节标识与已批准提纲不一致",
    ):
        await report_workspace._run_report_runtime(
            "validate_pdf",
            {"job": {}},
            RunContext(run_id="report-runtime-run", session_id="report-runtime-package"),
        )


def test_normalize_cjk_strong_markers_supports_chinese_punctuation() -> None:
    from markdown_it import MarkdownIt

    markdown = "呈**“年初低位—3月跳升”**形态"

    normalized = _normalize_cjk_strong_markers(markdown)
    rendered = _strip_strong_boundaries(
        MarkdownIt("commonmark", {"html": False}).render(normalized)
    )

    assert normalized == "呈\u200a**“年初低位—3月跳升”**\u200a形态"
    assert "<p>呈<strong>“年初低位—3月跳升”</strong>形态</p>" in rendered
    assert "**" not in rendered


@pytest.mark.parametrize(
    ("markdown", "expected"),
    [
        ("收入**增长**明显", "<p>收入<strong>增长</strong>明显</p>"),
        ("总收入为**1,234.5万元**。", "<p>总收入为<strong>1,234.5万元</strong>。</p>"),
        ("门诊**（含急诊）**人次", "<p>门诊<strong>（含急诊）</strong>人次</p>"),
        ("由**总部院区**、分院", "<p>由<strong>总部院区</strong>、分院</p>"),
        (
            "| 指标 | 说明 |\n| --- | --- |\n| 收入 | 同比**“双升”**态势 |",
            "<td>同比<strong>“双升”</strong>态势</td>",
        ),
    ],
)
def test_cjk_strong_markers_render_without_visible_spaces(markdown: str, expected: str) -> None:
    from markdown_it import MarkdownIt

    rendered = _strip_strong_boundaries(
        MarkdownIt("commonmark", {"html": False})
        .enable("table")
        .render(_normalize_cjk_strong_markers(markdown))
    )

    assert expected in rendered
    assert "\u200a" not in rendered
    assert "**" not in rendered


@pytest.mark.parametrize(
    ("heading", "title"),
    [
        ("### 1.1 门诊**收入**分析", "门诊收入分析"),
        ("### 1.1 门诊** 收入 **分析", "门诊** 收入 **分析"),
        ("### 1.1 门诊收入（**同比**）", "门诊收入（同比）"),
        ("### 1.1 呈**“双升”**态势", "呈**“双升”**态势"),
    ],
)
def test_cjk_strong_markers_keep_heading_text_bound_to_draft_contract(
    heading: str, title: str
) -> None:
    """标题锚点绑定使用草稿装配时的标题文本；规范粗体不能让 PDF 渲染失败。"""

    from markdown_it import MarkdownIt

    from smart_reporting.reporting.delivery.draft_v1 import _inline_heading_text

    # 草稿装配先做粗体内侧空白规范，再按 CommonMark 文本生成标题契约。
    draft_title = _inline_heading_text(
        normalize_report_markdown_strong_spacing(heading.split(" ", 2)[2])
    )
    assert draft_title == title
    markdown = f"# 报告\n\n## 1. 概述\n\n{heading}\n\n正文**增长**明显。\n"
    tokens = MarkdownIt("commonmark", {"html": False}).parse(
        _normalize_cjk_strong_markers(markdown)
    )

    _bind_heading_anchors(
        tokens,
        [
            {"level": 2, "number": "1", "title": "概述", "anchor": "report-section-overview"},
            {"level": 3, "number": "1.1", "title": draft_title, "anchor": "report-heading-a"},
        ],
    )


def test_normalize_cjk_strong_markers_preserves_unmatched_stars() -> None:
    markdown = "数量增长 * 2，备注 **未闭合"

    assert _normalize_cjk_strong_markers(markdown) == markdown


def test_normalize_spaced_strong_markers_renders_budget_values() -> None:
    from markdown_it import MarkdownIt

    markdown = "预算为 ** 131.73 亿元 **，执行率 ** 95.14% **。"

    normalized = _normalize_cjk_strong_markers(markdown)
    rendered = MarkdownIt("commonmark", {"html": False}).render(normalized)

    assert normalized == "预算为 **131.73 亿元**，执行率 **95.14%**。"
    assert "<strong>131.73 亿元</strong>" in rendered
    assert "<strong>95.14%</strong>" in rendered
    assert "**" not in rendered


@pytest.mark.parametrize(
    ("source", "expected"),
    [
        ("由** 总部院区**、", "由**总部院区**、"),
        ("由**总部院区 **、", "由**总部院区**、"),
        ("由** 总部院区 **、", "由**总部院区**、"),
        ("由**总部院区**、", "由**总部院区**、"),
        ("** 总部院区**", "**总部院区**"),
    ],
)
def test_normalize_spaced_chinese_strong_markers(source: str, expected: str) -> None:
    from markdown_it import MarkdownIt

    normalized = normalize_report_markdown_strong_spacing(source)
    rendered = MarkdownIt("commonmark", {"html": False}).render(normalized)

    assert normalized == expected
    assert "<strong>总部院区</strong>" in rendered
    assert "**" not in rendered


@pytest.mark.parametrize(
    "markdown",
    [
        "说明 `由** 总部院区**、`。",
        "```markdown\n** 总部院区**\n```",
        "~~~~\n** 总部院区**\n~~~~",
        "    ** 总部院区**",
    ],
)
def test_normalize_strong_markers_does_not_change_code(markdown: str) -> None:
    assert normalize_report_markdown_strong_spacing(markdown) == markdown
    assert _normalize_cjk_strong_markers(markdown) == markdown


def test_normalize_strong_markers_preserves_math_stars() -> None:
    markdown = "幂运算 2 ** 3，未格式化 ** 2 **。"

    assert _normalize_cjk_strong_markers(markdown) == markdown


@pytest.mark.parametrize(
    ("physical_page", "body_start_page", "physical_page_count", "expected"),
    [
        (2, 3, 25, ("i", "i")),
        (2, 4, 25, ("i", "ii")),
        (3, 4, 25, ("ii", "ii")),
        (3, 3, 25, (1, 23)),
        (25, 3, 25, (23, 23)),
    ],
)
def test_page_number_context_uses_section_page_count(
    physical_page: int,
    body_start_page: int,
    physical_page_count: int,
    expected: tuple[str | int, str | int],
) -> None:
    assert (
        _page_number_context(
            physical_page,
            body_start_page=body_start_page,
            physical_page_count=physical_page_count,
        )
        == expected
    )


def _docx_postprocess_fixture(docx: Any) -> Any:
    document = docx.Document()
    for text in (
        "测试报告",
        "__REPORT_COVER_END__",
        "目录",
        "__REPORT_TOC_FIELD_START__",
        "__REPORT_TOC_FIELD_END__",
        "__REPORT_TOC_END__",
        "__REPORT_BODY_START__",
        "经营概览",
        "测试机构",
        "2026-08-17",
    ):
        document.add_paragraph(text)
    return document


def _docx_postprocess_context() -> dict[str, Any]:
    return {
        "title": "测试报告",
        "periodLabel": "2025 年",
        "organizationName": "测试机构",
        "generatedByLabel": "Reporting Agent",
        "watermarkText": "内部资料",
        "generatedDate": "2026-08-17",
        "sections": [{"code": "overview", "title": "经营概览"}],
        "headingNumbers": [],
    }


def test_postprocess_docx_uses_section_page_count_field(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")

    path = tmp_path / "report.docx"
    document = _docx_postprocess_fixture(docx)
    document.save(path)

    _postprocess_docx(
        path,
        context=_docx_postprocess_context(),
        layout=DEFAULT_PAGE_LAYOUT,
    )

    with zipfile.ZipFile(path) as package:
        footer_xml = "\n".join(
            package.read(name).decode("utf-8")
            for name in package.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )

    # 总页数引用本节末尾书签（LibreOffice 不可靠刷新 SECTIONPAGES）。
    assert "PAGEREF report_section_end_" in footer_xml
    assert "NUMPAGES" not in footer_xml


def test_postprocess_docx_scales_tall_image_within_page_bounds(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    from docx.shared import Mm
    from PIL import Image

    image_path = tmp_path / "tall.png"
    Image.new("RGB", (600, 1800), "white").save(image_path)
    path = tmp_path / "report.docx"
    document = _docx_postprocess_fixture(docx)
    document.add_picture(str(image_path))
    document.save(path)

    _postprocess_docx(
        path,
        context=_docx_postprocess_context(),
        layout=DEFAULT_PAGE_LAYOUT,
    )

    processed = docx.Document(path)
    assert len(processed.inline_shapes) == 1
    shape = processed.inline_shapes[0]
    assert shape.width <= Mm(174)
    assert shape.height <= Mm(180)
    assert abs(shape.width / shape.height - 1 / 3) < 0.001


def test_postprocess_docx_heading_styles_follow_pdf_hierarchy(tmp_path: Path) -> None:
    # 章节、小标题、四级标题字号逐级递减并加粗，与 PDF 一致；标题前留白并与下文同页。
    docx = pytest.importorskip("docx")
    from docx.shared import Pt

    path = tmp_path / "report.docx"
    _docx_postprocess_fixture(docx).save(path)
    _postprocess_docx(path, context=_docx_postprocess_context(), layout=DEFAULT_PAGE_LAYOUT)

    styles = docx.Document(path).styles
    sizes = [styles[name].font.size for name in ("Heading 1", "Heading 2", "Heading 3", "Normal")]
    assert sizes == [Pt(16), Pt(12.5), Pt(11), Pt(10.5)]
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        assert styles[name].font.bold is True
        assert styles[name].paragraph_format.space_before >= Pt(8)
        assert styles[name].paragraph_format.keep_with_next is True


def test_postprocess_docx_tables_fill_text_width_and_quotes_are_boxed(tmp_path: Path) -> None:
    # 表格与 PDF 一致铺满版心；引用块有左侧竖线与底纹。子元素按 OOXML 架构顺序插入。
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn

    path = tmp_path / "report.docx"
    document = _docx_postprocess_fixture(docx)
    document.add_table(rows=2, cols=3)
    if "Block Text" not in document.styles:
        from docx.enum.style import WD_STYLE_TYPE

        document.styles.add_style("Block Text", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("注：本节数值均来自冻结分析结果。", style="Block Text")
    document.save(path)
    _postprocess_docx(path, context=_docx_postprocess_context(), layout=DEFAULT_PAGE_LAYOUT)

    processed = docx.Document(path)
    table_properties = processed.tables[0]._tbl.tblPr
    width = table_properties.find(qn("w:tblW"))
    assert (width.get(qn("w:type")), width.get(qn("w:w"))) == ("pct", "5000")
    order = [child.tag for child in table_properties]
    assert order.index(qn("w:tblW")) < order.index(qn("w:tblBorders"))
    if qn("w:tblLook") in order:
        assert order.index(qn("w:tblBorders")) < order.index(qn("w:tblLook"))
    block = processed.styles["Block Text"]._element.pPr
    tags = [child.tag for child in block]
    assert qn("w:pBdr") in tags and qn("w:shd") in tags
    assert tags.index(qn("w:pBdr")) < tags.index(qn("w:shd"))


def test_postprocess_docx_source_appendix_entries_and_signature_match_pdf(tmp_path: Path) -> None:
    # 来源附录条目字段缩进、行距紧凑，落款与正文同字号。
    docx = pytest.importorskip("docx")
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Mm, Pt

    path = tmp_path / "report.docx"
    document = _docx_postprocess_fixture(docx)
    for name in ("Definition Term", "Definition"):
        if name not in document.styles:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    document.save(path)
    _postprocess_docx(path, context=_docx_postprocess_context(), layout=DEFAULT_PAGE_LAYOUT)

    processed = docx.Document(path)
    styles = processed.styles
    assert styles["Report Signature"].font.size == Pt(10.5)
    signature = [item for item in processed.paragraphs if item.style.name == "Report Signature"]
    assert [item.text for item in signature] == ["测试机构", "2026-08-17"]
    assert abs(signature[0].paragraph_format.space_before - Mm(12)) < Mm(0.1)
    assert signature[1].paragraph_format.space_before is None
    # 落款不单独占一页：末段与机构名都随下一段换页。
    position = next(index for index, item in enumerate(processed.paragraphs) if item._p is signature[0]._p)
    assert processed.paragraphs[position - 1].paragraph_format.keep_with_next is True
    assert signature[0].paragraph_format.keep_with_next is True
    assert styles["Definition Term"].font.bold is True
    assert abs(styles["Definition"].paragraph_format.left_indent - Mm(6)) < Mm(0.1)
    assert styles["Definition"].paragraph_format.space_after == Pt(2)


def test_word_page_fields_use_section_page_count() -> None:
    # “pages” 改由分节末尾书签的 PAGEREF 生成，不再使用 SECTIONPAGES。
    assert _WORD_PAGE_FIELDS == {"page": "PAGE"}


def _png(path: Path) -> None:
    from PIL import Image

    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (4, 4), "white").save(path)


def _image_tokens(markdown: str) -> list[Any]:
    from markdown_it import MarkdownIt

    return MarkdownIt("commonmark").parse(markdown)


def _draft_state(*paths: str) -> dict[str, Any]:
    return {"render": {"images": [{"path": path} for path in paths]}}


def test_draft_images_fall_back_to_render_manifest(tmp_path: Path) -> None:
    _png(tmp_path / "reports/r1/chart.png")
    draft = tmp_path / "reports/r1/revision-2/draft/report.md"
    runtime = runtime_module.ReportRuntime(tmp_path)

    images = runtime._images(
        draft, _image_tokens("![图](chart.png)"), _draft_state("reports/r1/chart.png")
    )

    assert images == {(tmp_path / "reports/r1/chart.png").resolve()}
    body = '<p><img src="chart.png" alt="图"></p>'
    inlined = runtime._inline_images(body, draft.parent, images, runtime._image_sources)
    assert "data:image/png;base64," in inlined


def test_manifest_fallback_rejects_paths_outside_workspace(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    (workspace / "reports/r1/draft").mkdir(parents=True)
    _png(tmp_path / "outside/chart.png")
    draft = workspace / "reports/r1/draft/report.md"
    runtime = runtime_module.ReportRuntime(workspace)

    with pytest.raises(ReportFailure, match="不存在或格式不受支持"):
        runtime._images(
            draft, _image_tokens("![图](chart.png)"), _draft_state("../outside/chart.png")
        )


def test_manifest_fallback_prefers_ancestor_and_rejects_ambiguity(tmp_path: Path) -> None:
    _png(tmp_path / "reports/r1/chart.png")
    _png(tmp_path / "reports/r2/chart.png")
    runtime = runtime_module.ReportRuntime(tmp_path)
    state = _draft_state("reports/r2/chart.png", "reports/r1/chart.png")

    images = runtime._images(
        tmp_path / "reports/r1/revision-2/draft/report.md",
        _image_tokens("![图](chart.png)"),
        state,
    )
    assert images == {(tmp_path / "reports/r1/chart.png").resolve()}

    with pytest.raises(ReportFailure, match="多个同名候选"):
        runtime._images(
            tmp_path / "other/draft/report.md", _image_tokens("![图](chart.png)"), state
        )


@pytest.mark.parametrize(
    ("physical_page", "body_start_page", "expected"),
    [(1, 1, (1, 5)), (5, 1, (5, 5)), (1, 2, ("i", "i")), (2, 2, (1, 4))],
)
def test_page_number_context_without_cover_numbers_from_first_page(
    physical_page: int, body_start_page: int, expected: tuple[str | int, str | int]
) -> None:
    assert (
        _page_number_context(
            physical_page,
            body_start_page=body_start_page,
            physical_page_count=5,
            first_numbered_page=1,
        )
        == expected
    )


def test_pdf_decorations_without_cover_cover_every_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import pypdf

    path = tmp_path / "report.pdf"
    writer = pypdf.PdfWriter()
    for _index in range(2):
        writer.add_blank_page(width=595.276, height=841.89)
    with path.open("wb") as stream:
        writer.write(stream)
    monkeypatch.setattr(
        "smart_reporting.reporting.delivery.report_runtime.pdf._pdf_section_pages",
        lambda _reader, _sections: {"section_001": 1},
    )

    _apply_pdf_page_decorations(
        path,
        context={
            "title": "Report",
            "organizationName": "Org",
            "watermarkText": "Internal",
            "sections": [{"code": "section_001"}],
        },
        layout=DEFAULT_PAGE_LAYOUT,
        include_cover=False,
    )

    texts = [page.extract_text() or "" for page in pypdf.PdfReader(str(path)).pages]
    assert all("Internal" in text for text in texts)


@pytest.mark.parametrize(
    ("include_cover", "include_toc"), [(False, True), (True, False), (False, False)]
)
def test_postprocess_docx_follows_cover_and_toc_export_settings(
    tmp_path: Path, include_cover: bool, include_toc: bool
) -> None:
    docx = pytest.importorskip("docx")
    from smart_reporting.reporting.delivery.report_runtime.markdown import _WORD_MARKERS

    context = {
        **_docx_postprocess_context(),
        "headingNumbers": [
            {
                "level": 2,
                "number": "1",
                "title": "经营概览",
                "sectionCode": "overview",
                "anchor": "report-heading-overview",
            }
        ],
    }
    document = docx.Document()
    texts = [
        *(["测试报告"] if include_cover else []),
        _WORD_MARKERS["cover_end"],
        *(["目录"] if include_toc else []),
        _WORD_MARKERS["toc_field_start"],
        *(["1. 经营概览"] if include_toc else []),
        _WORD_MARKERS["toc_field_end"],
        _WORD_MARKERS["toc_end"],
        _WORD_MARKERS["body_start"],
        *([] if include_cover else ["测试报告"]),
        "1. 经营概览",
        "测试机构",
        "2026-08-17",
    ]
    for text in texts:
        document.add_paragraph(text)
    path = tmp_path / "report.docx"
    document.save(path)

    _postprocess_docx(
        path,
        context=context,
        layout=DEFAULT_PAGE_LAYOUT,
        include_cover=include_cover,
        include_toc=include_toc,
    )

    processed = docx.Document(path)
    assert len(processed.sections) == 1 + include_cover + include_toc
    paragraph_texts = [item.text for item in processed.paragraphs]
    assert not any(marker in paragraph_texts for marker in _WORD_MARKERS.values())
    # 无封面时首个分节即带页眉页脚；有封面时封面分节不带页面元素。
    first_header = processed.sections[0].header
    assert first_header.is_linked_to_previous is include_cover
    with zipfile.ZipFile(path) as package:
        document_xml = package.read("word/document.xml").decode()
    assert ('TOC \\o "1-3"' in document_xml) is include_toc


def test_semantic_documents_without_cover_keep_title_in_body() -> None:
    from smart_reporting.reporting.delivery.report_runtime.markdown import (
        _WORD_MARKERS,
        _semantic_documents,
    )

    context = {
        **_docx_postprocess_context(),
        "headingNumbers": [
            {
                "level": 2,
                "number": "1",
                "title": "经营概览",
                "sectionCode": "overview",
                "anchor": "report-heading-overview",
            }
        ],
    }
    pdf_document, word_document = _semantic_documents(
        "<h2>1 经营概览</h2>",
        context=context,
        layout=DEFAULT_PAGE_LAYOUT,
        include_cover=False,
        include_toc=False,
    )

    assert 'class="report-cover"' not in pdf_document
    assert 'class="report-toc"' not in pdf_document
    assert '<h1 class="report-title">测试报告</h1>' in pdf_document
    assert all(marker in word_document for marker in _WORD_MARKERS.values())
    assert "<h1>目录</h1>" not in word_document
    body = word_document.split(_WORD_MARKERS["body_start"], 1)[1]
    assert "<h1>测试报告</h1>" in body


@pytest.mark.anyio
@pytest.mark.parametrize(
    ("output", "message"),
    [
        ("Error: command timed out after 600 seconds", "报表运行时执行超时"),
        ("Error (exit 137): Killed", "报表运行时进程异常退出"),
    ],
)
async def test_report_runtime_classifies_process_failures(output: str, message: str) -> None:
    class Workspace:
        async def awrite_bytes(self, *_args: object) -> None:
            return None

        async def arun_command(self, *_args: object, **_kwargs: object) -> str:
            return output

        async def adelete_file(self, *_args: object) -> None:
            return None

    with pytest.raises(WorkspaceError, match=message):
        await WorkspaceReportService(Workspace())._run_report_runtime(  # type: ignore[arg-type]
            "render_markdown",
            {"job": {}},
            RunContext(run_id="report-runtime-run", session_id="report-runtime-package"),
        )


def test_docx_usable_width_falls_back_when_template_section_lacks_page_setup() -> None:
    from types import SimpleNamespace

    from docx.shared import Mm

    from smart_reporting.reporting.delivery.report_runtime.docx import _usable_width

    bare = SimpleNamespace(page_width=None, left_margin=None, right_margin=None)
    configured = SimpleNamespace(page_width=Mm(210), left_margin=Mm(20), right_margin=Mm(20))

    assert _usable_width(bare) == Mm(210) - Mm(31.8) - Mm(31.8)
    assert _usable_width(configured) == Mm(170)


@pytest.mark.parametrize(("markdown", "expected"), [
    ("收入增长5.2%,成本下降。", "收入增长5.2%，成本下降。"),
    ("收入, 成本均有变化;详见下表。", "收入，成本均有变化；详见下表。"),
    ("门诊收入:1,234元。", "门诊收入：1,234元。"),
    # 千分位、时间、英文、协议标记、链接与代码保持原样。
    ("收入1,234,567元,时间12:30。", "收入1,234,567元，时间12:30。"),
    ("采用 Plan A, Plan B 对比。", "采用 Plan A, Plan B 对比。"),
    ("收入稳定[[citation:cite_001]],详见附录。", "收入稳定[[citation:cite_001]]，详见附录。"),
    ("见[说明](https://example.com/a,b:c)。", "见[说明](https://example.com/a,b:c)。"),
    ("配置 `a,b:c` 中文。", "配置 `a,b:c` 中文。"),
    ("```\n中文,代码:不变\n```", "```\n中文,代码:不变\n```"),
    # 含中文或紧跟中文/全角标点的成对半角括号改为全角；英文、链接地址与代码不变。
    ("收入结构稳定(门诊占比约45%)。", "收入结构稳定（门诊占比约45%）。"),
    ("包括(1)门诊;(2)住院", "包括（1）门诊；（2）住院"),
    ("| 收入(万元) | 1,234 |", "| 收入（万元） | 1,234 |"),
    ("函数f(x)与 A (B) C 不变。", "函数f(x)与 A (B) C 不变。"),
    ("见[收入(万元)](https://example.com/x(1))。", "见[收入（万元）](https://example.com/x(1))。"),
    ("配置 `f(收入)` 不变。", "配置 `f(收入)` 不变。"),
    # 全角标点两侧的空格去掉；列表标记后与表格竖线两侧的空格不动。
    ("见附表 (表1)，（1） 门诊增长**明显** 。", "见附表（表1），（1）门诊增长**明显**。"),
    ("- （1）门诊\n* （2）住院\n1. （3）医技", "- （1）门诊\n* （2）住院\n1. （3）医技"),
    ("| 收入（万元） | 1,234 |", "| 收入（万元） | 1,234 |"),
    # 中文语境的半角引号改为中文引号；英文引语、撇号、代码、链接标题与协议标记不变。
    ('所谓"门诊收入"指诊疗收入，称为"DRG"付费，\'门诊\'口径。', "所谓“门诊收入”指诊疗收入，称为“DRG”付费，‘门诊’口径。"),
    ('He said "hello", don\'t stop.', 'He said "hello", don\'t stop.'),
    ('见`"收入"`与[说明](a.md "门诊收入")。', '见`"收入"`与[说明](a.md "门诊收入")。'),
    ('[1]: https://example.com/a "门诊收入"', '[1]: https://example.com/a "门诊收入"'),
    # 中文语境的省略号写作“……”；英文与数字中的点不变。
    ("包括内科、外科等...，以及医技等…。", "包括内科、外科等……，以及医技等……。"),
    ("wait... ok，区间1...3", "wait... ok，区间1...3"),
    # 全角括号、粗体收尾或数字之后的半角标点，只要处于中文语境同样改写。
    ("详见附表(表1),同比持平", "详见附表（表1），同比持平"),
    ("收入(元):123", "收入（元）：123"),
    ("**门诊量(万人次)**:35.3", "**门诊量（万人次）**：35.3"),
    ("会议时间10:30,地点会议室", "会议时间10:30，地点会议室"),
    ("门诊与住院比为3:1,较上年持平", "门诊与住院比为3:1，较上年持平"),
    ("(1)门诊量增长", "（1）门诊量增长"),
    ("函数f(x)为正,Q1(2025)", "函数f(x)为正，Q1(2025)"),
    # 数值占位是机器文本。
    ("收入{{value:fact-0123456789abcdef:total:亿元}}", "收入{{value:fact-0123456789abcdef:total:亿元}}"),
    # 中文句末的半角句号、问号、叹号改为全角（含百分号、粗体收尾、数值占位与表格单元格之后）。
    ("收入1.37亿元.占比30%. 是否持续?需观察!", "收入1.37亿元。占比30%。是否持续？需观察！"),
    ("**收入增长**.\n\n收入{{value:fact-0123456789abcdef:total:亿元}}.", "**收入增长**。\n\n收入{{value:fact-0123456789abcdef:total:亿元}}。"),
    ("| 收入增长. | 1.5 |", "| 收入增长。 | 1.5 |"),
    # 条目编号、小数、文件名、省略号、图片语法与英文句子不变。
    ("一.收入分析\n第二.成本见报告.pdf与v1.2版", "一.收入分析\n第二.成本见报告.pdf与v1.2版"),
    ("如图![趋势](a.png)\nGrowth was 5%.", "如图![趋势](a.png)\nGrowth was 5%."),
])
def test_cjk_punctuation_is_normalized_only_in_chinese_prose(markdown, expected):
    from smart_reporting.reporting.delivery.report_runtime.markdown import normalize_cjk_punctuation

    assert normalize_cjk_punctuation(markdown) == expected



@pytest.mark.parametrize(("markdown", "expected"), [
    # 模型照抄数值目录的 ISO 期间：改为中文日期，同年区间省略后一个年份。
    ("2025-01至2025-12期间，门诊收入稳定增长。", "2025年1月至12月期间，门诊收入稳定增长。"),
    ("报告期为2025-01-01至2025-12-31。", "报告期为2025年1月1日至12月31日。"),
    ("2024-07至2025-06跨年度对比。", "2024年7月至2025年6月跨年度对比。"),
    ("2025-01~2025-06收入增长。", "2025年1月至6月收入增长。"),
    ("**2025-03**门诊量最高。", "**2025年3月**门诊量最高。"),
    ("| 2025-01 | 123 |", "| 2025年1月 | 123 |"),
    ("2025年01月最高，2025年3月05日最低。", "2025年1月最高，2025年3月5日最低。"),
    # 文件名、地址、编号、代码、协议标记、数值占位与非日期保持原样。
    ("见 report-2025-01.pdf 与 https://x.com/2025-01/a。", "见 report-2025-01.pdf 与 https://x.com/2025-01/a。"),
    ("版本v2025-01与编号A2025-01不变。", "版本v2025-01与编号A2025-01不变。"),
    ("配置 `2025-01` 不变。", "配置 `2025-01` 不变。"),
    ("[[analysis:2025-01]]", "[[analysis:2025-01]]"),
    ("2024-2025年度与2025-13不变。", "2024-2025年度与2025-13不变。"),
    # 数值目录的比较口径字段值改为中文；英文单词、字段名、标记与代码不变。
    ("收入yoy增长5.03%，MoM下降2.1%。", "收入同比增长5.03%，环比下降2.1%。"),
    ("moment、yoy_rate、[[analysis:yoy]]与`mom`不变。", "moment、yoy_rate、[[analysis:yoy]]与`mom`不变。"),
    # 数字与汉字、单位之间的空格去掉，与数值目录写法一致；英文单词两侧、表格分隔与代码不变。
    ("收入 3,600 万元，同比增长 8.6 %。", "收入3,600万元，同比增长8.6%。"),
    ("**收入** 3,600 万元，共 12 个科室。", "**收入**3,600万元，共12个科室。"),
    ("采用 Plan A 方案，CMI 为 1.05。", "采用 Plan A 方案，CMI 为1.05。"),
    ("| 内科 | 3,600 |", "| 内科 | 3,600 |"),
    ("见 `收入 3 万` 代码。", "见 `收入 3 万` 代码。"),
    # 汉字之间（含粗体两侧）的空格去掉，日期改写后留下的空格同样去掉。
    ("2025-01 收入最高，内科 门诊增长。", "2025年1月收入最高，内科门诊增长。"),
    ("**2025-01** 收入最高。", "**2025年1月**收入最高。"),
    # 全角数字、百分号与全角千分位改为半角；条目编号“1．”与半角数字的全角逗号列举不变。
    ("收入１３，６５４．３２万元，同比增长５．０３％。", "收入13,654.32万元，同比增长5.03%。"),
    ("１．床位分别为１２０，１３５，１５０张；2．增长5.03％。", "1．床位分别为120，135，150张；2．增长5.03%。"),
    # 同一年内升序的月份区间统一为“1—10月”；跨年、编号与非月份区间不变。
    ("2025年1-10月收入，1~3月、1月-10月与2025年1月～3月。", "2025年1—10月收入，1—3月、1—10月与2025年1—3月。"),
    ("12-1月、ICD-10月度、v1-2月、3-5人与第1-3名不变。", "12-1月、ICD-10月度、v1-2月、3-5人与第1-3名不变。"),
    # 数字前的 Unicode 减号、全角正负号改为半角；不在数字前的破折号不变。
    ("同比−3.00%，变化额－195.00万元，同比＋11.00%。", "同比-3.00%，变化额-195.00万元，同比+11.00%。"),
    ("收入结构−门诊为主，见附表－说明。", "收入结构−门诊为主，见附表－说明。"),
    # 百分点缩写 pp/ppt/pct 改为“个百分点”；标识符、行内代码与英文单词不变。
    ("占比提高2pp，提升1.5个pct，下降0.8 PPT。", "占比提高2个百分点，提升1.5个百分点，下降0.8个百分点。"),
    ("v2pp、fact.3pct、`2pp`与2 apps不变。", "v2pp、fact.3pct、`2pp`与2 apps不变。"),
])
def test_iso_dates_in_prose_are_written_in_chinese(markdown, expected):
    from smart_reporting.reporting.delivery.report_runtime.markdown import normalize_cjk_wording

    assert normalize_cjk_wording(markdown) == expected

@pytest.mark.parametrize(("markdown", "expected"), [
    # 单个换行分隔的中文条目会并成一个段落、渲染成一行连写：各自成段。
    ("主要发现如下：\n1、门诊收入增长；\n2、住院收入下降。", "主要发现如下：\n\n1、门诊收入增长；\n\n2、住院收入下降。"),
    ("（1）门诊收入增长；\n（2）住院收入下降。", "（1）门诊收入增长；\n\n（2）住院收入下降。"),
    ("一是门诊增长；\n二是住院下降。", "一是门诊增长；\n\n二是住院下降。"),
    ("第一，收入增长。\n第二，成本下降。", "第一，收入增长。\n\n第二，成本下降。"),
    ("**1、门诊**增长；\n**2、住院**下降。", "**1、门诊**增长；\n\n**2、住院**下降。"),
    # Markdown 列表、表格、代码块、已分段条目与普通换行保持原样。
    ("发现：\n1. 门诊增长；\n2. 住院下降。", "发现：\n1. 门诊增长；\n2. 住院下降。"),
    ("| 科室 | 说明 |\n| --- | --- |\n| 1、内科 | 增长 |", "| 科室 | 说明 |\n| --- | --- |\n| 1、内科 | 增长 |"),
    ("```\n说明\n1、代码\n```", "```\n说明\n1、代码\n```"),
    ("发现：\n\n1、门诊", "发现：\n\n1、门诊"),
    ("第一行\n第二行", "第一行\n第二行"),
])
def test_cjk_enumerated_lines_become_separate_paragraphs(markdown, expected):
    from smart_reporting.reporting.delivery.report_runtime.markdown import separate_enumerated_lines

    assert separate_enumerated_lines(markdown) == expected


@pytest.mark.parametrize("fence", ["````", "~~~~"])
@pytest.mark.parametrize("interior", ["short", "info", "other", "indented"])
@pytest.mark.parametrize("closing", ["same", "longer"])
def test_cjk_enumeration_preserves_fence_contents_and_resumes_after_close(fence, interior, closing):
    from smart_reporting.reporting.delivery.report_runtime.markdown import (
        _normalize_report_markdown_segments,
        separate_enumerated_lines,
    )

    candidates = {
        "short": fence[:-1], "info": fence + "python",
        "other": "~~~~" if fence[0] == "`" else "````", "indented": "    " + fence,
    }
    close = "   " + fence + (fence[0] if closing == "longer" else "") + " \t"
    code = f"{fence}\n{candidates[interior]}\n一、literal\n{close}"
    markdown = code + "\n正文\n一、条目"
    assert separate_enumerated_lines(markdown) == code + "\n正文\n\n一、条目"
    for protected in (code, code.replace("\n", "\r\n")):
        assert _normalize_report_markdown_segments(
            protected, lambda line: line.replace("literal", "changed"),
        ) == protected



@pytest.mark.parametrize(("markdown", "expected"), [
    # WeasyPrint 把中文之间的段内换行渲染成空格：两侧均为宽字符时去掉换行。
    ("门诊收入持续增长\n住院收入有所回落", "<p>门诊收入持续增长住院收入有所回落</p>\n"),
    ("收入增长，\n住院回落", "<p>收入增长，住院回落</p>\n"),
    ("**门诊收入**\n住院回落", "<p><strong>门诊收入</strong>住院回落</p>\n"),
    # 英文、数字或代码一侧保留换行（渲染为空格）；代码块不变。
    ("Revenue grew\nslightly", "<p>Revenue grew\nslightly</p>\n"),
    ("收入增长5%\n住院回落", "<p>收入增长5%\n住院回落</p>\n"),
    ("门诊`code`\n住院", "<p>门诊<code>code</code>\n住院</p>\n"),
    ("```\n说明\n代码\n```", "<pre><code>说明\n代码\n</code></pre>\n"),
])
def test_cjk_soft_breaks_render_without_spaces(markdown, expected):
    from markdown_it import MarkdownIt

    from smart_reporting.reporting.delivery.report_runtime.markdown import _join_cjk_soft_breaks

    parser = MarkdownIt("commonmark", {"html": False}).enable("table")
    assert parser.renderer.render(_join_cjk_soft_breaks(parser.parse(markdown)), parser.options, {}) == expected


@pytest.mark.parametrize(("title", "expected"), [
    # 回放：按宽度自然折行剩下单字孤行“…运营分析报 / 告”。
    ("2025年瑞金医院综合运营分析报告", "2025年瑞金医院<br>综合运营分析报告"),
    ("运营报告", "运营报告"),
    ("2025年瑞金医院综合运营与财务预算执行情况分析报告", "2025年瑞金医院综合运营<br>与财务预算执行情况分析报告"),
    ("<脚本>报告名称超过十五个字的测试标题啊", "&lt;脚本&gt;报告名称超过<br>十五个字的测试标题啊"),
])
def test_cover_title_breaks_into_balanced_lines(title, expected):
    from smart_reporting.reporting.delivery.report_runtime.markdown import _balanced_cover_title

    assert _balanced_cover_title(title) == expected


def test_server_tables_move_after_body_with_their_captions_only():
    from smart_reporting.reporting.delivery.report_runtime.markdown import (
        APPENDIX_TABLES_SENTINEL,
        relocate_server_tables,
    )

    table = "[[table:{0}]]\n| | 金额 |\n| --- | ---: |\n| 1月 | 1 |\n\n[[/table:{0}]]"
    markdown = (
        "## 1. 收入\n\n正文结论。\n\n2025年收入：分期间指标汇总\n\n" + table.format("t1")
        + "\n\n## 2. 成本\n\n次均收入120元[[claim:c]]。\n\n" + table.format("t2") + "\n\n| 模型表 |\n| --- |\n| 留在正文 |\n"
    )
    body, appendix = relocate_server_tables(markdown).split(APPENDIX_TABLES_SENTINEL)
    assert "[[table:" not in body and "分期间指标汇总" not in body
    assert "次均收入120元[[claim:c]]。" in body and "| 留在正文 |" in body
    assert appendix.index("分期间指标汇总") < appendix.index("[[table:t1]]") < appendix.index("[[table:t2]]")
    assert relocate_server_tables("## 1. 收入\n\n正文。\n") == "## 1. 收入\n\n正文。\n"


def test_executive_summary_takes_each_section_conclusion():
    from smart_reporting.reporting.delivery.report_runtime.markdown import (
        SUMMARY_END_SENTINEL,
        SUMMARY_SENTINEL,
        insert_executive_summary,
    )

    sections = [{"code": "a", "title": "收入预算执行情况"}, {"code": "b", "title": "业务工作量变化"},
                {"code": "c", "title": "指标统计关系说明"}]
    markdown = (
        "# 报告\n\n[[section:a]]\n\n## 1. 收入预算执行情况\n\n### 1.1 总体\n\n"
        "全院收入预算执行整体未达预算。全年完成率为79.28%[[claim:x]]，收入端滞后。\n\n"
        "[[section:b]]\n\n## 2. 业务工作量变化\n\n"
        "门诊人次3月最高，为644,217人次，10月最低。其余月份平稳。\n\n"
        "[[section:c]]\n\n## 3. 指标统计关系说明\n\n相关不等于因果，仅作观察。\n"
    )
    result = insert_executive_summary(markdown, sections)
    summary = result[result.index(SUMMARY_SENTINEL):result.index(SUMMARY_END_SENTINEL)]
    assert result.index(SUMMARY_SENTINEL) < result.index("[[section:a]]")
    assert "- **收入预算执行情况**：1.1 总体： 全院收入预算执行整体未达预算。全年完成率为79.28%，收入端滞后。" in summary
    assert "- **业务工作量变化**：门诊人次3月最高，为644,217人次，10月最低。" in summary
    assert "指标统计关系说明" not in summary and "[[" not in summary
    # 只有一章可取时不生成。
    assert insert_executive_summary(markdown, sections[:1]) == markdown


@pytest.mark.parametrize("qualification", [
    "收入归属和成本分摊口径尚未匹配，暂不能据此评价经营结余。",
    "\n\n该层级组织映射缺失率80%，不支持完整科室间比较。",
    "\n\n> 该结果只覆盖登记范围，不能外推全院。",
    "\n### 口径边界\n该分类未确认，不支持业务归因。",
    "\n\n### 预算版本未核实，不支持完成率评价",
])
def test_summary_keeps_same_and_later_paragraph_qualifications(qualification):
    from smart_reporting.reporting.delivery.report_runtime.markdown import _section_conclusion

    result = _section_conclusion("## 收支\n\n收支差额为100万元。" + qualification)
    assert "100万元" in result
    assert qualification.splitlines()[-1].lstrip("> #") in result


def test_summary_omits_whole_chapter_if_it_cannot_preserve_its_context():
    from smart_reporting.reporting.delivery.report_runtime.markdown import _section_conclusion

    assert _section_conclusion("收入100万元。\n\n" + "完整口径限制。" * 100) is None
    assert _section_conclusion("收入100万元。\n\n| 口径限制 |\n| --- |\n| 不可比 |") is None


def test_executive_summary_keeps_qualifications_and_compacts_raw_values_and_long_titles():
    from smart_reporting.reporting.delivery.report_runtime.markdown import (
        SUMMARY_END_SENTINEL,
        SUMMARY_SENTINEL,
        insert_executive_summary,
    )

    # 回放：领导首页出现“数值待核实”、括号原值和拼接成的长章节名。
    long_title = "预算人次比较变化，也称预算门诊人次；预算总医疗收入比较变化；预算药品收入比较变化；实际人次趋势"
    sections = [{"code": "a", "title": long_title}, {"code": "b", "title": "成本结构"}]
    markdown = (
        "[[section:a]]\n\n## 1. 预算\n\n实际支出相当于预算支出的数值待核实。其中1月偏差最大。\n\n"
        "全院收入合计111.24亿元（11,123,541,503元），较上年增长。\n\n"
        "[[section:b]]\n\n## 2. 成本结构\n\n成本合计数值待核实。该口径为含负值的净额口径。\n\n人员支出占比最高，为35.32%。\n"
    )
    result = insert_executive_summary(markdown, sections)
    summary = result[result.index(SUMMARY_SENTINEL):result.index(SUMMARY_END_SENTINEL)]
    assert "- **预算人次比较变化等**：实际支出相当于预算支出的数值待核实。" in summary
    assert "全院收入合计111.24亿元，较上年增长。" in summary
    assert "- **成本结构**：成本合计数值待核实。该口径为含负值的净额口径。" in summary
    assert "人员支出占比最高，为35.32%。" in summary
    assert "11,123,541,503元" not in summary


def test_render_collapses_identical_listed_values():
    from smart_reporting.reporting.delivery.report_runtime.markdown import collapse_repeated_values

    assert collapse_repeated_values("全年合计分别为0元[[claim:a]]、0元[[claim:b]]、0元、0元。") == "全年合计均为0元[[claim:a]]。"
    assert collapse_repeated_values("分别为0元、0元、5元。") == "分别为0元、0元、5元。"
    assert collapse_repeated_values("分别为10%、10.5%。") == "分别为10%、10.5%。"
    assert collapse_repeated_values("三项分别为10%、10%和20%。") == "三项分别为10%、10%和20%。"
    assert collapse_repeated_values("分别为0元、0元及5元") == "分别为0元、0元及5元"
    assert collapse_repeated_values("分别为0元、0元，合计0元") == "均为0元，合计0元"


def test_render_transposes_tall_single_column_server_tables():
    from smart_reporting.reporting.delivery.report_runtime.markdown import (
        compact_single_column_tables,
    )

    # 回放附表：“收入金额（亿元）”逐月一行，11行只有一列数字，占半页。
    rows = "".join(f"| 2025年{month}月 | {month}.00 |\n" for month in range(1, 12))
    block = f"[[table:t1]]\n|  | 收入金额（亿元） |\n| --- | ---: |\n{rows}\n[[/table:t1]]"
    compact = compact_single_column_tables("表名\n\n" + block + "\n")
    assert compact.startswith("表名\n\n[[table:t1]]\n| 2025年 | 1月 | 2月 | 3月 | 4月 | 5月 | 6月 |")
    assert "| 收入金额（亿元） | 1.00 | 2.00 | 3.00 | 4.00 | 5.00 | 6.00 |" in compact
    assert "| 2025年 | 7月 | 8月 | 9月 | 10月 | 11月 |" in compact and compact.rstrip().endswith("[[/table:t1]]")
    # 行数少或多列的表不动。
    short = "[[table:t2]]\n|  | 收入 |\n| --- | ---: |\n| 2025年1月 | 1.00 |\n\n[[/table:t2]]"
    wide = block.replace("收入金额（亿元） |", "收入 | 成本 |").replace("---: |", "---: | ---: |")
    assert compact_single_column_tables(short) == short
    assert compact_single_column_tables(wide) == wide


def test_model_table_numeric_columns_are_right_aligned():
    from smart_reporting.reporting.delivery.report_runtime.markdown import (
        align_numeric_table_columns,
    )

    table = (
        "| 月份 | 收入总额 | 环比差额（万元） | 环比幅度 | 说明 |\n"
        "| --- | --- | --- | :---: | --- |\n"
        "| 2025年1月 | 8.90亿元 | — | — | 起始月 |\n"
        "| 2025年2月 | **8.99亿元** | -322.71 | 1.05% | 回升 |\n"
    )
    aligned = align_numeric_table_columns("正文。\n\n" + table)
    # 数值列右对齐；文本列与已声明居中的列不动。
    assert "| --- | ---: | ---: | :---: | --- |" in aligned
    assert align_numeric_table_columns("| a |\n| --- |") == "| a |\n| --- |"
    # 平均住院日、倍数、例数与带来源标记的单元格同样是数值列。
    hospital = (
        "| 科室 | 平均住院日 | 手术例数 | 收入 |\n| --- | --- | --- | --- |\n"
        "| 内科 | 8.5天 | 1,320例 | 1.20亿元[[claim:a]] |\n| 外科 | 6天 | 980例 | **0.98亿元** |\n"
    )
    assert "| --- | ---: | ---: | ---: |" in align_numeric_table_columns(hospital)


def test_pdf_tables_keep_edge_rows_together_and_units_whole() -> None:
    from markdown_it import MarkdownIt

    from smart_reporting.reporting.delivery.report_runtime.markdown import _paginate_pdf_tables

    rows = "".join(f"| 2025年{month}月 | {month} |\n" for month in range(1, 8))
    rendered = _paginate_pdf_tables(
        MarkdownIt("commonmark").enable("table").render(f"| 月份 | 实际人次（人次） |\n| --- | ---: |\n{rows}")
    )
    # 回放附表：末行“2025年12月”单独落到下一页；首尾各两行放进不可拆分的 tbody。
    assert rendered.count('<tbody class="report-table-keep">') == 2 and rendered.count("<tbody") == 3
    assert rendered.count("<tr>") + rendered.count('<tr class="report-row-even">') == 8
    # 斑马纹按全表行序标注，不随 tbody 分组重新计数。
    assert rendered.count('class="report-row-even"') == 3
    assert "实际人次<span class=\"report-unit\">（人次）</span></th>" in rendered
    short = _paginate_pdf_tables("<table>\n<tbody>\n<tr>\n<td>1</td>\n</tr>\n</tbody>\n</table>")
    assert short.count("<tbody") == 1 and 'class="report-table-keep"' in short


def test_render_shows_large_yuan_amounts_in_readable_units() -> None:
    from smart_reporting.reporting.delivery.report_runtime.markdown import readable_large_amounts

    # 回放：正文 68 处“3,046,407,434元”式全位数金额，领导与财务读者难以读出量级。
    text = (
        "药品收入3,046,407,434元[[claim:number-1]]、占27.39%；其他费用91,342,503元，"
        "差额为-7,977,183元，零星支出52,300元。预算人次7,275,108人次，门诊644,217人次。\n\n"
        "合计99.12亿元（9,911,520,144元），覆盖9个月。\n\n"
        "| 项目 | 金额 |\n| --- | --- |\n| 药品 | 3,046,407,434元 |\n\n"
        "![图：药品3,046,407,434元](chart.png)"
    )
    rendered = readable_large_amounts(text)
    assert "药品收入30.46亿元[[claim:number-1]]" in rendered
    assert "其他费用9,134.25万元" in rendered and "差额为-797.72万元" in rendered
    assert "零星支出52,300元" in rendered
    assert "预算人次727.51万人次" in rendered and "门诊644,217人次" in rendered
    assert "（9,911,520,144元）" in rendered
    assert "| 药品 | 3,046,407,434元 |" in rendered and "![图：药品3,046,407,434元]" in rendered


def test_executive_summary_strips_every_list_marker_in_a_list_paragraph():
    from smart_reporting.reporting.delivery.report_runtime.markdown import _section_conclusion

    # 审阅：列表段落先拼行再去前缀，只去掉第一项的“- ”，后续项的符号混进核心结论。
    conclusion = _section_conclusion("## 1. 收入\n\n收入总体增长。\n\n- 药品收入增长5%。\n- 材料收入下降2%。")
    assert conclusion == "收入总体增长。 药品收入增长5%。材料收入下降2%。"
    assert "- " not in conclusion


def test_word_tables_keep_edge_rows_and_header_units_together() -> None:
    from docx import Document
    from docx.oxml.ns import qn

    from smart_reporting.reporting.delivery.report_runtime.docx import (
        _break_before_header_unit,
        _keep_table_together_at_edges,
    )

    document = Document()
    document.add_paragraph("2025年1—9月收入汇总：分期间指标汇总")
    table = document.add_table(rows=8, cols=2)
    table.rows[0].cells[1].paragraphs[0].add_run("实际人次")
    table.rows[0].cells[1].paragraphs[0].add_run("（人次）")
    _keep_table_together_at_edges(table)
    _break_before_header_unit(table.rows[0].cells[1])

    # 与 PDF 一致：行不跨页拆分；表头与前两行、倒数第二行与末行同页；表名随表格。
    assert all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in table.rows)
    keeps = [bool(row.cells[0].paragraphs[0].paragraph_format.keep_with_next) for row in table.rows]
    assert keeps == [True, True, False, False, False, False, True, False]
    assert document.paragraphs[0].paragraph_format.keep_with_next
    # 表头单位整体换行：“实际人次 / （人次）”，不在括号中间折断。
    assert table.rows[0].cells[1].text == "实际人次\n（人次）"
