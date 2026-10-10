"""Editor 原版、手改版、历史恢复版的真实双格式渲染与提交验收。"""

from __future__ import annotations

import asyncio
import hashlib
import io
from pathlib import Path
from uuid import UUID

import pypdf
import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont

from smart_reporting.report_editor import ReportEditorContext
from smart_reporting.reporting.delivery.artifacts_v1 import ArtifactFile, ChartArtifact
from smart_reporting.reporting.delivery.publishing import (
    ReportArtifactPersistenceService,
    ReportDownloadGrantService,
)
from smart_reporting.reporting.delivery.report_runtime.runtime import ReportRuntime
from smart_reporting.reporting.trace.chart_subjects import freeze_chart_presentations
from smart_reporting.reporting.trace.contracts_v1 import (
    ChartTraceV1,
    TableTraceV1,
    TraceFileRefV1,
    derive_resource_id,
)
from smart_reporting.reporting.trace.index_builder import encode_trace_index
from smart_reporting.reporting.trace.table_builder import render_table_markdown
from smart_reporting.reporting.workflow.state import ReportingRunState
from smart_reporting.reporting.workspace import WorkspaceReportService

from .delivery_fakes import InMemoryDownloadGrantRepository, InMemoryReportArtifactRepository
from .lineage_fixtures.manifest import register_trace_manifest
from .test_report_editor_lineage_export import _StateRepository
from .test_report_runtime_full_render import _DEPENDENCIES_OK, _SKIP_REASON
from .test_trace_subject_validate import _make_editor_with_subject

pytestmark = [pytest.mark.integration, pytest.mark.skipif(not _DEPENDENCIES_OK, reason=_SKIP_REASON)]


def _identity(path: str, content: bytes) -> dict:
    return {"path": path, "size": len(content), "sha256": hashlib.sha256(content).hexdigest()}


@pytest.mark.anyio
async def test_editor_real_dual_export_preserves_edited_and_restored_snapshot(tmp_path: Path) -> None:
    editor, grants, context = await _make_editor_with_subject(tmp_path)
    workspace = editor.workspace
    thread = context.scope["threadId"]
    index = await editor.trace.load_index(context)
    assert index is not None
    records = [[f"记录{number:02d}", "3,600"] for number in range(1, 49)]
    table = render_table_markdown("tbl-1", ("income_total",), records)
    original = (
        '# 收入快照追溯报告\n\n'
        '## 1. Report\n\n[[section:section_002]]\n[[analysis:analysis_001]]\n\n'
        '2025-09 收入 3600 万元[[claim:claim-1]]。\n\n'
        '![收入快照](chart.png) [[citation:citation_001]]\n\n'
        '*图表：冻结收入快照*\n\n'
        '下表逐行引用同一冻结事实，不对重复展示的记录求和。\n\n' + table + '\n'
    )
    await workspace.awrite_text(thread, context.markdown_path, original, overwrite=True)
    image = Image.new("RGB", (900, 260), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=24)
    draw.text((30, 20), "Frozen revenue / CNY 10,000", fill="black", font=font)
    draw.rectangle((30, 85, 750, 155), fill="#247e91")
    draw.text((765, 105), "3600", fill="black", font=font)
    draw.text((30, 195), "2025-09", fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    chart_bytes = buffer.getvalue()
    chart_path = "reports/revision-1/chart.png"
    plot_path = "reports/revision-1/chart.chart-input.json"
    plot_bytes = b'{"schema":"chart-input/v1","chartId":"chart_001","role":"measure","columns":["period","revenue"],"rows":[["2025-09",3600]],"rowCount":1}'
    await workspace.awrite_bytes(thread, chart_path, chart_bytes)
    await workspace.awrite_bytes(thread, plot_path, plot_bytes)
    indexed_markdown = next(file for file in index.files if file.resource_id == index.markdown_file_resource_id)
    files = tuple(file.model_copy(update=_identity(file.path, original.encode())) if file == indexed_markdown else file for file in index.files)
    for path, content, media_type in ((chart_path, chart_bytes, "image/png"), (plot_path, plot_bytes, "application/json")):
        files += (TraceFileRefV1(resourceId=derive_resource_id(path), mediaType=media_type, **_identity(path, content)),)
    cells = tuple(index.tables[0].cells[0].model_copy(update={"row_key": f"row:{number:02d}"}) for number in range(1, 49))
    index = index.model_copy(update={
        "files": files,
        "tables": (TableTraceV1(tableId="tbl-1", rowKeys=tuple(cell.row_key for cell in cells), columnKeys=("income_total",), cells=cells),),
        "chart_traces": (ChartTraceV1(chartId="chart_001", imageFileResourceId=derive_resource_id(chart_path), plotDataFileResourceIds=(derive_resource_id(plot_path),), datasetIds=("dataset-url-abc0001",)),),
    })
    index = freeze_chart_presentations(index, original)
    await workspace.awrite_bytes(thread, "reports/revision-1/trace-index-v1.json", encode_trace_index(index), overwrite=True)
    await workspace.adelete_file(thread, context.artifact_manifest.path)
    registered = await register_trace_manifest(workspace, thread, index)
    context = context.model_copy(update={"artifact_manifest": registered})
    manifest = await editor.trace.load_manifest(context)
    assert manifest is not None
    chart = ChartArtifact(chartId="chart_001", mediaType="image/png", datasetIds=("dataset-url-abc0001",), **_identity(chart_path, chart_bytes))
    manifest = manifest.model_copy(update={"charts": (chart,), "analysis_ids": ("analysis_001",)})
    encoded = manifest.model_dump_json(by_alias=True).encode()
    await workspace.awrite_bytes(thread, registered.path, encoded, overwrite=True)
    context = context.model_copy(update={"artifact_manifest": ArtifactFile(mediaType="application/json", **_identity(registered.path, encoded))})
    dataset = next(file for file in index.files if file.media_type == "text/csv")
    job_id = str(UUID(int=1))
    job = {
        "jobId": job_id,
        "_threadBinding": hashlib.sha256(thread.encode()).hexdigest(),
        "sources": [dataset.model_dump(mode="json", by_alias=True, exclude={"resource_id", "media_type"})],
        "render": {"images": [_identity(chart_path, chart_bytes)]},
        "_documentContext": {
            "title": "收入快照追溯报告", "periodLabel": "2025-09", "organizationName": "快照验收机构",
            "generatedByLabel": "Reporting Agent", "generatedDate": "2026-09-30", "watermarkText": "内部资料",
            "sectionNumbers": list(manifest.section_numbers),
            "sections": [{"code": "section_002", "sectionNumber": "1", "title": "Report"}],
            "headingNumbers": [heading.model_dump(mode="json", by_alias=True) for heading in manifest.heading_numbers],
        },
        "_citationPresentations": [{"citationId": "citation_001", "label": "集团医院营业收入按月冻结快照长中文业务名称", "coverageItems": [{"label": "收入明细", "periods": ["2025-09"]}]}],
    }
    root = editor.workspace_registry.get(thread).root
    rendered = await asyncio.to_thread(ReportRuntime(root).render_markdown, job, context.markdown_path, "reports/revision-1/report.pdf", ".reporting-tmp/workspace-report-initial-real-export-render/render.pdf")
    job["render"] = {key: rendered["render"][key] for key in ("markdown", "pdf", "word", "images")}
    context = context.model_copy(update={"job_id": job_id, "job": job})
    durable = ReportingRunState.initial(report_run_id=context.workflow_run_id, external_run_id="external-1", thread_id="caller-thread", owner_user_id="user-1", payload={"reportEditorContexts": {"1": context.model_dump(mode="json", by_alias=True)}})
    repository = _StateRepository(durable, failure=None)
    artifacts = InMemoryReportArtifactRepository()
    editor.state_repository = repository
    editor.report_tools = WorkspaceReportService(workspace)
    editor.artifact_persistence = ReportArtifactPersistenceService(artifacts, workspace)
    editor.download_grants = ReportDownloadGrantService(InMemoryDownloadGrantRepository())
    editor.editor_grants = grants
    editor.public_base_url = "https://reports.example.com"
    edited = original.replace("收入 3600 万元", "收入 3800 万元").replace("冻结收入快照*", "编辑后收入快照*")
    before = await editor.read_document(context)
    await editor.save_draft(context, markdown=edited, expected_sha256=before.sha256)

    for revision, expected, source_status in ((2, edited, "stale"), (3, original, "valid")):
        if revision == 3:
            before = await editor.read_document(context)
            await editor.restore_history(context, 1, expected_sha256=before.sha256)
        document = await editor.read_document(context)
        assert document.markdown == expected
        result = await editor.export_revision(context, expected_sha256=document.sha256, settings={"cover": True, "toc": True})
        assert result["revision"] == revision
        context = ReportEditorContext.model_validate(repository.durable.payload["reportEditorContexts"][str(revision)])
        current_manifest = await editor.trace.load_manifest(context)
        current_index = await editor.trace.load_index(context)
        assert current_manifest is not None and current_index is not None
        assert current_manifest.revision == current_index.revision == context.revision == revision
        assert current_manifest.markdown.sha256 == document.sha256
        assert await workspace.aread_text(thread, context.markdown_path) == expected
        assert current_index.subject_bindings[0].subject_sha256 == index.subject_bindings[0].subject_sha256
        validation = await editor.trace.validate(context, expected, document.sha256)
        assert validation["subjects"][0]["status"] == source_status
        assert validation["charts"][0]["status"] == source_status
        assert validation["tableSummary"]["valid"] == 48
        pdf = root / f"reports/revision-{revision}/report.pdf"
        word = pdf.with_suffix(".docx")
        reader = pypdf.PdfReader(pdf)
        pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
        word_document = Document(word)
        word_text = "\n".join(paragraph.text for paragraph in word_document.paragraphs)
        for text in (pdf_text, word_text):
            assert ("收入 3800 万元" if revision == 2 else "收入 3600 万元") in text
            assert "事实值：3,600万元" in text
            assert "数据来源附录" in text
            assert ("状态：待复核" if revision == 2 else "状态：有效") in text
            citation_appendix = text.split("实际引用附录", 1)[1].split("数据来源附录", 1)[0]
            assert ("状态：待复核" if revision == 2 else "状态：有效") in citation_appendix
            if revision == 2:
                assert "状态：有效" not in citation_appendix
            assert "[[claim:" not in text and "[[table:" not in text
            assert '<a id="report-section-' not in text
            assert ("图表：编辑后收入快照" if revision == 2 else "图表：冻结收入快照") + " [数据来源 002]" in text
        assert len(reader.pages) >= 4
        data_tables = [item for item in word_document.tables if [cell.text for cell in item.rows[0].cells] == ["", "income_total"]]
        assert len(data_tables) == 1
        assert len(data_tables[0].rows) == 49
        assert len(word_document.inline_shapes) == 1
        assert "PAGEREF report_section_end_2" in word_document.sections[-1].footer._element.xml
        assert "SECTIONPAGES" not in word_document.sections[-1].footer._element.xml
        for path, kind in ((pdf, "pdf"), (word, "word")):
            stored = next(record for record in artifacts.records.values() if record.revision == revision and record.artifact == kind)
            assert stored.sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
            assert b"".join(artifacts.chunks[stored.artifact_key]) == path.read_bytes()
        expected_link = f"https://reports.example.com/reports/v1/editor/report-1/{revision}?subject=sub-" + "0" * 16
        links = [str(annotation.get_object().get("/A", {}).get("/URI", "")) for page in reader.pages for annotation in page.get("/Annots") or ()]
        assert expected_link in links
        assert expected_link in {relationship.target_ref for relationship in word_document.part.rels.values() if relationship.is_external}

    historical = await editor.read_history_revision(context, 1)
    assert historical["markdown"] == original
    assert historical["sources"]["revision"] == 1
    assert not list(root.glob(".workspace-report-runtime-*"))
    assert not list((root / "reports").glob("revision-*/.export-*"))
