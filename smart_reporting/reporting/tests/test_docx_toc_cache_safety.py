"""TOC cache write failures must preserve the actual readable Word artifact."""

from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.document import Document as DocumentClass

from smart_reporting.reporting.delivery.report_runtime import docx as runtime


@pytest.mark.parametrize('fault', ['partial_write', 'invalid_zip', 'replace'])
def test_toc_cache_failure_preserves_original_word(tmp_path, monkeypatch, fault):
    path = tmp_path / 'report.docx'
    document = Document()
    document.add_paragraph('原始报告正文')
    document.save(path)
    original = path.read_bytes()

    def convert(args, **kwargs):
        output = Path(args[args.index('--outdir') + 1])
        (output / 'report.pdf').write_bytes(b'converted')
        return SimpleNamespace(returncode=0)

    def broken_save(self, target):
        Path(target).write_bytes(b'PK-incomplete')
        if fault == 'partial_write':
            raise OSError('injected write failure')

    def broken_replace(*args):
        raise OSError('injected replace failure')

    reader = SimpleNamespace(
        named_destinations={'/ReportBodyStart': object()},
        get_destination_page_number=lambda _: 1, pages=[None, None],
    )
    monkeypatch.setattr(runtime.shutil, 'which', lambda _: '/fake/soffice')
    monkeypatch.setattr(runtime.subprocess, 'run', convert)
    monkeypatch.setattr('pypdf.PdfReader', lambda _: reader)
    monkeypatch.setattr(runtime, '_heading_body_pages', lambda *args, **kwargs: [2])
    if fault == 'replace':
        monkeypatch.setattr(runtime.os, 'replace', broken_replace)
    else:
        monkeypatch.setattr(DocumentClass, 'save', broken_save)
    runtime._cache_toc_page_numbers(path, context={'headingNumbers': [
        {'anchor': 'report-income', 'number': '1', 'level': 2, 'title': '收入'},
    ]}, include_cover=True)
    assert path.read_bytes() == original
    assert Document(path).paragraphs[0].text == '原始报告正文'
    assert not list(tmp_path.glob('.reporting-toc-*'))
