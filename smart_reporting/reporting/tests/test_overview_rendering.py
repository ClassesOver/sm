"""A source-bound matrix replaces excerpts without inheriting business approval."""

import copy
import hashlib
import json

import pytest

from smart_reporting.reporting.delivery.report_runtime.markdown import (
    SUMMARY_END_SENTINEL,
    SUMMARY_SENTINEL,
)
from smart_reporting.reporting.delivery.report_runtime.projection import (
    build_render_projection,
    verify_render_projection,
)
from smart_reporting.reporting.workflow.management_overview import management_overview_projection
from smart_reporting.reporting.workflow.overview_sources import registered_management_overview

from .test_quality_answer_review import case, evaluate, register_quality


@pytest.mark.parametrize("inline", [True, False])
def test_explicit_null_overview_never_falls_back_to_unverified_excerpts(tmp_path, inline):
    from smart_reporting.reporting.delivery.report_runtime.overview import load_overview

    raw = b"null"
    (tmp_path / "overview.json").write_bytes(raw)
    state = ({"_managementOverviewProjection": None} if inline else {
        "_managementOverviewProjectionFile": {"path": "overview.json", "size": len(raw),
                                             "sha256": hashlib.sha256(raw).hexdigest()}})
    assert load_overview({}, tmp_path) is None  # Legacy callers have no binding at all.
    payload = load_overview(state, tmp_path)
    projection = build_render_projection("[[section:income]]\n\n## 1. 收入\n\n收入999999万元。",
        [{"code": "income", "title": "收入"}], management_overview=payload)
    summary = projection.markdown.split(SUMMARY_SENTINEL, 1)[1].split(SUMMARY_END_SENTINEL)[0]
    assert "暂不列示汇总数值" in summary and "999999" not in summary
    assert projection.overview_check["status"] == "unknown"


@pytest.mark.anyio
@pytest.mark.parametrize("anyio_backend", ["asyncio"])
@pytest.mark.parametrize("fault", ["none", "source", "rule", "protocol", "citation", "nan"])
async def test_structured_projection_refuses_stale_or_malformed_overviews(tmp_path, fault):
    _, _, index, markdown, sources, questions, plan = await case(tmp_path)
    quality = evaluate(index, markdown, sources, questions, plan)
    payload = management_overview_projection(quality.management_overview, hashlib.sha256(markdown.encode()).hexdigest())
    with pytest.raises(ValueError):
        management_overview_projection(quality.management_overview, "f" * 64)
    if fault == "source":
        payload["sourceSha256"] = "0" * 64
    elif fault == "rule":
        payload["version"] = "old"
    elif fault == "protocol":
        payload["markdown"] += "[[section:fake]]"
    elif fault == "citation":
        payload["markdown"] += "[[citation:unknown]]"
    elif fault == "nan":
        payload["markdown"] = float("nan")
    sections = [{"code": "section_002", "title": "报告"}]
    projection = build_render_projection(markdown, sections, management_overview=payload)
    summary = projection.markdown.split(SUMMARY_SENTINEL, 1)[1].split(SUMMARY_END_SENTINEL)[0]
    assert projection.overview_check["status"] == ("verified" if fault == "none" else "unknown")
    assert ("2400" in summary) == (fault == "none")
    assert "该项待核实" in summary if fault == "none" else "暂不列示汇总数值" in summary
    assert verify_render_projection(projection.receipt, markdown, sections, management_overview=payload)["status"] == "verified"
    changed = copy.deepcopy(payload)
    changed["sourceSha256"] = "f" * 64
    if fault == "nan":
        changed["markdown"] = "repaired JSON content"
    assert verify_render_projection(projection.receipt, markdown, sections, management_overview=changed)["status"] == "unknown"


@pytest.mark.anyio
@pytest.mark.parametrize("anyio_backend", ["asyncio"])
@pytest.mark.parametrize("fault", ["none", "source", "authorization", "markdown", "quality", "profile"])
async def test_first_publication_checks_registered_overview_sources(tmp_path, fault):
    editor, context, index, markdown, sources, questions, plan = await case(tmp_path)
    # This fixture's legacy facts use the compiler's zero-profile fallback.
    # Register that profile explicitly, as production indexes do.
    index = index.model_copy(update={"profile_hash": "0" * 64})
    quality = evaluate(index, markdown, sources, questions, plan)
    context = await register_quality(editor, context, index, quality)
    manifest, index = await editor.trace.load_manifest_and_index(context)
    if fault != "profile":
        manifest = manifest.model_copy(update={"effective_profile_hash": index.profile_hash})
    if fault in {"source", "markdown", "quality"}:
        file = (next(file for file in index.files if file.resource_id == index.fact_files[0].file_resource_id)
                if fault == "source" else next(file for file in index.files if file.path.endswith("quality-bundle-v1.json"))
                if fault == "quality" else manifest.markdown)
        await editor.workspace.awrite_bytes(context.scope["threadId"], file.path, b"changed", overwrite=True)
    if fault == "none":
        overview = await registered_management_overview(editor.workspace, context.scope["threadId"], manifest, "c" * 64)
        assert overview.rows[0].status == "met"
    else:
        expected = ("provenance mismatch" if fault == "authorization" else "index identity mismatch"
                    if fault == "profile" else "input identity changed")
        with pytest.raises(ValueError, match=expected):
            await registered_management_overview(editor.workspace, context.scope["threadId"], manifest,
                                                 "d" * 64 if fault == "authorization" else "c" * 64)


@pytest.mark.anyio
@pytest.mark.parametrize("anyio_backend", ["asyncio"])
@pytest.mark.parametrize("fault", ["none", "source", "missing", "file"])
async def test_real_pdf_and_word_share_only_current_overview_facts(tmp_path, fault):
    import pypdf
    from docx import Document

    from smart_reporting.reporting.delivery.report_runtime.runtime import ReportRuntime

    from .test_report_runtime_full_render import _DEPENDENCIES_OK, _SKIP_REASON, _build_workspace

    if not _DEPENDENCIES_OK:
        pytest.skip(_SKIP_REASON)
    _, _, index, original, sources, questions, plan = await case(tmp_path / "proof")
    workspace = tmp_path / "workspace"
    state = _build_workspace(workspace)
    state["_editorExportSettings"]["sources"] = True
    markdown = ("# 智能运营全链路渲染验证报告\n\n[[section:section_002]]\n\n## 1. Report\n\n"
                + original)
    (workspace / "report.md").write_text(markdown, encoding="utf-8")
    quality = evaluate(index, markdown, sources, questions, plan)
    assert quality.management_overview.rows[0].status == "met"
    state["_documentContext"].update({
        "sections": [{"code": "section_002", "sectionNumber": "1", "title": "Report"}],
        "headingNumbers": [{"level": 2, "number": "1", "title": "Report",
                            "sectionCode": "section_002", "anchor": "report-section-section_002"}],
    })
    # An approved evidence subject may reuse only a server-registered source.
    subject = quality.management_overview.rows[0].evidence[0].subject_ids[0]
    state["_citationPresentations"].append({"citationId": "approved", "subjectId": subject,
        "label": "北院收入冻结数据", "coverageItems": [{"label": "收入", "periods": ["2025-10"]}]})
    payload = management_overview_projection(None if fault == "missing" else quality.management_overview,
        hashlib.sha256(markdown.encode()).hexdigest(), state["_citationPresentations"])
    if fault == "source":
        payload["sourceSha256"] = "f" * 64
    elif fault == "missing":
        payload["markdown"] = "未核验收入999999万元，不得进入总览。"
    raw = json.dumps(payload, ensure_ascii=False).encode()
    (workspace / "overview.json").write_bytes(raw if fault != "file" else b"{}")
    state["_managementOverviewProjectionFile"] = {"path": "overview.json", "size": len(raw),
                                                "sha256": hashlib.sha256(raw).hexdigest()}
    runtime = ReportRuntime(workspace)
    result = runtime.render_markdown(state, "report.md", "out/report.pdf",
        ".reporting-tmp/workspace-report-overview-render/render.pdf")
    assert result["status"] == "rendered"
    pdf = pypdf.PdfReader(workspace / result["pdfPath"])
    word = Document(workspace / result["wordPath"])
    pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    word_text = "\n".join(p.text for p in word.paragraphs) + "\n" + "\n".join(
        cell.text for table in word.tables for row in table.rows for cell in row.cells)
    for text in (pdf_text, word_text):
        compact = "".join(text.split())
        assert "核心结论" in compact
        assert "智能运营验证机构" in compact
        assert "999999" not in compact
        if fault == "none":
            assert "2025-10" in compact and "2400.0万元" in compact
            assert "该项已核验" in compact and "该项待核实" in compact
            assert "北院收入冻结数据" in compact
            assert "暂不列示汇总数值" not in compact
        else:
            assert "暂不列示汇总数值" in compact
            assert "该项已核验" not in compact
    assert result["render"]["projection"]["overviewCheck"]["status"] == (
        "unknown" if fault in {"source", "file"} else "verified")
    state["render"] = result["render"]
    validation = runtime.validate_pdf(state, result["pdfPath"],
        ".reporting-tmp/workspace-report-overview-validate", word_path=result["wordPath"])
    assert validation["projectionCheck"]["status"] == "verified"
    if fault in {"source", "file"}:
        assert validation["ok"] is False
        assert any(issue["code"] == "management_overview_projection_unknown" for issue in validation["issues"])
    else:
        assert validation["ok"] is True
    assert (workspace / result["pdfPath"]).is_file() and (workspace / result["wordPath"]).is_file()
